"""识别格式不固定的 Excel：找表头行、展开合并单元格、剔除合计行。"""
from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

import openpyxl

TOTAL_WORDS = ("合计", "小计", "总计", "总额")
SCAN_ROWS = 20


@dataclass
class SheetTable:
    file: str
    sheet: str
    header_row: int                         # Excel 行号
    headers: list[str]
    rows: list[tuple[int, list]] = field(default_factory=list)    # (Excel 行号, 值)
    totals: list[tuple[int, list]] = field(default_factory=list)  # 合计行，用于对账
    confidence: float = 1.0

    def column(self, header: str) -> list:
        i = self.headers.index(header)
        return [vals[i] for _, vals in self.rows]

    def sample(self, n: int = 5) -> dict[str, list]:
        return {h: [vals[i] for _, vals in self.rows[:n]] for i, h in enumerate(self.headers) if h}


def _grid(ws) -> list[list]:
    """读出整个 sheet，并把合并单元格的值填到每个格子里。"""
    grid = [list(r) for r in ws.iter_rows(values_only=True)]
    for rng in ws.merged_cells.ranges:
        v = ws.cell(rng.min_row, rng.min_col).value
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                if r - 1 < len(grid) and c - 1 < len(grid[r - 1]):
                    grid[r - 1][c - 1] = v
    return grid


def _is_text(v) -> bool:
    if not isinstance(v, str):
        return False
    s = v.strip().replace(",", "").rstrip("%")
    try:
        float(s)
        return False
    except ValueError:
        return bool(s)


def _filled(row) -> int:
    return sum(1 for v in row if v not in (None, ""))


def _is_total(row) -> bool:
    return any(isinstance(v, str) and any(w in v for w in TOTAL_WORDS) for v in row[:3])


def detect_sheet(ws, file: str) -> SheetTable | None:
    grid = _grid(ws)
    if not grid:
        return None
    best, best_score = None, 0.0
    for i, row in enumerate(grid[:SCAN_ROWS]):
        n = _filled(row)
        # 合并单元格展开的标题行（如跨整表的表名、说明）所有格子是同一个值，不是表头
        if n < 2 or len({str(v).strip() for v in row if v not in (None, "")}) < 2:
            continue
        text = sum(1 for v in row if _is_text(v)) / n
        nxt = next((r for r in grid[i + 1:i + 4] if _filled(r) >= 2), None)
        data_like = 0.0
        if nxt:
            data_like = sum(1 for v in nxt if v not in (None, "") and not _is_text(v)) / max(_filled(nxt), 1)
        score = text * n * (0.5 + data_like)
        if score > best_score:
            best, best_score = i, score
    if best is None:
        return None
    rows = [best]
    # 两行表头：上一行是分组标题（如「金额（万元）」跨两列），下一行是具体列名
    # 只有上一行像分组标题（有合并展开后的重复值，或下一行有值处上一行为空）才算两行表头，
    # 避免把全是文字的数据行（如 部门编号 | 部门名称）当成表头
    if best + 1 < len(grid):
        head, nxt = grid[best], grid[best + 1]
        vals = [v for v in head if v not in (None, "")]
        grouped = len(set(vals)) < len(vals) or any(
            (head[i] if i < len(head) else None) in (None, "") and v not in (None, "") for i, v in enumerate(nxt))
        if grouped and _filled(nxt) >= 2 and all(_is_text(v) or v in (None, "") for v in nxt):
            rows.append(best + 1)
    t = _build(grid, file, ws.title, rows)
    t.confidence = min(1.0, best_score / max(len(t.headers), 1))
    return t


def _build(grid, file: str, title: str, header_idx: list[int]) -> SheetTable:
    """按指定的表头行（0 起的下标，可多行）组装表格；多行表头用“上级_下级”合并。"""
    header = [str(v).strip() if v is not None else "" for v in grid[header_idx[0]]]
    for i in header_idx[1:]:
        nxt = list(grid[i]) + [None] * max(0, len(header) - len(grid[i]))
        header = [f"{a}_{str(b).strip()}" if b not in (None, "") and str(b).strip() != a else a
                  for a, b in zip(header, nxt)]
    start = header_idx[-1] + 1
    width = len(header)
    t = SheetTable(file, title, header_idx[0] + 1, header)
    for j, row in enumerate(grid[start:], start=start + 1):
        row = (list(row) + [None] * width)[:width]
        if _filled(row) == 0:
            continue
        (t.totals if _is_total(row) else t.rows).append((j, row))
    return t


def table_at(path: str | Path, sheet: str, header_rows: list[int]) -> SheetTable:
    """按人工（或智能体）指定的表头行读取工作表，header_rows 为 Excel 行号，如 [3] 或 [3, 4]。"""
    path = Path(path)
    wb = openpyxl.load_workbook(path, data_only=True)
    if sheet not in wb.sheetnames:
        raise KeyError(f"{path.name} 中没有工作表「{sheet}」")
    grid = _grid(wb[sheet])
    if not header_rows or max(header_rows) > len(grid) or min(header_rows) < 1:
        raise ValueError(f"表头行 {header_rows} 超出范围（工作表共 {len(grid)} 行）")
    return _build(grid, path.name, sheet, [r - 1 for r in sorted(header_rows)])


def detect(path: str | Path) -> list[SheetTable]:
    path = Path(path)
    wb = openpyxl.load_workbook(path, data_only=True)
    out = []
    for ws in wb.worksheets:
        t = detect_sheet(ws, path.name)
        if t is not None and t.rows:
            out.append(t)
    return out
