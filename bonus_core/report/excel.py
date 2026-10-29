"""结果 Excel：按分配树动态生成 sheet。结果只写数值不写公式，避免打开后被 Excel 重算；
只有「过程复核」表用公式，让财务在 Excel 里自己重算一遍，和引擎结果逐个核对。"""
from __future__ import annotations

from decimal import Decimal
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from bonus_core.dsl.models import RuleSet
from bonus_core.engine.run import Result

MONEY = "#,##0.00"
HEAD_FILL = PatternFill("solid", fgColor="DDEBF7")
LOCK_FILL = PatternFill("solid", fgColor="FFF2CC")


def _sheet(wb, title, headers, rows, money_cols=(), widths=None):
    ws = wb.create_sheet(title[:31])
    ws.append(headers)
    for c in ws[1]:
        c.font, c.fill = Font(bold=True), HEAD_FILL
        c.alignment = Alignment(horizontal="center")
    for r in rows:
        ws.append(r)
    for col in money_cols:
        for (cell,) in ws.iter_rows(min_row=2, min_col=col, max_col=col):
            cell.number_format = MONEY
    for i, h in enumerate(headers, start=1):
        w = (widths or {}).get(h) or max(10, min(60, max(len(str(h)) * 2, *(len(str(r[i - 1])) if len(r) >= i else 0 for r in rows) if rows else [0])))
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    return ws


def _d(s) -> Decimal:
    return Decimal(str(s))


def write_excel(rs: RuleSet, res: Result, path: str | Path, ledger: list[dict] | None = None,
                extra_sheets: list[tuple[str, list[str], list[list], dict]] | None = None) -> Path:
    """extra_sheets：(标题, 表头, 行, 列宽) 列表，放在「汇总」之后（如智能体的规划说明、待确认问题）。"""
    path = Path(path)
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    nodes = res.nodes
    leaf = res.levels[-1]

    # 汇总
    total_held = sum((_d(p.held) for p in res.payouts.values()), Decimal(0))
    total_paid = sum((_d(p.paid) for p in res.payout_rows()), Decimal(0)) if res.payouts else _d(res.pool)
    rows = [
        ["规则", f"{res.rule_set}（{res.rule_version}）"],
        ["奖金池（精确值）", float(_d(res.pool_exact))],
        ["奖金池（保留到分）", float(_d(res.pool))],
        ["本期实发合计", float(total_paid)],
        ["本期挂起合计", float(total_held)],
        ["调整次数", len([a for a in res.adjustments if not a.revoked])],
    ]
    for lv in res.levels:
        cnt = sum(1 for n in nodes.values() if n.level == lv)
        rows.append([f"「{lv}」节点数", cnt])
    ws = _sheet(wb, "汇总", ["项目", "值"], rows, widths={"项目": 22, "值": 30})
    for r in (3, 4, 5, 6):
        ws.cell(r, 2).number_format = MONEY
    for title, headers, extra_rows, widths in extra_sheets or []:
        _sheet(wb, title, headers, extra_rows, widths=widths)

    # 各层级明细
    for lv in res.levels:
        lv_rows = []
        for n in nodes.values():
            if n.level != lv:
                continue
            p = nodes.get(n.parent)
            t = res.traces.get(f"{n.id}/权重")
            lv_rows.append([
                n.id, n.label, p.label if p else "", float(_d(n.weight)),
                float(_d(n.exact)), float(_d(n.final)), "是" if n.locked else "", t.rendered if t else "",
            ])
        ws = _sheet(wb, f"{lv}明细", ["节点", "名称", "上级", "权重", "精确金额", "金额（到分）", "已调整锁定", "权重算式"],
                    lv_rows, money_cols=(5, 6), widths={"权重算式": 80})
        for i, r in enumerate(lv_rows, start=2):
            if r[6]:
                for c in ws[i]:
                    c.fill = LOCK_FILL

    _review_sheets(wb, rs, res)

    # 个人发放
    conds = [c.id for c in rs.payout]
    pay_rows = []
    for nid, p in res.payouts.items():
        n = nodes[nid]
        pay_rows.append([nid, n.label, nodes[n.parent].label, float(_d(p.final)), float(_d(p.held)),
                         float(_d(p.cap or p.paid)), float(_d(p.offset)), float(_d(p.paid)), float(_d(p.overpaid))]
                        + [float(_d(p.holds[c])) for c in conds])
    _sheet(wb, f"{leaf}发放", ["节点", "名称", "上级", "应得", "暂缓", "可发上限", "已发抵扣", "本次实发", "超发待核对"]
           + [f"挂起：{c}" for c in conds], pay_rows, money_cols=range(4, 10 + len(conds)))
    if res.rollups:
        ent = next(iter(res.rollups.values())).entity
        _sheet(wb, f"按{ent}汇总发放", ["节点", "名称", "包含", "应得", "暂缓", "可发上限", "已发抵扣", "实际抵扣", "本次实发", "超发待核对"],
               [[r.node_id, r.label, "、".join(r.members), float(_d(r.final)), float(_d(r.held)), float(_d(r.cap)),
                 float(_d(r.offset)), float(_d(r.deducted)), float(_d(r.paid)), float(_d(r.overpaid))]
                for r in res.rollups.values()], money_cols=range(4, 11), widths={"包含": 40})

    # 调整记录
    adj_rows = [[a.created_at, a.node_id, nodes[a.node_id].label if a.node_id in nodes else "", float(a.amount),
                 a.actor, a.reason, "已撤销" if a.revoked else "生效"] for a in res.adjustments]
    _sheet(wb, "调整记录", ["时间", "节点", "名称", "调整为", "调整人", "原因", "状态"], adj_rows, money_cols=(4,))

    # 未发放台账
    led = ledger if ledger is not None else [
        {"node_id": nid, "condition": c, "held": p.holds[c], "rule_version": res.rule_version}
        for nid, p in res.payouts.items() for c in conds if _d(p.holds[c]) > 0
    ]
    _sheet(wb, "未发放台账", ["节点", "名称", "挂起条件", "挂起金额", "规则版本"],
           [[e["node_id"], nodes[e["node_id"]].label, e["condition"], float(_d(e["held"])), e["rule_version"]] for e in led],
           money_cols=(4,))

    # 计算链
    _sheet(wb, "计算链", ["追踪 ID", "规则元素", "公式", "代入后算式", "结果", "依赖", "明细行"],
           [[t.id, t.rule_item, t.expr, t.rendered, t.value, ", ".join(t.inputs), ", ".join(t.rows[:20])]
            for t in res.traces.values()], widths={"公式": 50, "代入后算式": 60, "依赖": 40, "明细行": 40})

    # 异常与版本
    _sheet(wb, "异常与提示", ["级别", "代码", "说明", "位置"],
           [[i.severity, i.code, i.message, i.where] for i in res.issues], widths={"说明": 70, "位置": 40})
    _sheet(wb, "规则与数据版本", ["项目", "值"], [
        ["规则集", res.rule_set], ["规则版本", res.rule_version], ["引擎版本", res.engine_version],
        ["输入哈希（规则+数据+调整）", res.input_hash],
        ["数据哈希（规则+数据）", res.data_hash],
    ], widths={"项目": 14, "值": 70})

    wb.save(path)
    _self_check(path, res)
    return path


GREY = PatternFill("solid", fgColor="EDEDED")
GREEN = PatternFill("solid", fgColor="E2EFDA")
RED = PatternFill("solid", fgColor="F8CBAD")


def _cell_value(v):
    if isinstance(v, Decimal):
        return float(v)
    return v


def _label(rs: RuleSet, name: str) -> str:
    v = rs.variables.get(name)
    if v is not None:
        return f"{v.desc}（{name}）" if v.desc else name
    for ds in rs.datasets.values():
        f = ds.fields.get(name)
        if f is not None:
            return f"{f.desc}（{name}）" if f.desc else name
    return name


def _review_sheets(wb, rs: RuleSet, res: Result) -> None:
    """每个层级一张「过程复核」表：每个对象一行；灰色列是引擎取数（对明细的汇总、字段值），绿色列是 Excel 公式，
    最后用公式复算应得（或权重），和引擎结果比对。参数表另起一张，公式引用它。"""
    from openpyxl.formatting.rule import CellIsRule, FormulaRule

    from bonus_core.dsl import excel_formula as xf
    plans = xf.plan(rs)
    if not plans:
        return
    used = set()
    for p in plans.values():
        for src in [rs.variables[n].expr for n in p["computed"]] + [p[k] for k in ("pool", "weight") if k in p]:
            used |= {c.args[0].id for c in xf.walk(xf.parse(src)) if isinstance(c, xf.Call) and c.fn in ("tier", "tier_rate", "lookup")}
    params = xf.Params()
    params.place(rs, used)
    adjusted = {a.node_id for a in res.adjustments if not a.revoked}
    for lv, p in plans.items():
        ns = [n for n in res.nodes.values() if n.level == lv]
        if not ns:
            continue
        cols = ["编号", "名称"] + p["inputs"] + p["computed"]
        target = "pool" if "pool" in p else "weight" if "weight" in p else None
        if target:
            cols += ["@复算", "@引擎", "@核对"]
        letter = {c: get_column_letter(i) for i, c in enumerate(cols, start=1)}
        ws = wb.create_sheet(f"{lv}过程复核"[:31])
        kinds = ["", ""] + ["取数（引擎）"] * len(p["inputs"]) + ["公式"] * len(p["computed"])
        heads = ["编号", "名称"] + [_label(rs, c) for c in p["inputs"] + p["computed"]]
        if target:
            what = "应得" if target == "pool" else "权重"
            kinds += ["复核", "复核", "复核"]
            heads += [f"{what}（Excel 公式复算）", f"{what}（引擎结果）", "核对"]
        ws.append(kinds)
        ws.append(heads)
        for i, (k, h) in enumerate(zip(kinds, heads), start=1):
            for r in (1, 2):
                c = ws.cell(r, i)
                c.font = Font(bold=True)
                c.alignment = Alignment(horizontal="center", wrap_text=True)
                c.fill = GREY if k.startswith("取数") else GREEN if k == "公式" else HEAD_FILL
            ws.column_dimensions[get_column_letter(i)].width = max(12, min(30, len(h) * 2))
        for r, n in enumerate(ns, start=3):
            vals = res.inputs.get(n.id, {})
            f = xf.Formula(rs, params, lambda name, r=r: f"{letter[name]}{r}" if name in letter else "NA()")
            row = [n.key or n.id, n.label] + [_cell_value(vals.get(c)) for c in p["inputs"]]
            row += [f(rs.variables[c].expr) for c in p["computed"]]
            if target:
                re_ = f(p[target])
                re_ = f"=ROUND({re_[1:]},2)" if target == "pool" else re_
                engine = float(_d(n.final if target == "pool" else n.weight))
                a, b = letter["@复算"], letter["@引擎"]
                tol = "0.005" if target == "pool" else f"0.000001*MAX(1,ABS({b}{r}))"
                check = "人工调整" if n.id in adjusted and target == "pool" else \
                    f'=IF(ISERROR({a}{r}),"公式出错",IF(ABS({a}{r}-{b}{r})<{tol},"一致","不一致"))'
                row += [re_, engine, check]
            ws.append(row)
        last = 2 + len(ns)
        lo, hi = 3 + len(p["inputs"]), len(cols) - (1 if target else 0)    # 公式所在的列（不含核对列）
        first, end = get_column_letter(lo), get_column_letter(max(lo, hi))
        if hi >= lo:
            ws.conditional_formatting.add(f"{first}3:{end}{last}", FormulaRule(formula=[f"ISERROR({first}3)"], fill=RED))
        if target:
            k = letter["@核对"]
            for bad in ('"不一致"', '"公式出错"'):
                ws.conditional_formatting.add(f"{k}3:{k}{last}", CellIsRule(operator="equal", formula=[bad], fill=RED))
            for name in ("@复算", "@引擎"):
                col = cols.index(name) + 1
                for (cell,) in ws.iter_rows(min_row=3, max_row=last, min_col=col, max_col=col):
                    cell.number_format = MONEY if target == "pool" else "0.000000"
        ws.append([])
        if target:
            ws.append(["复核未通过的个数", f'=COUNTIF({k}3:{k}{last},"不一致")+COUNTIF({k}3:{k}{last},"公式出错")'
                                         f'+SUMPRODUCT(--ISERROR({k}3:{k}{last}))'])
        if hi >= lo:
            ws.append(["公式出错的单元格", f"=SUMPRODUCT(--ISERROR({first}3:{end}{last}))"])
        ws.freeze_panes = "C3"
    if params.rows:
        ps = wb.create_sheet("参数表")
        for r in params.rows:
            ps.append(r)
        ps.column_dimensions["A"].width, ps.column_dimensions["B"].width, ps.column_dimensions["C"].width = 50, 16, 16


def _self_check(path: Path, res: Result) -> None:
    """重新读取文件，核对各层级合计与快照一致。"""
    wb = openpyxl.load_workbook(path, data_only=True)
    pool = _d(res.pool)
    for lv in res.levels:
        ws = wb[f"{lv}明细"[:31]]
        s = sum((_d(round(r[5], 2)) for r in ws.iter_rows(min_row=2, values_only=True)), Decimal(0))
        if s.quantize(Decimal("0.01")) != pool:
            raise AssertionError(f"结果 Excel 自检失败：「{lv}」合计 {s} ≠ 奖金池 {pool}")
