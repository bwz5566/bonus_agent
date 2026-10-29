"""给结果文件里的公式单元格写入计算值（缓存值）。

openpyxl 只写公式、不写计算结果：Excel 打开时会重算，但不重算的程序（上传到其他系统、pandas、openpyxl data_only）
读到的是空单元格。这里按 Excel 的规则把本项目生成的公式（过程复核表、参数表引用）真正算一遍，把结果写进文件，
Excel 打开时仍会整本重算。只支持本项目生成的公式用到的函数；遇到不认识的函数，该单元格不写缓存值。
"""
from __future__ import annotations

import re
import zipfile
from datetime import date, datetime
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path
from xml.sax.saxutils import escape, unescape

import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter


class XlErr(Exception):
    """Excel 的错误值（#N/A、#DIV/0!、#VALUE!）。作为值传递，也可以抛出。"""

    def __init__(self, code: str):
        super().__init__(code)
        self.code = code

    def __eq__(self, other):
        return isinstance(other, XlErr) and other.code == self.code

    def __hash__(self):
        return hash(self.code)


class Unsupported(Exception):
    pass


_TOKEN = re.compile(r"""
    (?P<ws>\s+)
  | (?P<str>"(?:[^"]|"")*")
  | (?P<ref>(?:'(?:[^']|'')+'!|[A-Za-z_][\w.]*!)?\$?[A-Z]{1,3}\$?\d+(?::\$?[A-Z]{1,3}\$?\d+)?)
  | (?P<num>\d+(?:\.\d+)?(?:[Ee][+-]?\d+)?)
  | (?P<name>[A-Za-z_][\w.]*)
  | (?P<op><=|>=|<>|[-+*/^&=<>(),])
""", re.X)


def _tokens(src: str) -> list[tuple[str, str]]:
    out, i = [], 0
    while i < len(src):
        m = _TOKEN.match(src, i)
        if not m:
            raise Unsupported(f"无法解析：{src[i:i + 20]}")
        i = m.end()
        if m.lastgroup != "ws":
            out.append((m.lastgroup, m.group()))
    return out


_BIN = {"=": 1, "<>": 1, "<": 1, ">": 1, "<=": 1, ">=": 1, "&": 2, "+": 3, "-": 3, "*": 4, "/": 4, "^": 5}


class _Parser:
    def __init__(self, src: str):
        self.t, self.i = _tokens(src), 0

    def peek(self):
        return self.t[self.i] if self.i < len(self.t) else ("end", "")

    def take(self, val=None):
        k, v = self.peek()
        if val is not None and v != val:
            raise Unsupported(f"缺少 {val}")
        self.i += 1
        return k, v

    def expr(self, prec=0):
        left = self.unary()
        while True:
            k, v = self.peek()
            p = _BIN.get(v) if k == "op" else None
            if p is None or p <= prec:          # 同级左结合
                return left
            self.take()
            left = ("bin", v, left, self.expr(p))

    def unary(self):
        k, v = self.peek()
        if k == "op" and v in "+-":
            self.take()
            return ("neg", self.unary()) if v == "-" else self.unary()
        return self.atom()

    def atom(self):
        k, v = self.take()
        if k == "num":
            return ("num", float(v))
        if k == "str":
            return ("str", v[1:-1].replace('""', '"'))
        if k == "ref":
            return ("ref", v)
        if k == "op" and v == "(":
            e = self.expr()
            self.take(")")
            return e
        if k == "name":
            up = v.upper()
            if up in ("TRUE", "FALSE") and self.peek()[1] != "(":
                return ("bool", up == "TRUE")
            self.take("(")
            args = []
            if self.peek()[1] != ")":
                args.append(self.expr())
                while self.peek()[1] == ",":
                    self.take()
                    args.append(self.expr())
            self.take(")")
            return ("fn", up, args)
        raise Unsupported(f"无法解析：{v}")


def _split_ref(ref: str, sheet: str) -> tuple[str, str]:
    if "!" in ref:
        s, a = ref.rsplit("!", 1)
        sheet = s[1:-1].replace("''", "'") if s.startswith("'") else s
        ref = a
    return sheet, ref.replace("$", "")


def _cells(rng: str) -> list[str]:
    if ":" not in rng:
        return [rng]
    a, b = rng.split(":")
    ca, ra = re.match(r"([A-Z]+)(\d+)", a).groups()
    cb, rb = re.match(r"([A-Z]+)(\d+)", b).groups()
    return [f"{get_column_letter(c)}{r}" for r in range(int(ra), int(rb) + 1)
            for c in range(column_index_from_string(ca), column_index_from_string(cb) + 1)]


def _serial(d) -> float:
    if isinstance(d, datetime):
        return (d - datetime(1899, 12, 30)).total_seconds() / 86400
    return float((d - date(1899, 12, 30)).days)


def _text(v) -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, float):
        if v == int(v) and abs(v) < 1e15:
            return str(int(v))
        return format(Decimal(format(v, ".15g")).normalize(), "f")
    return str(v)


def _num(v) -> float:
    if isinstance(v, XlErr):
        raise v
    if v is None or v == "":
        return 0.0
    if isinstance(v, bool):
        return 1.0 if v else 0.0
    if isinstance(v, (int, float)):
        return float(v)
    raise XlErr("#VALUE!")


def _round(x: float, n: int) -> float:
    q = Decimal(1).scaleb(-n)
    return float(Decimal(format(x, ".15g")).quantize(q, ROUND_HALF_UP))   # Excel 按 15 位有效数字四舍五入（远离 0）


def _cmp(op: str, a, b) -> bool:
    for x in (a, b):
        if isinstance(x, XlErr):
            raise x

    def key(v):     # Excel：数字 < 文字 < 逻辑值；空白按对方类型
        if isinstance(v, bool):
            return (2, v)
        if isinstance(v, str):
            return (1, v.lower())
        return (0, float(v))
    if a is None:
        a = "" if isinstance(b, str) else False if isinstance(b, bool) else 0.0
    if b is None:
        b = "" if isinstance(a, str) else False if isinstance(a, bool) else 0.0
    x, y = key(a), key(b)
    return {"=": x == y, "<>": x != y, "<": x < y, ">": x > y, "<=": x <= y, ">=": x >= y}[op]


class Book:
    """按 Excel 规则计算工作簿里的公式。"""

    def __init__(self, wb):
        self.wb = wb
        self.memo: dict[tuple[str, str], object] = {}
        self.busy: set[tuple[str, str]] = set()

    def value(self, sheet: str, coord: str):
        key = (sheet, coord)
        if key in self.memo:
            return self.memo[key]
        if sheet not in self.wb.sheetnames:
            raise XlErr("#REF!")
        v = self.wb[sheet][coord].value
        if isinstance(v, str) and v.startswith("="):
            if key in self.busy:
                raise Unsupported("循环引用")
            self.busy.add(key)
            try:
                v = self.formula(v[1:], sheet)
            except XlErr as e:
                v = e
            finally:
                self.busy.discard(key)
        elif isinstance(v, (datetime, date)):
            v = _serial(v)
        elif isinstance(v, int) and not isinstance(v, bool):
            v = float(v)
        self.memo[key] = v
        return v

    def formula(self, src: str, sheet: str):
        p = _Parser(src)
        e = p.expr()
        if p.peek()[0] != "end":
            raise Unsupported("多余内容")
        return self.ev(e, sheet)

    def ev(self, e, sheet, array=False):
        kind = e[0]
        if kind in ("num", "str", "bool"):
            return e[1]
        if kind == "ref":
            s, r = _split_ref(e[1], sheet)
            cells = _cells(r)
            if ":" not in r:
                return self.value(s, cells[0])
            return [self.value(s, c) for c in cells]
        if kind == "neg":
            v = self.ev(e[1], sheet, array)
            if isinstance(v, list):
                return [x if isinstance(x, XlErr) else -_num(x) for x in v]
            return -_num(v)
        if kind == "bin":
            op, a, b = e[1], self.ev(e[2], sheet), self.ev(e[3], sheet)
            if op in ("=", "<>", "<", ">", "<=", ">="):
                return _cmp(op, a, b)
            if op == "&":
                for x in (a, b):
                    if isinstance(x, XlErr):
                        raise x
                return _text(a) + _text(b)
            x, y = _num(a), _num(b)
            if op == "/" and y == 0:
                raise XlErr("#DIV/0!")
            if op == "+":
                return x + y
            if op == "-":
                return x - y
            if op == "*":
                return x * y
            if op == "/":
                return x / y
            return x ** y
        if kind == "fn":
            return self.call(e[1], e[2], sheet)
        raise Unsupported(kind)

    def call(self, fn: str, args: list, sheet: str):
        if fn == "IF":
            c = self.ev(args[0], sheet)
            if isinstance(c, XlErr):
                raise c
            ok = bool(_num(c)) if not isinstance(c, str) else (_ for _ in ()).throw(XlErr("#VALUE!"))
            if ok:
                return self.ev(args[1], sheet)
            return self.ev(args[2], sheet) if len(args) > 2 else False
        if fn == "ISERROR":
            try:
                v = self.ev(args[0], sheet, array=True)
            except XlErr:
                return True
            if isinstance(v, list):
                return [isinstance(x, XlErr) for x in v]
            return isinstance(v, XlErr)
        if fn == "NA":
            raise XlErr("#N/A")
        vals = [self.ev(a, sheet, array=fn in ("SUMPRODUCT", "COUNTIF", "MIN", "MAX", "SUM")) for a in args]
        if fn in ("MIN", "MAX", "SUM"):
            flat = []
            for v in vals:
                if isinstance(v, list):
                    for x in v:
                        if isinstance(x, XlErr):
                            raise x
                        if isinstance(x, (int, float)) and not isinstance(x, bool):
                            flat.append(float(x))
                else:
                    flat.append(_num(v))
            if fn == "SUM":
                return sum(flat)
            return (min if fn == "MIN" else max)(flat) if flat else 0.0
        if fn == "ROUND":
            return _round(_num(vals[0]), int(_num(vals[1])))
        if fn == "ABS":
            return abs(_num(vals[0]))
        if fn == "NOT":
            return not bool(_num(vals[0]))
        if fn == "DATE":
            y, m, d = (int(_num(v)) for v in vals)
            return _serial(date(y, m, d))
        if fn == "COUNTIF":
            rng, crit = vals
            crit = _text(crit).lower()
            return float(sum(1 for x in rng if not isinstance(x, XlErr) and _text(x).lower() == crit))
        if fn == "SUMPRODUCT":
            arr = vals[0] if isinstance(vals[0], list) else [vals[0]]
            return sum(_num(x) for x in arr)
        raise Unsupported(f"函数 {fn}")


def compute(path: str | Path) -> dict[tuple[str, str], object]:
    """工作簿里每个公式单元格的计算值：(工作表, 单元格) -> 值（数字、文字、逻辑值或 XlErr）。不支持的公式不在结果里。"""
    wb = openpyxl.load_workbook(path)
    book = Book(wb)
    out = {}
    for ws in wb:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    try:
                        out[(ws.title, c.coordinate)] = book.value(ws.title, c.coordinate)
                    except Unsupported:
                        continue
    return out


def _cell_xml(attrs: str, f: str, v) -> str | None:
    attrs = re.sub(r'\s+t="[^"]*"', "", attrs)
    if isinstance(v, XlErr):
        return f'<c{attrs} t="e">{f}<v>{v.code}</v></c>'
    if isinstance(v, bool):
        return f'<c{attrs} t="b">{f}<v>{int(v)}</v></c>'
    if isinstance(v, str):
        return f'<c{attrs} t="str">{f}<v>{escape(v)}</v></c>'
    if isinstance(v, (int, float)):
        return f'<c{attrs}>{f}<v>{float(v)!r}</v></c>'
    return None


_FORMULA_CELL = re.compile(r'<c( r="([A-Z]+\d+)"[^>]*)>(<f>.*?</f>)(?:<v>[^<]*</v>|<v\s*/>)?</c>', re.S)


def fill_cached(path: str | Path) -> dict[tuple[str, str], object]:
    """把公式的计算值写进文件（缓存值），返回计算值。Excel 打开时是否整本重算由写文件的一方设置（calcPr fullCalcOnLoad）。"""
    path = Path(path)
    values = compute(path)
    with zipfile.ZipFile(path) as z:
        items = [(i, z.read(i.filename)) for i in z.infolist()]
    data = {i.filename: d for i, d in items}
    rels = {}
    for tag in re.findall(r"<Relationship\b[^>]*>", data["xl/_rels/workbook.xml.rels"].decode()):
        attr = dict(re.findall(r'(\w+)="([^"]*)"', tag))
        rels[attr.get("Id")] = attr.get("Target", "")
    sheet_of = {}
    for tag in re.findall(r"<sheet\b[^>]*>", data["xl/workbook.xml"].decode()):
        attr = dict(re.findall(r'([\w:]+)="([^"]*)"', tag))
        name, target = attr.get("name", ""), rels.get(attr.get("r:id"), "")
        sheet_of[target.lstrip("/") if target.startswith("/xl/") else "xl/" + target] = unescape(name, {"&quot;": '"'})

    def fill(m: re.Match, sheet: str) -> str:
        v = values.get((sheet, m.group(2)), None)
        out = _cell_xml(m.group(1), m.group(3), v) if (sheet, m.group(2)) in values else None
        return out or m.group(0)
    tmp = path.with_name(path.stem + ".tmp" + path.suffix)
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as z:
        for info, d in items:
            sheet = sheet_of.get(info.filename)
            if sheet is not None:
                d = _FORMULA_CELL.sub(lambda m: fill(m, sheet), d.decode("utf-8")).encode("utf-8")
            z.writestr(info, d)
    tmp.replace(path)
    return values
