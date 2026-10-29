"""交付结算奖回归用例（虚构数据）：一人参与多个项目、按参与记录分配后按人汇总抵扣预发；
V2 补充收款 / 费用（表头行、列名、单位、取值写法都和首轮不同）、项目长表更正、评分修订（核对旧值）、迟延规则修订；
V3 在 V2 上锁定两条参与记录的应得。

make(dir) 生成材料（01~04）和独立手算的期望结果 expected.xlsx（不经过引擎，不交给智能体）。
"""
from __future__ import annotations

import random
from datetime import date
from decimal import ROUND_FLOOR, ROUND_HALF_UP, Decimal as D
from pathlib import Path

import openpyxl

RULES = """# 2026 年下半年项目交付结算奖（虚构测试制度）

## 1. 数据范围
考核期为2026年7月1日至12月31日，首尾日期均包含。同一员工可以参与多个项目，先按参与编号分配，再按工号汇总，不得去重。
所有工作表第4行为表头。

## 2. 项目奖金包
只有“验收已批准=是”且实际验收日在考核期内的项目才有奖金，否则为0元。
回款取截至2026年12月31日（含）已审批通过的本合同累计收款，负收款是退款要扣除。净回款=max(min(合格收款合计,合同额),0)。
费用取截至同一截止日已审批通过的全部费用，成本费用表单位是万元。贡献毛利=max(净回款-合格费用合计,0)。
贡献毛利按超额累进计提：前100,000元（含）部分5%；超过100,000至300,000元（含）部分8%；超过300,000元部分10%。
交付评级系数：优=1.15，合格=1.00，待改进=0.80。
迟延天数=max(实际验收日-计划验收日,0)。迟延系数=max(0.70,1-迟延天数×0.02)。
项目奖金包=ROUND(min(分段计提额×交付评级系数×迟延系数,合同额×4%),2)。

## 3. 项目内分配
有效工时为考核期内且审批通过的小时数。角色系数：项目经理1.20，实施工程师1.00，顾问1.10，质量工程师0.90。
只有“用工性质=正式”且“参与资格=参与”的记录可以分配，其余权重为0。
原始权重=有效工时×角色系数×负责人确认贡献分÷100。按原始权重分到参与记录，整数分最大余数法。

## 4. 发放
“质保回访完成=是”的项目可发比例100%，否则80%。每条参与记录：可发上限=ROUND(应得×可发比例,2)，暂缓=应得-可发上限。
同一员工跨项目汇总可发上限后抵扣本批已预发：抵扣=min(可发上限合计,已预发)，本次可发=可发上限合计-抵扣，超发待核对=max(已预发-可发上限合计,0)。

## 5. 版本
V1只用首轮数据。V2按补充通知处理，V3在V2基础上执行人工指定，三个版本分开输出、不得相加。
"""

NOTICE = """# 补充通知

## V2
1. 补充收款按“流水编号”并入首轮收款流水：合同项目号即项目编号，到账日即入账日期，复核结果“通过/驳回”即“是/否”。同号重复提交只算一次。
2. 补充费用按“凭证编号”并入成本费用，金额单位为元。
3. 项目更正按项目编号和更正字段替换原值，新值是完整替代值。
4. 评分修订按参与编号覆盖负责人确认贡献分，先核对原贡献分，对不上要报错。
5. 迟延规则修订：V2迟延系数=max(0.80,1-迟延天数×0.01)。其他规则不变。
补充文件所有工作表第5行为表头。

## V3
在V2基础上执行人工指定：J01项目下U002的应得锁定为6,000.00元；J03项目下U008的应得锁定为12,000.00元。
项目总包不变，其余参与者按原始权重分剩余金额（整数分最大余数法）。
"""

S, E = date(2026, 7, 1), date(2026, 12, 31)
P = [("J01", "仓储升级", 2000000, "是", date(2026, 9, 15), date(2026, 9, 10), "优", "是"),
     ("J02", "门店联网", 800000, "是", date(2026, 11, 20), date(2026, 12, 1), "合格", "否"),
     ("J03", "生产看板", 1500000, "是", date(2026, 12, 20), date(2026, 10, 1), "待改进", "是"),
     ("J04", "报表系统", 600000, "否", date(2026, 10, 8), date(2026, 10, 1), "合格", "是"),
     ("J05", "门户升级", 500000, "是", date(2026, 6, 20), date(2026, 6, 30), "优", "是")]
EMP = [("U001", "测试01", "正式"), ("U002", "测试02", "正式"), ("U003", "测试03", "正式"), ("U004", "测试04", "正式"),
       ("U005", "测试05", "正式"), ("U006", "测试06", "正式"), ("U007", "测试07", "实习"), ("U008", "测试08", "正式"),
       ("U009", "测试09", "正式")]
PART = [("J01", "U001", "项目经理", 95, "参与"), ("J01", "U002", "实施工程师", 88, "参与"), ("J01", "U003", "顾问", 76, "参与"),
        ("J01", "U007", "实施工程师", 90, "参与"), ("J02", "U001", "顾问", 70, "参与"), ("J02", "U004", "项目经理", 92, "参与"),
        ("J02", "U005", "质量工程师", 80, "不参与"), ("J02", "U008", "实施工程师", 85, "参与"),
        ("J03", "U002", "实施工程师", 81, "参与"), ("J03", "U006", "项目经理", 90, "参与"), ("J03", "U008", "实施工程师", 77, "参与"),
        ("J03", "U003", "质量工程师", 0, "参与"), ("J04", "U004", "项目经理", 90, "参与"), ("J05", "U006", "项目经理", 88, "参与")]
REC = [("J01", date(2026, 3, 1), 600000, "是"), ("J01", date(2026, 9, 20), 900000, "是"), ("J01", date(2026, 10, 5), -50000, "是"),
       ("J01", date(2026, 11, 1), 120000, "否"), ("J01", date(2027, 1, 3), 200000, "是"),
       ("J02", date(2026, 8, 1), 500000, "是"), ("J02", date(2026, 12, 30), 400000, "是"),
       ("J03", date(2026, 12, 25), 700000, "是"), ("J04", date(2026, 10, 9), 600000, "是"), ("J05", date(2026, 7, 2), 500000, "是")]
EXP = [("J01", date(2026, 5, 1), 30.5, "是"), ("J01", date(2026, 9, 1), 42.25, "是"), ("J01", date(2026, 10, 1), -3.1, "是"),
       ("J01", date(2026, 10, 2), 8, "否"), ("J02", date(2026, 9, 9), 51.38, "是"),
       ("J03", date(2026, 11, 11), 35.6, "是"), ("J03", date(2026, 12, 1), 12.4, "是"), ("J04", date(2026, 10, 1), 20, "是")]
PRE = [("U001", 3000), ("U002", 50000), ("U004", 2000.5), ("U006", 8000), ("U009", 1200)]
# V2 补充资料（流水编号 R002 是首轮已有的单号、内容相同：重传不重复累加）
SUP_REC = [("R002", "J01", date(2026, 9, 20), 900000, "通过"), ("S001", "J03", date(2026, 12, 28), 300000, "通过"),
           ("S002", "J03", date(2026, 12, 29), 50000, "驳回"), ("S003", "J01", date(2027, 1, 2), 80000, "通过")]
SUP_EXP = [("V001", "J02", date(2026, 10, 10), -20000, "通过"), ("V002", "J01", date(2026, 12, 1), 15000, "通过")]
FIX = [("J02", "交付评级", "优", "V2"), ("J03", "计划验收日", "2026-11-20", "V2"), ("J01", "交付评级", "优", "V2")]
SCORE = [("J01-U003", 76, 86, "张三", "复核"), ("J02-U008", 85, 60, "李四", "复核")]
LOCKS = [("J01-U002", D("6000.00")), ("J03-U008", D("12000.00"))]


def _sheet(wb, name, head, rows, header_row):
    ws = wb.create_sheet(name)
    ws.append([f"{name}（虚构测试数据）"])
    for _ in range(header_row - 2):
        ws.append([])
    ws.append(head)
    for r in rows:
        ws.append(list(r))


def _timesheets():
    rnd, ts = random.Random(7), []
    for pj, e, *_ in PART:
        pid = f"{pj}-{e}"
        for _ in range(3):
            ts.append((pid, date(2026, 7 + rnd.randrange(6), 1 + rnd.randrange(28)), rnd.choice([16, 24, 32, 40, 37.5]), "是"))
        ts.append((pid, date(2026, 6, 28), 20, "是"))
        ts.append((pid, date(2026, 8, 3), 12, "否"))
    return ts


def make(out: Path) -> Path:
    out = Path(out)
    (out / "materials").mkdir(parents=True, exist_ok=True)
    (out / "materials" / "01_交付奖金规则.md").write_text(RULES, encoding="utf-8")
    (out / "materials" / "03_补充通知.md").write_text(NOTICE, encoding="utf-8")
    ts = _timesheets()
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    _sheet(wb, "项目台账", ["项目编号", "项目名称", "不含税合同额（元）", "验收已批准", "实际验收日", "计划验收日", "交付评级", "质保回访完成"], P, 4)
    _sheet(wb, "人员档案", ["工号", "姓名", "用工性质"], EMP, 4)
    _sheet(wb, "参与分配", ["参与编号", "项目编号", "工号", "角色", "负责人确认贡献分", "参与资格"],
           [(f"{p}-{e}", p, e, r, s, q) for p, e, r, s, q in PART], 4)
    _sheet(wb, "工时明细", ["工时单号", "参与编号", "工作日期", "小时数", "审批通过"], [(f"T{i + 1:04d}", *t) for i, t in enumerate(ts)], 4)
    _sheet(wb, "收款流水", ["收款单号", "项目编号", "入账日期", "收款金额（元）", "审批通过"], [(f"R{i + 1:03d}", *r) for i, r in enumerate(REC)], 4)
    _sheet(wb, "成本费用", ["费用单号", "项目编号", "发生日期", "费用金额（万元）", "审批通过"], [(f"F{i + 1:03d}", *r) for i, r in enumerate(EXP)], 4)
    _sheet(wb, "已预发", ["工号", "本批已预发（元）"], PRE, 4)
    wb.save(out / "materials" / "02_首轮业务数据.xlsx")
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    _sheet(wb, "补充收款", ["流水编号", "合同项目号", "到账日", "到账数（元）", "复核结果"], SUP_REC, 5)
    _sheet(wb, "补充费用", ["凭证编号", "合同项目号", "业务日期", "补记费用（元）", "复核结果"], SUP_EXP, 5)
    _sheet(wb, "项目更正", ["项目编号", "更正字段", "新值", "生效范围"], FIX, 5)
    _sheet(wb, "评分修订", ["参与编号", "原贡献分", "新贡献分", "负责人", "原因"], SCORE, 5)
    _sheet(wb, "人工调整", ["参与编号", "指定应得（元）"], [(k, float(v)) for k, v in LOCKS], 5)
    wb.save(out / "materials" / "04_补充业务数据.xlsx")
    _expected(out / "expected.xlsx", ts)
    return out


# ---------------- 独立手算 ----------------
C = D("0.01")


def _r2(x):
    return x.quantize(C, ROUND_HALF_UP)


def _tier(x):
    return min(x, D(100000)) * D("0.05") + max(min(x, D(300000)) - 100000, 0) * D("0.08") + max(x - 300000, 0) * D("0.10")


def _split(cents_total: int, weights: dict[str, D]) -> dict[str, int]:
    tot = sum(weights.values())
    if not weights or cents_total == 0:
        return {k: 0 for k in weights}
    assert tot > 0
    ex = {k: D(cents_total) * w / tot for k, w in weights.items()}
    fl = {k: int(ex[k].to_integral_value(ROUND_FLOOR)) for k in weights}
    for k in sorted(weights, key=lambda k: (-(ex[k] - fl[k]), k))[:cents_total - sum(fl.values())]:
        fl[k] += 1
    return fl


def _calc(version: str, ts):
    projects = {p[0]: list(p) for p in P}
    rec = [(f"R{i + 1:03d}", p, d, D(str(a)), k) for i, (p, d, a, k) in enumerate(REC)]
    exp = [(p, d, D(str(a)) * 10000, k) for p, d, a, k in EXP]
    scores = {f"{p}-{e}": s for p, e, _, s, _ in PART}
    if version != "V1":
        ids = {r[0] for r in rec}
        rec += [(i, p, d, D(str(a)), "是" if k == "通过" else "否") for i, p, d, a, k in SUP_REC if i not in ids]
        exp += [(p, d, D(str(a)), "是" if k == "通过" else "否") for _, p, d, a, k in SUP_EXP]
        for pj, f, v, _ in FIX:
            if f == "交付评级":
                projects[pj][6] = v
            else:
                projects[pj][5] = date.fromisoformat(v)
        for pid, old, new, *_ in SCORE:
            assert scores[pid] == old
            scores[pid] = new
    rate, floor = (D("0.02"), D("0.70")) if version == "V1" else (D("0.01"), D("0.80"))
    pool = {}
    for pid, _, contract, ok, act, plan, grade, war in projects.values():
        if ok != "是" or not (S <= act <= E):
            pool[pid] = D(0)
            continue
        net = max(min(sum((a for _, p, d, a, k in rec if p == pid and k == "是" and d <= E), D(0)), D(contract)), D(0))
        cost = sum((a for p, d, a, k in exp if p == pid and k == "是" and d <= E), D(0))
        gm = max(net - cost, D(0))
        dc = max(floor, 1 - D(max((act - plan).days, 0)) * rate)
        coef = {"优": D("1.15"), "合格": D(1), "待改进": D("0.80")}[grade]
        pool[pid] = _r2(min(_tier(gm) * coef * dc, D(contract) * D("0.04")))
    etype = {e: t for e, _, t in EMP}
    role = {"项目经理": D("1.2"), "实施工程师": D(1), "顾问": D("1.1"), "质量工程师": D("0.9")}
    weight, owner = {}, {}
    for pj, e, r, _, q in PART:
        pid = f"{pj}-{e}"
        hrs = sum((D(str(h)) for p, d, h, k in ts if p == pid and k == "是" and S <= d <= E), D(0))
        weight[pid] = hrs * role[r] * D(scores[pid]) / 100 if etype[e] == "正式" and q == "参与" else D(0)
        owner[pid] = (pj, e)
    locks = dict(LOCKS) if version == "V3" else {}
    amt = {}
    for pj in pool:
        ms = [k for k in weight if owner[k][0] == pj]
        rest = int(pool[pj] * 100) - sum(int(locks[k] * 100) for k in ms if k in locks)
        got = _split(rest, {k: weight[k] for k in ms if k not in locks})
        amt.update({k: D(v) / 100 for k, v in got.items()})
        amt.update({k: locks[k] for k in ms if k in locks})
    war = {p: v[7] for p, v in projects.items()}
    pers = {}
    for k, v in amt.items():
        pj, e = owner[k]
        cap = _r2(v * (1 if war[pj] == "是" else D("0.8")))
        a = pers.setdefault(e, [D(0)] * 3)
        a[0], a[1], a[2] = a[0] + v, a[1] + v - cap, a[2] + cap
    pre = {e: D(str(x)) for e, x in PRE}
    rows = []
    for e, *_ in EMP:
        f, h, cap = pers.get(e, [D(0)] * 3)
        p = pre.get(e, D(0))
        d = min(cap, p)
        rows.append([e, f, h, p, d, cap - d, max(p - cap, D(0))])
    return pool, amt, rows


def _expected(path: Path, ts):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for v, sheet in (("V1", "V1"), ("V2", "V2-R0"), ("V3", "V2-R1")):
        pool, amt, rows = _calc(v, ts)
        ws = wb.create_sheet(f"{sheet}-项目")
        ws.append(["项目编号", "期望应得"])
        for k, x in pool.items():
            ws.append([k, float(x)])
        ws = wb.create_sheet(f"{sheet}-参与记录")
        ws.append(["参与编号", "期望应得"])
        for k, x in sorted(amt.items()):
            ws.append([k, float(x)])
        ws = wb.create_sheet(f"{sheet}-个人")
        ws.append(["工号", "期望应得", "暂缓", "已预发", "实际抵扣", "本次可发", "超发待核对"])
        for r in rows:
            ws.append([r[0]] + [float(x) for x in r[1:]])
    wb.save(path)


if __name__ == "__main__":
    import sys
    make(Path(sys.argv[1]))
