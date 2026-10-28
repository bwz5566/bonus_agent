"""生成示例数据 Excel：故意带标题行、合并单元格、万元单位和合计行，模拟人工导出的表。

运行：python examples/sales_q3/make_data.py
"""
import random
from pathlib import Path

import openpyxl

HERE = Path(__file__).parent
rng = random.Random(20260930)

wb = openpyxl.Workbook()

# 人员表：第 1 行是标题，第 2 行空，第 3 行才是表头
ws = wb.active
ws.title = "人员名单"
ws["A1"] = "销售部 2026 年第三季度人员名单"
ws.merge_cells("A1:E1")
ws.append([])
ws.append(["工号", "姓名", "所属小组", "工龄（年）", "付出程度评分"])
groups = {"G01": "华东一组", "G02": "华南二组", "G03": "西北三组"}
people = []
names = "赵钱孙李周吴郑王冯陈褚卫蒋沈韩杨朱秦尤许"
for i, ch in enumerate(names):
    eid = f"E{1001 + i}"
    gid = list(groups)[i % 3]
    people.append((eid, gid))
    ws.append([eid, ch + "某", gid, rng.randint(1, 12), rng.choice([3, 3.5, 4, 4.5, 5])])

# 小组表
ws = wb.create_sheet("小组绩效")
ws.append(["小组编号", "小组名称", "部门绩效得分"])
for gid, gname, score in [("G01", groups["G01"], 92), ("G02", groups["G02"], 85), ("G03", groups["G03"], 78)]:
    ws.append([gid, gname, score])

# 销售明细：两行表头（「金额（万元）」跨三列），末尾有合计行
ws = wb.create_sheet("销售明细")
ws.append(["合同编号", "负责人工号", "项目类型", "区域", "金额（万元）", None, None, "毛利率"])
ws.append([None, None, None, None, "合同金额", "净利润", "已回款", None])
for col in "ABCDH":
    ws.merge_cells(f"{col}1:{col}2")
ws.merge_cells("E1:G1")
region_of = {"G01": "华东", "G02": "华南", "G03": "西北"}
tot_s = tot_p = tot_c = 0.0
for k in range(60):
    eid, gid = people[k % len(people)]
    amount = round(rng.uniform(15, 80), 2)
    margin = round(rng.uniform(0.12, 0.45), 3)
    profit = round(amount * margin * 0.8, 2)
    collected = amount if rng.random() < 0.7 else round(amount * rng.choice([0, 0.3, 0.5, 0.8]), 2)
    tot_s, tot_p, tot_c = tot_s + amount, tot_p + profit, tot_c + collected
    ptype = "新项目" if rng.random() < 0.4 else "老项目"
    ws.append([f"HT2026{k + 1:04d}", eid, ptype, region_of[gid], amount, profit, collected, margin])
ws.append(["合计", None, None, None, round(tot_s, 2), round(tot_p, 2), round(tot_c, 2), None])

out = HERE / "data.xlsx"
wb.save(out)
print(f"已生成 {out}，销售额合计 {tot_s:.2f} 万元，净利润合计 {tot_p:.2f} 万元")
