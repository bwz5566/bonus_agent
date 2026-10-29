"""把规则文档（Word / PPT / PDF / Excel / 图片）解析为有序的内容块。

- 文字和表格由程序直接读取，不经过大模型；
- PPT 里的原生图表直接读出数据序列，比看图更准；
- 图片（截图、扫描件、导出的图表）留给多模态模型转写。
"""
from __future__ import annotations

import mimetypes
from dataclasses import dataclass, field
from pathlib import Path


@dataclass
class Block:
    doc: str
    page: int | None
    kind: str                 # text | table | chart | image | transcript
    text: str = ""
    image: bytes | None = None
    mime: str | None = None

    def label(self) -> str:
        loc = f"{self.doc}" + (f" 第{self.page}页" if self.page else "")
        return f"[{loc} · {self.kind}]"


@dataclass
class ParsedDoc:
    name: str
    blocks: list[Block] = field(default_factory=list)

    def full_text(self) -> str:
        return "\n".join(b.text for b in self.blocks if b.text)


def md_table(rows: list[list]) -> str:
    rows = [["" if v is None else str(v).replace("\n", " ").strip() for v in r] for r in rows if any(v not in (None, "") for v in r)]
    if not rows:
        return ""
    w = max(len(r) for r in rows)
    rows = [r + [""] * (w - len(r)) for r in rows]
    lines = ["| " + " | ".join(rows[0]) + " |", "|" + " --- |" * w]
    lines += ["| " + " | ".join(r) + " |" for r in rows[1:]]
    return "\n".join(lines)


def parse_docx(path: Path) -> ParsedDoc:
    import docx
    from docx.oxml.ns import qn

    d = docx.Document(str(path))
    out = ParsedDoc(path.name)
    body = d.element.body
    for el in body.iterchildren():
        if el.tag == qn("w:p"):
            p = docx.text.paragraph.Paragraph(el, d)
            txt = p.text.strip()
            if txt:
                style = (p.style.name or "") if p.style is not None else ""
                prefix = "#" * int(style[-1]) + " " if style.startswith("Heading") and style[-1:].isdigit() else ""
                out.blocks.append(Block(path.name, None, "text", prefix + txt))
            for blip in el.iter(qn("a:blip")):
                rid = blip.get(qn("r:embed"))
                part = d.part.related_parts.get(rid)
                if part is not None:
                    out.blocks.append(Block(path.name, None, "image", image=part.blob, mime=part.content_type))
        elif el.tag == qn("w:tbl"):
            t = docx.table.Table(el, d)
            out.blocks.append(Block(path.name, None, "table", md_table([[c.text for c in r.cells] for r in t.rows])))
    return out


def parse_pptx(path: Path) -> ParsedDoc:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    prs = Presentation(str(path))
    out = ParsedDoc(path.name)

    def visit(shape, page):
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            for s in shape.shapes:
                visit(s, page)
            return
        if getattr(shape, "has_chart", False) and shape.has_chart:
            ch = shape.chart
            cats = list(ch.plots[0].categories) if ch.plots else []
            rows = [["类别"] + [s.name for s in ch.series]]
            for i, c in enumerate(cats):
                rows.append([c] + [s.values[i] for s in ch.series])
            title = ch.chart_title.text_frame.text if ch.has_title else ""
            out.blocks.append(Block(path.name, page, "chart", (f"图表：{title}\n" if title else "") + md_table(rows)))
        elif getattr(shape, "has_table", False) and shape.has_table:
            rows = [[c.text for c in r.cells] for r in shape.table.rows]
            out.blocks.append(Block(path.name, page, "table", md_table(rows)))
        elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            out.blocks.append(Block(path.name, page, "image", image=shape.image.blob, mime=shape.image.content_type))
        elif getattr(shape, "has_text_frame", False) and shape.has_text_frame:
            txt = shape.text_frame.text.strip()
            if txt:
                out.blocks.append(Block(path.name, page, "text", txt))

    for i, slide in enumerate(prs.slides, start=1):
        for shape in slide.shapes:
            visit(shape, i)
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                out.blocks.append(Block(path.name, i, "text", f"备注：{notes}"))
    return out


def parse_xlsx(path: Path, max_rows: int = 200) -> ParsedDoc:
    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=True)
    out = ParsedDoc(path.name)
    for i, ws in enumerate(wb.worksheets, start=1):
        rows = [list(r) for r in ws.iter_rows(values_only=True)][:max_rows]
        t = md_table(rows)
        if t:
            out.blocks.append(Block(path.name, i, "table", f"工作表：{ws.title}\n{t}"))
    return out


SCAN_CHARS = 30          # PDF 一页读出的文字少于这么多，当作扫描件，整页转成图片交给多模态模型转写
FIGURE_SHARE = 0.15      # 页面里图片占页面面积超过这个比例（截图、图表），整页另附一张图片


def parse_pdf(path: Path, dpi: int = 150) -> ParsedDoc:
    """文字型 PDF 由程序直接读出文字和表格；扫描件、含截图或图表的页面整页转成图片，交给多模态模型转写。"""
    import pdfplumber

    out = ParsedDoc(path.name)
    render = []
    with pdfplumber.open(str(path)) as pdf:
        for i, page in enumerate(pdf.pages, start=1):
            tables = page.find_tables()
            boxes = [t.bbox for t in tables]

            def outside(obj, boxes=boxes):
                x = (obj.get("x0", 0) + obj.get("x1", 0)) / 2
                y = (obj.get("top", 0) + obj.get("bottom", 0)) / 2
                return not any(b[0] <= x <= b[2] and b[1] <= y <= b[3] for b in boxes)
            text = (page.filter(outside).extract_text() or "").strip() if boxes else (page.extract_text() or "").strip()
            if text:
                out.blocks.append(Block(path.name, i, "text", text))
            for t in tables:
                md = md_table(t.extract())
                if md:
                    out.blocks.append(Block(path.name, i, "table", md))
            area = float(page.width * page.height) or 1.0
            pics = sum(abs((im["x1"] - im["x0"]) * (im["bottom"] - im["top"])) for im in page.images)
            if len(text) + sum(len(b.text) for b in out.blocks if b.page == i and b.kind == "table") < SCAN_CHARS \
                    or pics / area > FIGURE_SHARE:
                render.append(i)
    if render:
        import io

        import pypdfium2 as pdfium
        doc = pdfium.PdfDocument(str(path))
        try:
            for i in render:
                img = doc[i - 1].render(scale=dpi / 72).to_pil()
                buf = io.BytesIO()
                img.save(buf, format="PNG")
                out.blocks.append(Block(path.name, i, "image", image=buf.getvalue(), mime="image/png"))
        finally:
            doc.close()
    out.blocks.sort(key=lambda b: (b.page or 0, b.kind == "image"))
    return out


def parse_image(path: Path) -> ParsedDoc:
    mime = mimetypes.guess_type(path.name)[0] or "image/png"
    return ParsedDoc(path.name, [Block(path.name, None, "image", image=path.read_bytes(), mime=mime)])


def parse_text(path: Path) -> ParsedDoc:
    """纯文本类文件（txt / md / json / csv），如财务确认答复、参数说明。"""
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "gbk"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        text = raw.decode("utf-8", errors="replace")
    return ParsedDoc(path.name, [Block(path.name, None, "text", text)])


PARSERS = {".docx": parse_docx, ".pptx": parse_pptx, ".pdf": parse_pdf, ".xlsx": parse_xlsx, ".xlsm": parse_xlsx,
           ".png": parse_image, ".jpg": parse_image, ".jpeg": parse_image,
           ".txt": parse_text, ".md": parse_text, ".json": parse_text, ".csv": parse_text}


def parse_any(path: str | Path) -> ParsedDoc:
    path = Path(path)
    fn = PARSERS.get(path.suffix.lower())
    if fn is None:
        raise ValueError(f"暂不支持的文件类型：{path.suffix}（支持 {', '.join(PARSERS)}；.doc/.ppt 请先另存为新格式）")
    return fn(path)
