"""坏数据回归：把智能体做好的规则（不调用大模型）在一批改坏的数据上重算，看每种坏数据能不能被拦住。

坏数据用例的格式（negative_cases.json）：
  [{"case_id": "N001", "name": "销售主键重复",
    "patch": {"target": "data", "op": "append_copy" | "set", "table": "sales", "index": 0, "field": "...", "value": ...}}]
table 是数据表名；Excel 里的工作表名不一样时用 sheets 对照（如 {"sales": "销售明细"}）。
target 不是 data 的用例（如政策未确认）不在这里测，由智能体的阻断问题处理。
"""
from __future__ import annotations

import copy
import shutil
import tempfile
from dataclasses import dataclass
from pathlib import Path

import openpyxl


@dataclass
class NegResult:
    case_id: str
    name: str
    blocked: bool
    detail: str


def patch_workbook(src: Path, dst: Path, patch: dict, sheets: dict[str, str] | None = None) -> None:
    wb = openpyxl.load_workbook(src)
    ws = wb[(sheets or {}).get(patch["table"], patch["table"])]
    head = [c.value for c in ws[1]]
    r = patch["index"] + 2
    if patch["op"] == "append_copy":
        ws.append([c.value for c in ws[r]])
    elif patch["op"] == "set":
        v = patch["value"]
        if isinstance(v, str):
            try:
                v = float(v)
            except ValueError:
                pass
        ws.cell(r, head.index(patch["field"]) + 1).value = v
    else:
        raise ValueError(f"不支持的 op：{patch['op']}")
    wb.save(dst)


def _target(files: list[Path], patch: dict, sheets: dict[str, str] | None) -> Path:
    sheet = (sheets or {}).get(patch["table"], patch["table"])
    for f in files:
        if f.suffix.lower() in (".xlsx", ".xlsm"):
            if sheet in openpyxl.load_workbook(f, read_only=True).sheetnames:
                return f
    raise ValueError(f"材料里没有工作表「{sheet}」")


def replay(snapshot: dict, files: list[Path], cases: list[dict], sheets: dict[str, str] | None = None) -> list[NegResult]:
    """snapshot 是案卷 case.json 里的 workspace。每个用例都在一份复制的材料上改一处、重算全部方案。"""
    from bonus_agent import case as cs
    from bonus_agent.agent import ToolError, Workspace
    from bonus_agent.auto import build_inventory
    from bonus_agent.llm import FakeLLM

    out = []
    for c in cases:
        p = c.get("patch") or {}
        if p.get("target") != "data":
            continue
        with tempfile.TemporaryDirectory() as tmp:
            src = _target(files, p, sheets)
            copies = []
            for f in files:
                dst = Path(tmp) / f.name
                if f == src:
                    patch_workbook(f, dst, p, sheets)
                else:
                    shutil.copy(f, dst)
                copies.append(dst)
            ws = Workspace(FakeLLM([]), build_inventory(copies))
            cs.restore(ws, copy.deepcopy(snapshot))
            detail = ""
            try:
                for s in ws.schemes:
                    ws._trial(s)
            except ToolError as e:
                detail = " ".join(str(e).split())[:300]
            out.append(NegResult(c.get("case_id", ""), c.get("name", ""), bool(detail), detail or "没有拦住，照常算出了结果"))
    return out


def report(results: list[NegResult]) -> list[str]:
    n = sum(r.blocked for r in results)
    lines = [f"坏数据拦截 {n}/{len(results)}"]
    lines += [f"{'✓' if r.blocked else '✗'} {r.case_id} {r.name}：{r.detail}" for r in results]
    return lines
