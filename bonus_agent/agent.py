"""大模型主导的奖金计算智能体：不写死流程，模型通过工具自己决定看什么、怎么建模、怎么核对。

模型能做的事都封装成工具（看文件、看表、看图、定义数据表、写规则、试算、查算式、算表达式、提问、完成），
程序只负责：执行工具、校验输入、保存工作区状态。金额始终由计算引擎按模型写的规则算出。

长任务的上下文控制：工作区状态（笔记、已定义的数据表、规则、问题、试算结论）每轮重新附在请求末尾，
所以模型不会因为对话变长而“忘记”已经做过的事。对话历史超过上限时才整理一次：较早的工具结果只留摘要，
仍然太长就丢弃最早的几轮。两次整理之间历史只追加、不改动，请求前缀保持不变，便于前缀缓存命中
（私有化部署的 vLLM 自动缓存；百炼部分模型可在配置里开启显式缓存）。
"""
from __future__ import annotations

import json
import re
from decimal import Decimal
from dataclasses import dataclass, field
from pathlib import Path
from typing import Callable

import yaml
from pydantic import ValidationError

from bonus_agent.auto import (DATA_SUFFIXES, MAX_FILES, Inventory, PlanDataset, TooManyFiles, build_inventory,
                              check_dataset, collect_files, filter_sheet)
from bonus_agent import case as cs, failure
from bonus_agent import facts as fx
from bonus_agent.extract_rules import Extraction, Question, check_quotes, prompt
from bonus_agent.llm import LLM, LLMError, image_part, parse_json
from bonus_core.data.mapping import DatasetMapping, apply_mapping
from bonus_core.data.table_detect import table_at
from bonus_core.dsl.functions import describe as describe_functions
from bonus_core.dsl.models import DatasetDef, RuleSet
from bonus_core.dsl.validator import validate
from bonus_core.engine.allocation import AdjustConflict
from bonus_core.engine.evaluator import EvalError, Evaluator, fmt
from bonus_core.engine.issues import EngineError, Issue
from bonus_core.engine.run import Adjustment, Engine, Result, compute
from bonus_core.engine.tree import ROOT_LEVEL, build_tree
from bonus_core.report.excel import write_excel

MAX_RESULT_CHARS = 6000
DEFAULT_SCHEME = "默认"        # 只有一个方案时的方案名
SCHEME_REF = re.compile(r"""scheme_paid\(\s*["']([^"']+)["']\s*(?:,\s*["']([^"']+)["']\s*)?\)""")   # 规则里引用别的方案
SHARED = "共用"                # 多个方案共用的部分（权重、校验、共用变量），合并进每个方案
MAX_STUCK = 3                  # 同一调用原样重复、或同一工具连续失败达到这个次数就停下，把卡住的原因报给财务
KEEP_FULL_RESULTS = 6          # 整理历史时，最近几条工具结果保留全文
BRIEF_CHARS = 240              # 较早的工具结果保留的开头字数
READ_ONLY = {"list_files", "read_document", "read_sheet", "view_image", "explain", "evaluate"}
WRITE_TOOLS = {"define_dataset", "remove_dataset", "set_rules", "record_facts", "remove_facts", "ask_question",
               "answer_question"}
RULE_SUFFIXES = {".docx", ".pptx", ".png", ".jpg", ".jpeg", ".txt", ".md"}   # 规则文档：要登记参数
PIN_DOC_CHARS = 12000          # 不超过这么长的文档全文附在任务里（材料包），不必用工具读
PIN_TOTAL_CHARS = 40000        # 材料包里文档全文的总长上限
PIN_SHEET_ROWS = 30            # 不超过这么多行的工作表整表附上，更长的只附表头和前几行


class ToolError(Exception):
    """工具参数或执行出错，错误信息原样返回给模型。"""


# ---------------- 工具定义 ----------------
def _obj(props: dict, required: list[str] = ()) -> dict:
    return {"type": "object", "properties": props, "required": list(required)}


_S = {"type": "string"}
_I = {"type": "integer"}
FIELD_SCHEMA = _obj({"column": {**_S, "description": "Excel 表头原文"},
                     "type": {"type": "string", "enum": ["money", "number", "int", "ratio", "text", "enum", "date", "id"]},
                     "unit": {"type": "string", "enum": ["元", "万元"], "description": "money 字段必填"},
                     "desc": _S}, ["column", "type"])

TOOLS: list[dict] = [
    {"name": "list_files", "description": "列出本批次所有文件；Excel 列出每个工作表及自动识别的表头。",
     "parameters": _obj({})},
    {"name": "read_document", "description": "读取 Word / PPT / 图片文件的文字和表格（图片显示为占位符，用 view_image 查看）。",
     "parameters": _obj({"file": _S, "offset": {**_I, "description": "从第几个字符开始，默认 0"}}, ["file"])},
    {"name": "view_image", "description": "查看文件中的第 index 张图片（从 1 开始，默认 1），返回图片内容的逐字转写。",
     "parameters": _obj({"file": _S, "index": _I, "focus": {**_S, "description": "可选：希望重点看清的内容"}},
                        ["file"])},
    {"name": "read_sheet", "description": "查看 Excel 工作表的原始内容（带 Excel 行号），用于判断表头位置、合计行、说明文字等。",
     "parameters": _obj({"file": _S, "sheet": _S, "start_row": _I, "rows": _I}, ["file", "sheet"])},
    {"name": "note",
     "description": "把读到的关键信息记在工作区（如某条规则、某个参数、文件用途、待核对事项）。笔记每轮都会附给你，"
                    "较早的对话会被裁剪，重要内容请及时记下。同一 topic 再次记录会覆盖，text 为空则删除。",
     "parameters": _obj({"topic": _S, "text": _S}, ["topic", "text"])},
    {"name": "record_facts",
     "description": "登记从材料里读到的参数和口径（比例、金额、阶梯、系数、门槛、日期、取数口径等），程序据此比对冲突。"
                    "每个文件里出现的都要登记，同一参数在几个文件（或一个文件的几处）出现就登记几次，用同一个 name；"
                    "取值不一致时程序会自动向财务提问，你不要自行取舍。value 写原文的写法（如 5%、30万元、1.2），"
                    "quote 逐字复制包含该取值的原文（不超过 200 字），程序会核对 quote 在文件里、value 在 quote 里。",
     "parameters": _obj({"facts": {"type": "array", "items": _obj({
         "name": {**_S, "description": "参数名，写清适用范围，如 季度奖-部门净利润提成比例"},
         "value": {**_S, "description": "原文写法，一条只写一个值（阶梯、系数表每一档分开登记）"}, "file": _S, "quote": _S,
         "meaning": {**_S, "description": "可选：文字口径的规范说法，如 含税"},
         "override": {"type": "boolean", "description": "原文明确写了修改旧规定（修改为 / 以本通知为准）时为 true，"
                                                         "quote 要包含这句话"}}, ["name", "value", "file", "quote"])}},
         ["facts"])},
    {"name": "remove_facts",
     "description": "删掉台账里某个参数名的全部登记（登记错了时用，如把阶梯几档登记成了同一个名字）；依据它的、还没答复的冲突问题会被撤回。",
     "parameters": _obj({"name": _S}, ["name"])},
    {"name": "define_dataset",
     "description": "定义（或重新定义）一张数据表。实体表写 key（每行一个对象，如员工、小组，以及被别的表按编号引用的项目、销售单）；"
                    "明细表写 link（每行挂到某层级对象或某个实体上，如回款 link=sales_id）；公司级表 key、link 都不写。"
                    "同名字段表示关联。一张工作表混放多类数据时（如目标值表按 entity_type 区分部门/小组/个人），"
                    "用 filter 只取其中一类，并把编号列映射成对应的字段名（如 group_id ← entity_id）。返回校验结果和换算后的样例行。",
     "parameters": _obj({"name": {**_S, "description": "中文表名，如 人员、销售明细"}, "file": _S, "sheet": _S,
                         "header_rows": {"type": "array", "items": _I, "description": "可选：表头所在的 Excel 行号，多行表头写多个"},
                         "key": _S, "link": _S, "desc": _S,
                         "filter": {"type": "object", "description": '可选：只取这些行，{"列名": "取值"} 或 {"列名": ["取值1", "取值2"]}'},
                         "fields": {"type": "object", "additionalProperties": FIELD_SCHEMA,
                                    "description": "字段名（英文小写加下划线） -> 定义"}},
                        ["name", "file", "sheet", "fields"])},
    {"name": "remove_dataset", "description": "删除一张已定义的数据表。", "parameters": _obj({"name": _S}, ["name"])},
    {"name": "set_rules",
     "description": "写入或修改计算规则。只给出要改的部分：tables / variables 按名字合并（值为 null 表示删除），"
                    "pool、allocation、payout、offsets、checks、revisions 整体替换。返回静态校验结果。"
                    "一批文件里有多个方案（如季度奖、过程奖、年度奖）时，每个方案用 scheme 分别写 pool 等；"
                    "各方案相同的部分（分配层级与权重、数据校验、共用的参数表和变量）写到 scheme=\"共用\"，会合并进每个方案。"
                    "只有一个方案时不写 scheme。",
     "parameters": _obj({"scheme": {**_S, "description": "方案名，如 QUARTER；共用部分写 共用"},
                         "tables": {"type": "object"}, "variables": {"type": "object"}, "pool": {"type": "object"},
                         "allocation": {"type": "array", "items": {"type": "object"}},
                         "payout": {"type": "array", "items": {"type": "object"},
                                    "description": '发放条件：[{"id": "cash", "scope": "个人", "pay_ratio": "可发比例表达式"}] '
                                                   '或 hold_ratio（挂起比例）'},
                         "offsets": {"type": "array", "items": {"type": "object"},
                                     "description": '已发抵扣：[{"id": "prior", "scope": "个人", "expr": "已发金额表达式"}]，'
                                                    '可用 scheme_paid("QUARTER", "R2") 引用另一个方案某修订的本次实发'},
                         "checks": {"type": "array", "items": {"type": "object"}},
                         "revisions": {"type": "array", "items": {"type": "object"},
                                       "description": '人工调整的修订版本，后一个包含前面的调整：[{"name": "R1", "adjust": '
                                                      '[{"node": "小组:G01", "amount": 900000, "reason": "..."}]}]；R0 为不调整的基准'}})},
    {"name": "trial_run", "description": "用当前数据表和规则试算，返回奖金池、各层级金额分布、各修订的变化和提示。不写 scheme 时试算全部方案。",
     "parameters": _obj({"scheme": _S})},
    {"name": "explain", "description": "查看计算链：节点 id（如 个人:E1001、项目:P001、全部）或变量追踪 id 的公式、代入数值后的算式和结果。",
     "parameters": _obj({"id": _S, "scheme": _S, "revision": {**_S, "description": "默认 R0"}}, ["id"])},
    {"name": "evaluate", "description": "在某个节点上用引擎计算任意表达式（如 sum(net_profit)），用于核对数据和规则。"
                                        "节点可以是层级对象或关联实体（如 项目:P001）。需要先在 set_rules 里设定 allocation。",
     "parameters": _obj({"expr": _S, "node": {**_S, "description": "节点 id，默认 全部"}, "scheme": _S}, ["expr"])},
    {"name": "ask_question",
     "description": "向财务提问：材料之间有冲突、缺少必要的信息（参数、口径、数据）、或能有两种理解时使用。"
                    "不要自行假设或补出算法；提问后继续做不受影响的部分，问完后调用 finish 结束本轮，"
                    "财务答复后下一轮接着做。options 列出可选的理解，suggested 是你的建议（只供财务参考，不会被采用）。"
                    "related 写受影响的方案或规则（如 ANNUAL.offsets）。blocking=false 只用于不影响任何金额的提示。",
     "parameters": _obj({"text": _S, "options": {"type": "array", "items": _S}, "suggested": _I, "related": _S,
                         "blocking": {"type": "boolean"}}, ["text", "options"])},
    {"name": "reply",
     "description": "对话时用：回复财务这条消息并结束本次对话。改过规则、数据表或台账时，程序会先自动试算全部方案并做完成前检查，"
                    "不通过会把原因返回给你修正。回复里的金额要来自 evaluate / explain / 试算结果，不要自己算。",
     "parameters": _obj({"text": _S, "ignored": {"type": "array", "items": _obj({"ref": _S, "reason": _S})}}, ["text"])},
    {"name": "answer_question",
     "description": "对话时用：财务在消息里明确答复了某个待确认问题时，把他的原话记为该问题的答复。他没说清楚的不要替他选，用 reply 反问。",
     "parameters": _obj({"id": _S, "answer": {**_S, "description": "财务的原话或他选的选项"}}, ["id", "answer"])},
    {"name": "finish", "description": "完成。要求每个方案在最近一次修改之后都已成功试算。",
     "parameters": _obj({"summary": {**_S, "description": "两三句话说明：什么奖金、奖金池怎么来、怎么分、发放条件"},
                         "ignored": {"type": "array", "items": _obj({"ref": _S, "reason": _S})}}, ["summary"])},
]


def _fix_arguments(text: str) -> str:
    """模型输出的参数不是合法 JSON 时尽量修复（后面多了内容就取第一个对象）；修不好就换成说明错误的合法 JSON。
    历史里的参数必须是合法 JSON，否则服务端会拒绝之后的所有请求。"""
    try:
        json.loads(text or "{}")
        return text or "{}"
    except ValueError:
        pass
    raw = (text or "").strip()
    err, pos = "参数不是 JSON 对象", None
    for cand in (raw, re.sub(r",\s*([}\]])", r"\1", raw)):        # 原样；去掉多余的结尾逗号
        try:
            obj, _ = json.JSONDecoder().raw_decode(cand)
            if isinstance(obj, dict):
                return json.dumps(obj, ensure_ascii=False)
        except json.JSONDecodeError as e:
            if cand is raw:
                err, pos = e.msg, e.pos
    near = f"，出错位置附近：…{raw[max(0, pos - 60):pos + 40]}…" if pos is not None else ""
    return json.dumps({"_invalid_json": f"参数不是合法 JSON（{err}{near}），未执行。请重新调用：字符串用双引号，"
                                        "表达式里的双引号要转义成 \\\"，不要有注释和多余的逗号；参数很长时分几次调用"
                                        "（如 set_rules 先写 variables，再写 allocation、pool），每次少写一些"}, ensure_ascii=False)


def _coerce_args(tool: str, args: dict) -> dict:
    """常见的参数写法偏差改回来：数组、对象被写成 JSON 字符串（思考模式下常见）；tables / variables 只写了一项、
    把名字放进了 name 字段（{"name": "x", "expr": ...} → {"x": {"expr": ...}}）。"""
    props = next((t["parameters"]["properties"] for t in TOOLS if t["name"] == tool), {})
    out = dict(args)
    for k, v in args.items():
        if isinstance(v, str) and props.get(k, {}).get("type") in ("array", "object") and v.strip()[:1] in "[{":
            try:
                out[k] = json.loads(v)
            except ValueError:
                pass
    if tool == "set_rules":
        for k in ("tables", "variables"):
            v = out.get(k)
            if isinstance(v, dict) and isinstance(v.get("name"), str) and ("expr" in v or "kind" in v):
                item = {a: b for a, b in v.items() if a != "name"}
                out[k] = {v["name"]: item}
    return out


def openai_tools() -> list[dict]:
    return [{"type": "function", "function": t} for t in TOOLS]


# ---------------- 工作区 ----------------
@dataclass
class Step:
    tool: str
    args: dict
    result: str
    ok: bool


@dataclass
class Workspace:
    llm: LLM
    inv: Inventory
    datasets: dict[str, DatasetDef] = field(default_factory=dict)
    mapping: dict[str, DatasetMapping] = field(default_factory=dict)
    drafts: dict[str, dict] = field(default_factory=dict)  # 方案 -> 规则草稿（tables / variables / pool / allocation ...）
    questions: list[Question] = field(default_factory=list)
    transcripts: dict[str, str] = field(default_factory=dict)
    results: dict[str, dict[str, Result]] = field(default_factory=dict)   # 方案 -> 修订（R0 为基准）-> 结果
    untested: set[str] = field(default_factory=set)        # 最近修改后还没试算成功的方案
    summary: str = ""
    ignored: list[dict] = field(default_factory=list)
    notes: dict[str, str] = field(default_factory=dict)
    done: bool = False
    chat_by: str = ""                                       # 对话模式：发消息的财务；空表示普通的一轮处理
    changed: bool = False                                   # 对话里改过规则、数据表、台账或问题
    reply_text: str = ""
    budget: str = ""
    round: int = 1                                          # 第几轮（同一案卷和财务来回的次数）
    facts: list[fx.Fact] = field(default_factory=list)      # 参数台账：从材料里登记的参数和口径
    conflicts: dict[str, fx.Conflict] = field(default_factory=dict)   # 已提问的参数冲突
    _sheet_corpus: dict[str, str] = field(default_factory=dict)
    progress: tuple[int, int] = (0, 0)                      # 已用步数、上限
    pinned: set[str] = field(default_factory=set)           # 全文已附在材料包里的文件 / 工作表

    @property
    def data_files(self) -> list[Path]:
        return [f for f in self.inv.files if f.suffix.lower() in DATA_SUFFIXES]

    # ---- 文件 ----
    def _file(self, name: str) -> Path:
        hit = [f for f in self.inv.files if f.name == name or str(f) == name]
        if not hit:
            raise ToolError(f"没有文件「{name}」，可用文件：{[f.name for f in self.inv.files]}")
        return hit[0]

    def list_files(self) -> str:
        lines = []
        for f in self.inv.files:
            if f.suffix.lower() in DATA_SUFFIXES:
                lines.append(f"- {f.name}（Excel）")
                for ref, body in self.inv.sheet_text.items():
                    if ref.startswith(f.name + "/"):
                        st = self.inv.sheet(ref)
                        if st is None:
                            lines.append(f"  · 工作表「{ref.split('/', 1)[1]}」：未识别出规整表格（{body.count(chr(10)) + 1} 行），可用 read_sheet 查看")
                        else:
                            lines.append(f"  · 工作表「{st.sheet}」：表头在第 {st.header_row} 行，{len(st.rows)} 行数据"
                                         + (f"，{len(st.totals)} 行合计" if st.totals else "")
                                         + f"；表头 {[h for h in st.headers if h]}")
            else:
                d = next(d for d in self.inv.docs if d.name == f.name)
                imgs = sum(1 for b in d.blocks if b.kind == "image")
                lines.append(f"- {f.name}（{f.suffix[1:]}，约 {len(d.full_text())} 字" + (f"，{imgs} 张图片" if imgs else "") + "）")
        return "\n".join(lines)

    def read_document(self, file: str, offset: int = 0, limit: int = MAX_RESULT_CHARS - 200) -> str:
        f = self._file(file)
        d = next((d for d in self.inv.docs if d.name == f.name), None)
        if d is None:
            raise ToolError(f"{file} 是 Excel，请用 read_sheet")
        if f.name in self.pinned and limit < 10 ** 9:
            return f"{f.name} 的全文已在任务开头的【材料包】里（每轮都在），直接看那里，不必再读。"
        summary = _json_data_summary(f)
        if summary and limit < 10 ** 9:
            return (summary + "\n这是结构化数据文件，不能逐页读、也不能直接定义成数据表。"
                    "同样的数据在 Excel 里有时用 Excel 的工作表定义数据表；它是唯一来源时请用 ask_question 请财务提供 Excel。")
        parts, k = [], 0
        for b in d.blocks:
            if b.kind == "image":
                k += 1
                t = self.transcripts.get(f"{d.name}#{k}")
                parts.append(f"{b.label()} 图片 #{k}" + (f"（已转写）\n{t}" if t else "（用 view_image 查看）"
                                                          if self.llm.image_llm else "（当前模型不能看图，未转写）"))
            elif b.text:
                parts.append(f"{b.label()}\n{b.text}")
        text = "\n\n".join(parts)
        chunk = text[offset:offset + limit]
        more = len(text) - offset - len(chunk)
        return chunk + (f"\n\n……还有 {more} 字，用 offset={offset + len(chunk)} 继续读取" if more > 0 else "")

    def view_image(self, file: str, index: int = 1, focus: str = "") -> str:
        f = self._file(file)
        d = next((d for d in self.inv.docs if d.name == f.name), None)
        imgs = [b for b in d.blocks if b.kind == "image"] if d else []
        if not 1 <= index <= len(imgs):
            raise ToolError(f"{file} 中有 {len(imgs)} 张图片，index 应在 1~{len(imgs)}")
        b = imgs[index - 1]
        seer = self.llm.image_llm
        if seer is None:
            raise ToolError(f"当前模型 {self.llm.name} 不能看图，也没有配置看图的模型（models.yaml 里的 vision）。"
                            "图片里的规则如果影响计算，用 ask_question 请财务提供文字版，不要猜")
        ask = "请转写这张图片。" + (f"重点看清：{focus}" if focus else "")
        text = seer.chat([{"role": "system", "content": prompt("transcribe")},
                              {"role": "user", "content": [image_part(b.image, b.mime), {"type": "text", "text": ask}]}],
                             temperature=0).strip()
        self.transcripts[f"{d.name}#{index}"] = text
        return text

    def read_sheet(self, file: str, sheet: str, start_row: int = 1, rows: int = 40) -> str:
        f = self._file(file)
        if f"{f.name}/{sheet}" in self.pinned:
            return f"{f.name}/{sheet} 整表已在任务开头的【材料包】里（每轮都在），直接看那里，不必再读。"
        if f"{f.name}/{sheet}" not in self.inv.sheet_text:
            sheets = [r.split("/", 1)[1] for r in self.inv.sheet_text if r.startswith(f.name + "/")]
            raise ToolError(f"{f.name} 中没有工作表「{sheet}」，可选：{sheets}")
        import openpyxl
        ws = openpyxl.load_workbook(f, data_only=True, read_only=True)[sheet]
        start_row, rows = max(1, start_row), max(1, min(rows, 80))
        out, total = [], 0
        for i, r in enumerate(ws.iter_rows(values_only=True), start=1):
            total = i
            if start_row <= i < start_row + rows:
                cells = ["" if v is None else str(v).replace("\n", " ")[:30] for v in r]
                while cells and cells[-1] == "":
                    cells.pop()
                out.append(f"第{i}行 | " + " | ".join(cells))
        st = self.inv.sheet(f"{f.name}/{sheet}")
        head = (f"工作表共 {total} 行" + (f"，自动识别表头在第 {st.header_row} 行、{len(st.rows)} 行数据" if st else "，未自动识别出表格")
                + f"；以下为第 {start_row}~{min(start_row + rows - 1, total)} 行")
        end = start_row + rows - 1
        tail = (f"\n（还有第 {end + 1}~{total} 行，用 start_row={end + 1} 继续；数据表不必读完，define_dataset 会读取全部行）"
                if end < total else "\n（已到末尾）")
        return head + "\n" + "\n".join(out) + tail

    def transcribe_all(self):
        """开工前把所有图片转写好，材料包里直接给文字，模型不必逐张查看。"""
        if self.llm.image_llm is None:
            return
        for d in self.inv.docs:
            for k, _ in enumerate([b for b in d.blocks if b.kind == "image"], start=1):
                if f"{d.name}#{k}" not in self.transcripts:
                    try:
                        self.view_image(d.name, k)
                    except (LLMError, ToolError):
                        pass       # 转写失败的图片留给模型自己用 view_image 看

    def briefing(self) -> str:
        """材料包：短文档全文（含图片转写）、小工作表整表、大工作表的表头和前几行。附在任务里，每轮都在，
        较早的对话被裁剪也不会丢，模型不用反复读。"""
        self.pinned = set()
        parts, used = [], 0
        for d in self.inv.docs:
            text = self.read_document(d.name, 0, limit=10 ** 9)
            summary = _json_data_summary(self._file(d.name)) if len(text) > PIN_DOC_CHARS else ""
            if summary:
                parts.append(f"## 文件 {d.name}（结构化数据，共 {len(text)} 字，不要逐页读）\n{summary}")
                continue
            if len(text) <= PIN_DOC_CHARS and used + len(text) <= PIN_TOTAL_CHARS:
                used += len(text)
                self.pinned.add(d.name)
                parts.append(f"## 文件 {d.name}（全文）\n{_strip_tags(text, d.name)}")
            else:
                parts.append(f"## 文件 {d.name}（共 {len(text)} 字，太长未全附，以下为开头；需要时用 read_document 按 offset 读）\n"
                             + text[:1500])
        for f in self.data_files:
            for ref in [r for r in self.inv.sheet_text if r.startswith(f.name + "/")]:
                sheet = ref.split("/", 1)[1]
                full = self.read_sheet(f.name, sheet, 1, PIN_SHEET_ROWS + 1)
                if "（已到末尾）" in full:
                    self.pinned.add(ref)
                    parts.append(f"## 工作表 {ref}（整表）\n{full.split(chr(10), 1)[1].removesuffix('（已到末尾）').rstrip()}")
                else:
                    parts.append(f"## 工作表 {ref}（数据表，只附前 5 行；define_dataset 会读取全部行）\n"
                                 + self.read_sheet(f.name, sheet, 1, 6))
        return "\n\n".join(parts)

    def note(self, topic: str, text: str) -> str:
        if text.strip():
            self.notes[topic] = text.strip()
        else:
            self.notes.pop(topic, None)
        return f"已记录「{topic}」（共 {len(self.notes)} 条笔记，每轮都会附在工作区状态里）"

    # ---- 参数台账 ----
    def _lines(self, name: str) -> list[str]:
        """文件的原文，按行：文档文字、图片转写、Excel 每行单元格。"""
        f = self._file(name)
        if f.suffix.lower() in DATA_SUFFIXES:
            if f.name not in self._sheet_corpus:
                import openpyxl
                wb = openpyxl.load_workbook(f, data_only=True, read_only=True)
                self._sheet_corpus[f.name] = "\n".join(
                    " | ".join("" if v is None else str(v) for v in r) for w in wb.worksheets
                    for r in w.iter_rows(values_only=True))
            return self._sheet_corpus[f.name].splitlines()
        d = next(d for d in self.inv.docs if d.name == f.name)
        return (d.full_text() + "\n" + "\n".join(t for k, t in self.transcripts.items()
                                                   if k.rsplit("#", 1)[0] == f.name)).splitlines()

    def _corpus(self) -> dict[str, str]:
        """每个文件规范化后的全文，用来核对引用。"""
        return {f.name: fx.norm("".join(self._lines(f.name))) for f in self.inv.files}

    def record_facts(self, facts: list[dict]) -> str:
        corpus, lines, added, dup = self._corpus(), [], 0, 0
        for i, a in enumerate(facts, start=1):
            try:
                f = fx.Fact.model_validate({**a, "round": self.round}) if isinstance(a, dict) else None
            except ValidationError as e:
                lines.append(f"第 {i} 条格式不对：{str(e)[:200]}")
                continue
            if f is None:
                lines.append(f"第 {i} 条应为对象")
                continue
            bad = fx.check_fact(f, corpus)
            if bad:
                near = fx.nearest(f, self._lines(f.file)) if f.file in corpus and "找不到" in bad else ""
                lines.append(f"第 {i} 条「{f.name}」未登记：{bad}" + (f"。原文里最接近的一句：{near}（照这句改 quote，不必重读文件）"
                                                                      if near else ""))
                continue
            if any((g.name, g.file, fx.norm(g.quote)) == (f.name, f.file, fx.norm(f.quote)) for g in self.facts):
                dup += 1
                continue
            self.facts.append(f)
            added += 1
        if not added:       # 一条也没登记上：作为失败返回，避免模型以为成功、反复提交同样的内容
            raise ToolError("没有登记任何新参数。" + (f"{dup} 条之前已登记过，不要重复登记。" if dup else "") + "\n".join(lines)
                            + f"\n下一步（程序判断）：{self.next_hint()}")
        new = fx.detect(self.facts, self.conflicts, self.questions, self.round)
        out = [f"已登记 {added} 条（台账共 {len(self.facts)} 条）。" + (f"另有 {dup} 条之前已登记过，已跳过。" if dup else "")] + lines
        for n in fx.multi_valued(self.facts):
            out.append(f"注意：「{n}」登记了几个不同的取值。如果是阶梯的几档、几个区域的系数这类一组参数，请用 remove_facts 删掉它，"
                       "每档（每项）用不同的名字重新登记（如 …-第1档）；如果同一参数在材料里前后写法矛盾，用 ask_question 问财务")
        for n in fx.mixed_names(self.facts):
            out.append(f"注意：「{n}」下既有数值又有文字说明，像是两个不同的参数（如权重 10% 和归一化公式），"
                       "请换成两个名字分别登记（如 …-权重、…-公式）；同一参数的不同写法才用同一个名字")
        if new:
            out.append("程序发现取值冲突，已自动向财务提问（阻断，不要自行取舍，相关规则先按待定处理）：")
            out += [f"- {q.id} {q.text}" for q in new]
        return "\n".join(out)

    # ---- 数据表 ----
    def define_dataset(self, name: str, file: str, sheet: str, fields: dict, key: str | None = None,
                       link: str | None = None, desc: str = "", header_rows: list[int] | None = None,
                       filter: dict | None = None) -> str:
        f = self._file(file)
        try:
            st = table_at(f, sheet, header_rows) if header_rows else self.inv.sheet(f"{f.name}/{sheet}")
        except (KeyError, ValueError) as e:
            raise ToolError(str(e)) from None
        if st is None:
            raise ToolError(f"{f.name}/{sheet} 没有自动识别出表格，请先 read_sheet 查看，再用 header_rows 指定表头行")
        fixed = []
        if isinstance(fields, dict):     # key / link 写成了 Excel 列名（如 entity_id），而字段已改名（dept_id ← entity_id）
            for which, val in (("key", key), ("link", link)):
                hit = [f for f, d in fields.items() if isinstance(d, dict) and d.get("column") == val] if val and val not in fields else []
                if len(hit) == 1:
                    fixed.append(f"{which} 写的是列名 {val}，已改成对应的字段名 {hit[0]}（key / link 要写字段名）")
                    key, link = (hit[0], link) if which == "key" else (key, hit[0])
        try:
            pd = PlanDataset(sheet=sheet, key=key or None, link=link or None, desc=desc, fields=fields)
        except ValidationError as e:
            raise ToolError(f"参数格式不对：{e}") from None
        types = {fn: (fd.type, n) for n, d in self.datasets.items() if n != name for fn, fd in d.fields.items()}
        if filter is not None and not (isinstance(filter, dict) and all(isinstance(v, (str, int, float, list))
                                                                         for v in filter.values())):
            raise ToolError('filter 的格式是 {"列名": "取值"} 或 {"列名": ["取值1", "取值2"]}')
        flt = {str(k): ([str(x) for x in v] if isinstance(v, list) else str(v)) for k, v in (filter or {}).items()}
        d, m, warnings, errors = check_dataset(name, pd, st, types, header_rows, flt)
        if errors:
            owner = next((n for n, o in self.datasets.items() if n != name and key and o.key == key), None)
            if owner and any("重复值" in e for e in errors):
                errors.append(f"{key} 是「{owner}」的主键，这张表应写 link={key}（不写 key），每行挂到对应的{owner}上")
            raise ToolError("未保存，请修正：\n" + "\n".join(f"- {e}" for e in fixed + errors))
        warnings = fixed + warnings
        self.datasets[name], self.mapping[name] = d, m
        self.untested |= set(self.schemes)
        kind = f"实体表，主键 {d.key}" if d.key else f"明细表，关联 {d.link}" if d.link else "公司级表"
        sample = []
        from bonus_core.data.table import convert
        st, _ = filter_sheet(st, flt)
        for _, vals in st.rows[:3]:
            sample.append({fn: fmt(convert(vals[st.headers.index(c)], d.fields[fn])) for fn, c in m.columns.items()})
        lines = [f"已保存「{name}」（{kind}），{len(st.rows)} 行" + (f"，另有 {len(st.totals)} 行合计（不计入）" if st.totals else "")]
        lines += [f"注意：{w}" for w in warnings]
        lines += [f"关联提示：{h}" for h in self._relation_hints(name)]
        lines.append("前 3 行（金额已换算为元）：" + json.dumps(sample, ensure_ascii=False))
        return "\n".join(lines)

    def _values(self, name: str, fld: str) -> list[str]:
        """数据表某个字段的全部取值（文字），用来判断表之间的关联。"""
        from bonus_core.data.mapping import cell_text
        m = self.mapping[name]
        st = table_at(Path(m.file) if Path(m.file).exists() else self._file(Path(m.file).name), m.sheet, m.header_rows) \
            if m.header_rows else self.inv.sheet(f"{Path(m.file).name}/{m.sheet}")
        if st is None or fld not in m.columns:
            return []
        st, _ = filter_sheet(st, m.filter or None)
        i = st.headers.index(m.columns[fld])
        return [cell_text(v[i]) for _, v in st.rows if v[i] not in (None, "")]

    def _relation_hints(self, name: str) -> list[str]:
        """程序按数据判断的关联问题：编号列取值是别的表的主键却没同名、link 找不到被关联的表、该写 key 的没写。"""
        d, out = self.datasets[name], []
        keys = {o.key: n for n, o in self.datasets.items() if o.key and n != name}
        try:
            for f, fd in d.fields.items():
                if fd.type not in ("id", "text") or f == d.key:
                    continue
                vals = set(self._values(name, f))
                if not vals:
                    continue
                for k, n in keys.items():
                    if k != f and k not in d.fields and vals <= set(self._values(n, k)):
                        out.append(f"字段 {f} 的取值都是「{n}」的主键 {k}（如 {sorted(vals)[0]}）。要和「{n}」关联，"
                                   f"请把这个字段命名为 {k}（column 不变）" + ("，并写 link" if not d.link and not d.key else ""))
            if d.link and d.link not in keys and not any(d.link == o.link for n, o in self.datasets.items() if n != name):
                own = [n for n, o in self.datasets.items() if n != name and d.link in o.fields]
                out.append(f"link={d.link} 目前没有以 {d.link} 为主键的表" + (
                    f"；「{own[0]}」有 {d.link} 列，如果它每行一个 {d.link}，把它重新定义为 key={d.link} 的实体表"
                    if own else "；它应该是某个层级（如部门、小组、个人）或关联实体的主键"))
            if not d.key and not d.link:
                ids = [f for f in d.fields if f in keys]
                if ids:
                    out.append(f"这张表有 {ids[0]} 字段却没写 link，会被当作公司级表；按 {ids[0]} 挂到「{keys[ids[0]]}」上请写 link={ids[0]}")
            for n, o in self.datasets.items():
                if n != name and o.link and o.link in d.fields and d.key != o.link and o.link not in keys:
                    vals = self._values(name, o.link)
                    if vals and len(set(vals)) == len(vals):
                        out.append(f"「{n}」按 {o.link} 关联，而本表每行一个不同的 {o.link}：本表应写 key={o.link}（成为关联实体）")
        except Exception:      # 提示只是辅助，取值失败不影响保存
            pass
        return out

    def remove_dataset(self, name: str) -> str:
        if name not in self.datasets:
            raise ToolError(f"没有数据表「{name}」")
        del self.datasets[name], self.mapping[name]
        self.untested |= set(self.schemes)
        return f"已删除「{name}」"

    # ---- 规则（可以有多个方案，"共用" 方案合并进每个方案） ----
    def _scheme(self, scheme: str | None, *, for_write: bool = False) -> str:
        if scheme:
            if not for_write and scheme not in self.schemes:
                raise ToolError(f"没有方案「{scheme}」，已有方案：{self.schemes or '无'}")
            return scheme
        if for_write or not self.schemes:
            return self.schemes[0] if len(self.schemes) == 1 else DEFAULT_SCHEME
        if len(self.schemes) == 1:
            return self.schemes[0]
        raise ToolError(f"有多个方案 {self.schemes}，请用 scheme 指定")

    @property
    def schemes(self) -> list[str]:
        """要计算的方案（不含共用部分）。"""
        return [s for s in self.drafts if s != SHARED]

    def set_rules(self, scheme: str | None = None, **parts) -> str:
        name = self._scheme(scheme, for_write=True)
        if name == SHARED and parts.get("pool") is not None:
            raise ToolError("pool 是每个方案自己的，不能写在共用部分")
        draft = json.loads(json.dumps(self.drafts.get(name, {})))
        for k in ("tables", "variables"):
            for n, v in (parts.get(k) or {}).items():
                if v is None:
                    draft.setdefault(k, {}).pop(n, None)
                else:
                    draft.setdefault(k, {})[n] = v
        for k in ("pool", "allocation", "payout", "offsets", "checks", "revisions"):
            if parts.get(k) is not None:
                draft[k] = parts[k]
        for lv in draft.get("allocation") or []:       # 最上层不需要权重，模型常写成 null
            if isinstance(lv, dict) and lv.get("weight") is None:
                lv["weight"] = ""
        try:
            _revisions(draft.get("revisions"))
            targets = (self.schemes or [SHARED]) if name == SHARED else [name]   # 还没有方案时也先校验共用部分的格式
            rss = {t: self._ruleset(t, {**self.drafts, name: draft}) for t in targets}
        except (ValidationError, ValueError) as e:
            hint = ""
            for k in ("variables", "tables"):
                v = draft.get(k)
                if isinstance(v, dict) and ("expr" in v or "kind" in v):
                    hint += f"\n{k} 要写成 {{名字: {{...}}}}，每一项以名字为键，如 {{\"emp_sales\": {{\"expr\": \"...\", \"scope\": \"个人\"}}}}"
            raise ToolError(f"规则格式不对，未保存：{str(e)[:2500]}{hint}") from None
        self.drafts[name] = draft
        self.untested |= set(targets) - {SHARED}
        out = [f"已保存{'共用部分' if name == SHARED else f'方案「{name}」'}。"]
        self.last_blocking = set()
        for t, rs in rss.items():
            full = self._merged(t)
            issues = self._validate(t, rs)
            if "pool" not in full:
                issues = [i for i in issues if i.where != "pool"]
            blocking = [str(i) for i in issues if i.severity == "blocking"]
            warns = [str(i) for i in issues if i.severity == "warning"]
            head = f"「{t}」" + ("（尚未设定 pool）" if "pool" not in full else "")
            self.last_blocking |= {(t, i.code, i.where) for i in issues if i.severity == "blocking"}
            if blocking:
                out.append(head + "静态校验发现问题，需要修正：\n" + "\n".join(f"- {e}" for e in blocking))
            else:
                out.append(head + "静态校验通过。" + ("\n提示：\n" + "\n".join(f"- {w}" for w in warns) if warns else ""))
        if not self.schemes:
            out.append("还没有具体方案；用 set_rules(scheme=方案名, pool=...) 写入各方案后再完整校验。")
        return "\n".join(out)

    def _merged(self, scheme: str, drafts: dict | None = None) -> dict:
        drafts = self.drafts if drafts is None else drafts
        base, own = drafts.get(SHARED, {}), drafts.get(scheme, {})
        out = {k: v for k, v in base.items() if k not in ("tables", "variables")}
        for k in ("tables", "variables"):
            if base.get(k) or own.get(k):
                out[k] = {**base.get(k, {}), **own.get(k, {})}
        out.update({k: v for k, v in own.items() if k not in ("tables", "variables")})
        return out

    def _ruleset(self, scheme: str | None = None, drafts: dict | None = None) -> RuleSet:
        scheme = scheme or self._scheme(None)
        draft = {k: v for k, v in self._merged(scheme, drafts).items() if k != "revisions"}
        ext = Extraction.model_validate({"pool": {"expr": "0"}, "allocation": [], **draft})
        return RuleSet(name=scheme if scheme != DEFAULT_SCHEME else (self.summary[:30] or "智能体草稿"),
                       version="agent-draft", datasets=self.datasets, **ext.model_dump(exclude={"questions"}))

    def _data(self, rs: RuleSet):
        data, issues = apply_mapping(rs, self.data_files, self.mapping)
        blocking = [i for i in issues if i.severity == "blocking"]
        if blocking:
            raise ToolError("数据问题：\n" + "\n".join(f"- {i}" for i in blocking[:8]))
        return data, issues

    def trial_run(self, scheme: str | None = None) -> str:
        if not self.datasets:
            raise ToolError("还没有定义数据表")
        names = [self._scheme(scheme)] if scheme or len(self.schemes) <= 1 else self._order(self.schemes)
        out = []
        for name in names:
            if "pool" not in self._merged(name) or not self._merged(name).get("allocation"):
                raise ToolError(f"方案「{name}」的规则还缺 pool 或 allocation")
            out.append(self._trial(name))
        return "\n\n".join(out)

    def _deps(self, name: str) -> set[tuple[str, str]]:
        """方案引用的其他方案（修订）：规则里 scheme_paid("方案", "修订") 的参数。"""
        out: set[tuple[str, str]] = set()

        def walk(v):
            if isinstance(v, str):
                out.update((m[0], m[1] or "R0") for m in SCHEME_REF.findall(v))
            elif isinstance(v, dict):
                for x in v.values():
                    walk(x)
            elif isinstance(v, list):
                for x in v:
                    walk(x)
        walk(self._merged(name))
        return out

    def _order(self, names: list[str]) -> list[str]:
        """被引用的方案排在前面。"""
        out: list[str] = []

        def visit(n, path):
            if n in out:
                return
            if n in path:
                raise ToolError(f"方案之间循环引用：{' → '.join(path + [n])}")
            for d, _ in sorted(self._deps(n)):
                if d in self.schemes:
                    visit(d, path + [n])
            out.append(n)
        for n in names:
            visit(n, [])
        return [n for n in out if n in names]

    def _refs(self, name: str, _path: tuple = ()) -> dict[str, Result]:
        """本方案引用的其他方案结果；还没试算或改过的，先试算。"""
        refs = {}
        for s, rev in sorted(self._deps(name)):
            if s == name or s in _path:
                raise ToolError(f"方案「{name}」不能引用自己或形成循环引用（{s}）")
            if s not in self.schemes:
                raise ToolError(f"方案「{name}」引用了不存在的方案「{s}」，现有方案：{self.schemes}")
            if s in self.untested or s not in self.results:
                self._trial(s, _path + (name,))
            if rev not in self.results[s]:
                raise ToolError(f"方案「{name}」引用了「{s}」的修订 {rev}，但该方案只有 {list(self.results[s])}")
            refs[f"{s}/{rev}"] = self.results[s][rev]
        return refs

    def _validate(self, scheme: str, rs: RuleSet) -> list[Issue]:
        """静态校验；全局变量在某层按对象引用、又没明确写 scope 的，必须先处理（写明 scope），不能只当提示。"""
        variables = self._merged(scheme).get("variables") or {}
        out = []
        for i in validate(rs):
            m = re.match(r"变量 (\S+) 没写 scope", i.message) if i.code == "global_in_level" else None
            if m and "scope" not in (variables.get(m.group(1)) or {}):
                i = Issue.block(i.code, i.message + "。必须明确写 scope（要全公司合计就写 scope: global）", i.where)
            out.append(i)
        return out

    def _trial(self, name: str, _path: tuple = ()) -> str:
        rs = self._ruleset(name)
        guard = [str(i) for i in self._validate(name, rs) if i.code == "global_in_level" and i.severity == "blocking"]
        if guard:
            raise ToolError(f"方案「{name}」试算前需要修正：\n" + "\n".join(f"- {g}" for g in guard))
        data, issues = self._data(rs)
        refs = self._refs(name, _path)
        try:
            eng = Engine(rs, data, refs)
            res = {"R0": eng.compute()}
            done: list[Adjustment] = []
            for rev in _revisions(self._merged(name).get("revisions")):
                done += [Adjustment(a["node"], Decimal(str(a["amount"])), reason=a.get("reason", ""), id=a.get("id", ""),
                                    created_at=f"{len(done) + i:06d}") for i, a in enumerate(rev["adjust"])]
                res[rev["name"]] = eng.compute(list(done))
        except EngineError as e:
            raise ToolError(f"方案「{name}」试算失败：\n" + "\n".join(f"- {i}" for i in e.issues[:8])) from None
        except (EvalError, AdjustConflict, ValueError) as e:
            raise ToolError(f"方案「{name}」试算失败：{e}") from None
        self.results[name] = res
        self.untested.discard(name)
        self.untested |= {s for s in self.schemes if s != name and name in {d for d, _ in self._deps(s)}}  # 引用本方案的要重算
        base = res["R0"]
        lines = [f"方案「{name}」试算成功。奖金池 {fmt(Decimal(base.pool))} 元（精确值 {base.pool_exact}）"]
        for lv in base.levels:
            ns = sorted((n for n in base.nodes.values() if n.level == lv), key=lambda n: -float(n.final))
            total = sum(float(n.final) for n in ns)
            zero = sum(1 for n in ns if float(n.final) == 0)
            top = "，".join(f"{n.id}（{n.label}）{fmt(Decimal(n.final))}" for n in ns[:8])
            lines.append(f"「{lv}」{len(ns)} 个，合计 {total:,.2f}，为 0 的 {zero} 个；最多的：{top}"
                         + ("……" if len(ns) > 8 else "") + f"；最少的：{ns[-1].id} {fmt(Decimal(ns[-1].final))}")
        for lv in base.levels:                          # 抽查：每层取一个中间对象，列出权重代入后的数值
            ns = sorted((n for n in base.nodes.values() if n.level == lv), key=lambda n: float(n.final))
            nid = ns[len(ns) // 2].id if ns else ""
            tr = base.traces.get(f"{nid}/权重")
            while tr is not None and tr.expr.isidentifier() and f"{nid}/{tr.expr}" in base.traces:   # 权重直接写成变量名时展开
                tr = base.traces[f"{nid}/{tr.expr}"]
            if tr is not None and tr.rendered != tr.expr:
                lines.append(f"抽查「{lv}」{tr.id}：{tr.rendered[:400]} = {tr.value}")
        for rev, r in list(res.items())[1:]:
            changed = [k for k, n in r.nodes.items() if n.final != base.nodes[k].final]
            lines.append(f"修订 {rev}：{len(changed)} 个对象金额变化，奖金池 {fmt(Decimal(r.pool))}；"
                         + "，".join(f"{k} {fmt(Decimal(base.nodes[k].final))}→{fmt(Decimal(r.nodes[k].final))}"
                                    for k in changed[:6]))
        held = sum(float(p.held) for p in base.payouts.values())
        if base.payouts:
            rows = base.payout_rows()
            paid = sum(Decimal(p.paid) for p in rows)
            off = sum(Decimal(p.offset or 0) for p in rows)
            over = sum(Decimal(p.overpaid or 0) for p in rows)
            lines.append(f"发放：本次实发合计 {paid:,.2f}，暂缓合计 {held:,.2f}"
                         + (f"，已发抵扣 {off:,.2f}，超发待核对 {over:,.2f}" if off or over else ""))
            if base.rollups:
                rs_ = sorted(base.rollups.values(), key=lambda r: r.key)
                lines.append(f"按「{rs_[0].entity}」汇总后抵扣：{len(rs_)} 个；"
                             + "，".join(f"{r.key}（{r.label}）可发上限 {fmt(Decimal(r.cap))} 抵扣 {fmt(Decimal(r.deducted))} "
                                        f"实发 {fmt(Decimal(r.paid))}" + (f" 超发 {fmt(Decimal(r.overpaid))}" if Decimal(r.overpaid) else "")
                                        for r in rs_[:8]) + ("……" if len(rs_) > 8 else ""))
            if held == 0 and any(p.get("pay_ratio") or p.get("hold_ratio") for p in self._merged(name).get("payout") or []):
                lines.append("请核对：写了按比例发放，但所有对象都全额可发、暂缓为 0。比例的分子分母是否写反"
                             "（如应为 回款 ÷ 销售），筛选条件是否对")
        if refs:
            lines.append("引用：" + "、".join(refs))
        mismatch = self._period_mismatch(name)
        if mismatch:
            lines.append("请核对：" + mismatch)
        warns = [str(i) for i in list(issues) + list(base.issues) if i.severity == "warning"]
        if warns:
            lines.append("数据提示：\n" + "\n".join(f"- {w}" for w in warns[:8]))
        return "\n".join(lines)

    def _period_mismatch(self, name: str) -> str:
        """奖金池和分配权重用的统计期间（日期、period == "Q3" 这类取值）对不上时给出提示，常见于把按季度算的权重放在共用部分、
        年度方案也跟着用了。只提示，不阻断。"""
        full = self._merged(name)
        variables = full.get("variables") or {}

        def periods(expr: str, depth: int = 0) -> set[str]:
            out = set(re.findall(r'"(\d{4}-\d{2}-\d{2})"', expr)) | set(re.findall(r'==\s*"([^"]{1,4})"', expr))
            if depth < 4:
                for n in set(re.findall(r"[A-Za-z_][A-Za-z0-9_]*", expr)) & set(variables):
                    out |= periods(str((variables[n] or {}).get("expr", "")), depth + 1)
            return out
        own = [full.get("pool") or {}] + list(full.get("payout") or []) + list(full.get("offsets") or [])
        pool = periods(" ".join(str(v) for p in own if isinstance(p, dict) for k, v in p.items()
                                if k in ("expr", "pay_ratio", "hold_ratio")))
        weight = periods(" ".join(str(lv.get("weight") or "") for lv in full.get("allocation") or []))
        if pool and weight and not pool & weight:
            src = "（allocation 来自共用部分）" if "allocation" not in self.drafts.get(name, {}) else ""
            return (f"方案「{name}」的奖金池和发放用的期间是 {sorted(pool)}，分配权重用的是 {sorted(weight)}{src}。"
                    "如果权重应按本方案的期间计算，把 allocation 写到这个方案里")
        return ""

    @property
    def result(self) -> Result | None:
        """第一个方案的基准结果（兼容只有一个方案的情况）。"""
        s = self.schemes[0] if self.schemes else None
        return self.results.get(s, {}).get("R0") if s else None

    @property
    def dirty(self) -> bool:
        return bool(self.untested) or not self.schemes or any(s not in self.results for s in self.schemes)

    def explain(self, id: str, scheme: str | None = None, revision: str = "R0") -> str:
        name = self._scheme(scheme)
        res = self.results.get(name, {}).get(revision)
        if res is None:
            raise ToolError(f"方案「{name}」{revision} 还没有试算结果，先 trial_run")
        tr = res.traces
        keys = [k for k in tr if k == id or k.startswith(id + "/")] or [k for k in tr if id in k]
        if not keys:
            raise ToolError(f"没有找到「{id}」的计算链，节点 id 形如 {next(iter(res.nodes), '')}")
        out = []
        for k in keys[:12]:
            t = tr[k]
            out.append(f"{t.id}：{t.expr}\n  = {t.rendered}\n  = {t.value}" + (f"\n  明细行：{', '.join(t.rows[:5])}" if t.rows else ""))
        if id in res.nodes:
            n = res.nodes[id]
            out.append(f"{id} 最终金额 {fmt(Decimal(n.final))}（权重 {n.weight}，上级 {n.parent}）")
        return "\n".join(out)

    def evaluate(self, expr: str, node: str = ROOT_LEVEL, scheme: str | None = None) -> str:
        name = self._scheme(scheme) if self.schemes else self._scheme(None, for_write=True)
        if not self._merged(name).get("allocation"):
            raise ToolError("需要先用 set_rules 设定 allocation（pool 可以先写 0）")
        rs = self._ruleset(name)
        data, _ = self._data(rs)
        root, by_level, issues = build_tree(rs, data)
        blocking = [i for i in issues if i.severity == "blocking"]
        if blocking:
            raise ToolError("建树失败：" + "；".join(str(i) for i in blocking[:5]))
        nodes = {n.id: n for n in root.walk()}
        nodes.update({n.id: n for tbl in (root.index or {}).values() for n in tbl.values()})
        if node not in nodes:
            raise ToolError(f"没有节点「{node}」，例：{list(nodes)[:5]}")
        ev = Evaluator(rs, root, by_level)
        try:
            v = ev.evaluate(expr, nodes[node], "probe", "probe")
        except (EvalError, ValueError) as e:
            raise ToolError(f"计算失败：{e}") from None
        return f"{expr} = {ev.traces['probe'].rendered} = {fmt(v)}"

    # ---- 问题与完成 ----
    def ask_question(self, text: str, options: list[str] | None = None, suggested: int | None = None, related: str = "",
                     blocking: bool = True) -> str:
        same = next((q for q in self.questions if q.text.strip() == text.strip()), None)
        if same is not None:
            if same.answer:
                return f"这个问题（{same.id}）财务已答复：{same.answer}。请按答复处理"
            return f"这个问题已经问过（{same.id}），还在等财务答复"
        q = Question(id=f"q{len(self.questions) + 1}", text=text, options=options or [], suggested=suggested,
                     related=related, blocking=blocking, round=self.round)
        self.questions.append(q)
        if blocking:
            return (f"已记录问题 {q.id}，等财务答复。不要自行选择口径或补出算法；"
                    "继续做不受影响的部分，全部问完后调用 finish 结束本轮")
        return f"已记录提示 {q.id}（不影响金额，不必等答复）"

    @property
    def pending(self) -> list[Question]:
        """还在等财务答复的问题。"""
        return [q for q in self.questions if q.blocking and not q.answer and not q.withdrawn]

    def finish(self, summary: str, ignored: list[dict] | None = None) -> str:
        # 上一轮登记过的忽略项继续有效，本轮同名的覆盖
        ignored = list({str(i.get("ref", "")): i for i in [*self.ignored, *(ignored or [])] if isinstance(i, dict)}.values())
        self.ignored = ignored                 # finish 被拦下时也保留，修正后再 finish 不必重写
        skipped = {str(i.get("ref", "")) for i in ignored}
        docs = [f.name for f in self.inv.files if f.suffix.lower() in RULE_SUFFIXES]
        bare = [d for d in docs if d not in skipped and not any(f.file == d for f in self.facts)]
        if bare:
            raise ToolError(f"这些规则文档还没有登记参数：{'、'.join(bare)}。请用 record_facts 登记其中的参数和口径"
                            "（程序靠台账比对冲突）；确实与计算无关的写进 finish 的 ignored（ref 写文件名，说明原因）")
        if self.pending:     # 有问题等答复：结束本轮，不出正式结果
            self.summary, self.ignored, self.done = summary, ignored, True
            return f"本轮结束，等待财务答复 {len(self.pending)} 个问题：{'、'.join(q.id for q in self.pending)}"
        if not self.schemes:
            raise ToolError("还没有写入规则")
        pending = [s for s in self.schemes if s in self.untested or s not in self.results]
        if pending:
            raise ToolError(f"方案 {pending} 最近一次修改之后还没有成功试算，请先 trial_run")
        miss = [n for n in fx.unused(self.facts, self.drafts, self.conflicts, self.questions) if n.split("（")[0] not in skipped]
        if miss:
            raise ToolError(f"台账里这些参数的数值在规则里没有出现：{'、'.join(miss)}。请核对规则是否漏了或写错；"
                            "确实不参与计算的写进 ignored（ref 写参数名，说明原因）")
        self.summary, self.ignored, self.done = summary, ignored, True
        return "已完成"

    def reply(self, text: str, ignored: list[dict] | None = None) -> str:
        if not self.chat_by:
            raise ToolError("现在不是对话，请用 finish 结束")
        if self.changed and self.schemes:
            if any(s in self.untested or s not in self.results for s in self.schemes):
                out = self.trial_run()
                if any(s in self.untested or s not in self.results for s in self.schemes):
                    raise ToolError("修改后试算没有通过，请修正后再回复：\n" + out[:2000])
            self.finish(self.summary or "对话中修改", ignored)
        self.reply_text, self.done = text.strip(), True
        return "已回复"

    def remove_facts(self, name: str) -> str:
        """删掉某个参数名下的全部登记（登记错了时用）；依据它的、还没答复的冲突问题由程序撤回。"""
        before = len(self.facts)
        self.facts = [f for f in self.facts if f.name.strip() != name.strip()]
        if len(self.facts) == before:
            raise ToolError(f"台账里没有「{name}」，已有：{'、'.join(dict.fromkeys(f.name for f in self.facts))[:800]}")
        gone = fx.withdraw_stale(self.facts, self.conflicts, self.questions)
        return (f"已删除「{name}」的 {before - len(self.facts)} 条登记" +
                (f"，撤回问题 {'、'.join(q.id for q in gone)}" if gone else "") + "。请按正确的名字重新登记")

    def answer_question(self, id: str, answer: str) -> str:
        if not self.chat_by:
            raise ToolError("现在不是和财务对话，待确认问题只能由财务答复，你不能替财务回答。如果是程序因为你登记错了而误报的冲突（如把阶梯几档登记成"
                            "同一个名字），用 remove_facts 删掉那个名字的登记、按正确的名字重新登记，程序会自动撤回问题")
        q = next((q for q in self.questions if q.id == id), None)
        if q is None or q.withdrawn:
            raise ToolError(f"没有待答复的问题 {id}")
        from datetime import datetime
        q.answer, q.answered_by = answer.strip(), f"{self.chat_by}（对话）"
        q.answered_at = datetime.now().isoformat(timespec="seconds")
        return f"已记录 {id} 的答复：{q.answer}。请按答复修改规则"

    def next_hint(self) -> str:
        """程序根据工作区判断的下一步（模型原地打转或步数用得多时提示）。"""
        docs = [f.name for f in self.inv.files if f.suffix.lower() in RULE_SUFFIXES]
        bare = [d for d in docs if not any(f.file == d for f in self.facts)
                and d not in {str(i.get("ref", "")) for i in self.ignored}]
        if not self.datasets:
            step = f"用 define_dataset 定义数据表（数据文件：{'、'.join(f.name for f in self.data_files) or '无'}）"
        elif not self.schemes:
            step = "用 set_rules 写入规则（笔记和参数台账里已有的内容不必重读）"
        elif any(s in self.untested or s not in self.results for s in self.schemes):
            step = "用 trial_run 试算"
        elif bare:
            step = f"用 record_facts 登记这些文档的参数：{'、'.join(bare)}"
        else:
            step = "核对后调用 finish"
        return step + (f"；还没登记参数的规则文档：{'、'.join(bare)}" if bare and "record_facts" not in step else "")

    def state(self) -> str:
        """每轮附在请求末尾的工作区摘要。"""
        lines = ["【当前工作区状态】" + (f"（{self.budget}）" if self.budget else "")]
        used, limit = self.progress
        if limit and used >= limit * 0.3:
            lines.append(f"- 已用 {used}/{limit} 步，下一步（程序判断）：{self.next_hint()}")
        if self.notes:
            lines.append("- 笔记：")
            lines += [f"  · {k}：{v}" for k, v in self.notes.items()]
        if self.datasets:
            for n, d in self.datasets.items():
                m = self.mapping[n]
                kind = (f"key={d.key}" + (f", link={d.link}" if d.link else "")) if d.key else f"link={d.link}" if d.link else "公司级"
                flt = f" 筛选 {m.filter}" if m.filter else ""
                lines.append(f"- 数据表「{n}」← {m.file}/{m.sheet}{flt}（{kind}）字段：{', '.join(d.fields)}")
        else:
            lines.append("- 尚未定义数据表")
        if self.drafts:
            for s, text in _rules_briefs(self.drafts).items():
                title = "共用部分（合并进每个方案）" if s == SHARED else f"方案「{s}」"
                lines.append(f"- {title}规则摘要（不含出处；修改时按名字只传要改的部分）：\n" + text)
        else:
            lines.append("- 尚未写入规则")
        if self.facts:
            by = {}
            for f in self.facts:
                by.setdefault(f.file, []).append(f.name)
            lines.append(f"- 参数台账 {len(self.facts)} 条（再登记同一参数请用相同的名字）：")
            lines += [f"  · {k}：{'、'.join(dict.fromkeys(v))}" for k, v in by.items()]
        if self.questions:
            lines.append("- 问题：")
            lines += [f"  · {q.id} {q.text} → " + (f"财务答复：{q.answer}" if q.answer else f"已撤回：{q.withdrawn}" if q.withdrawn
                                          else "等待答复" if q.blocking else "提示")
                      for q in self.questions]
        if self.transcripts:
            lines.append(f"- 已查看图片：{', '.join(self.transcripts)}")
        for s in self.schemes:
            r = self.results.get(s, {})
            lines.append(f"- 试算「{s}」：" + ("最近修改后尚未试算" if s in self.untested or not r else
                                             f"已试算，奖金池 {fmt(Decimal(r['R0'].pool))}" + (f"，修订 {list(r)[1:]}" if len(r) > 1 else "")))
        return "\n".join(lines)


def _strip_tags(text: str, name: str) -> str:
    """材料包里每段前的 [文件名 · text] 标记去掉文件名（标题已写明）；只剩 text 的整行删掉，页码和表格、图片标记保留。"""
    def sub(m):
        rest = m.group(1).removeprefix(name).strip()
        kind = m.group(2)
        label = " · ".join(x for x in (rest, "" if kind == "text" else kind) if x)
        return f"[{label}]\n" if label else ""
    return re.sub(r"^\[([^\]\n]*?) · (\w+)\]\n", sub, text, flags=re.M)


def _json_data_summary(path: Path) -> str:
    """大 JSON 数据文件的结构摘要：各数组的条数、字段和第一条记录。不是 JSON 或不是数据时返回空。"""
    if path.suffix.lower() != ".json":
        return ""
    try:
        obj = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return ""
    arrays = ({k: v for k, v in obj.items() if isinstance(v, list)} if isinstance(obj, dict)
              else {"（顶层）": obj} if isinstance(obj, list) else {})
    if not arrays:
        return ""
    lines = ["结构摘要（程序生成）："]
    for k, v in arrays.items():
        first = v[0] if v else None
        cols = list(first) if isinstance(first, dict) else []
        lines.append(f"- {k}：{len(v)} 条" + (f"，字段 {cols}" if cols else "")
                     + (f"，第一条 {json.dumps(first, ensure_ascii=False)[:300]}" if first is not None else ""))
    return "\n".join(lines)


def _revisions(revs) -> list[dict]:
    """修订版本：[{"name": "R1", "adjust": [{"node": "小组:G01", "amount": 900000, "reason": ...}]}, ...]。
    后一个修订包含前面所有修订的调整（逐步累加），R0 是不做调整的基准。"""
    out = []
    for r in revs or []:
        if not isinstance(r, dict) or not r.get("name") or not isinstance(r.get("adjust"), list):
            raise ValueError('revisions 的每一项应为 {"name": "R1", "adjust": [{"node": "小组:G01", "amount": 900000}]}')
        if r["name"] == "R0":
            raise ValueError("R0 是不做调整的基准结果，修订从 R1 开始命名")
        for a in r["adjust"]:
            if not isinstance(a, dict) or "node" not in a or "amount" not in a:
                raise ValueError(f"修订 {r['name']} 的调整项应为 {{\"node\": 节点 id, \"amount\": 金额}}，得到 {a}")
        out.append(r)
    return out


def _rules_lines(draft: dict) -> list[tuple[str, str]]:
    """规则摘要的每一行：(名字, 行)。"""
    out = []
    for n, t in (draft.get("tables") or {}).items():
        if t.get("kind") == "lookup" or "mapping" in t:
            out.append((f"表 {n}", f"  表 {n}：查找表 {json.dumps(t.get('mapping'), ensure_ascii=False)}"))
        else:
            out.append((f"表 {n}", f"  表 {n}：{t.get('mode', '')} 阶梯 {json.dumps(t.get('steps'), ensure_ascii=False)}"))
    for n, v in (draft.get("variables") or {}).items():
        out.append((n, f"  变量 {n}[{v.get('scope', 'global')}] = {v.get('expr')}"))
    if "pool" in draft:
        p = draft["pool"] or {}
        out.append(("pool", f"  pool[{p.get('scope', 'global')}] = {p.get('expr')}"))
    for lv in draft.get("allocation") or []:
        out.append((f"分配 {lv.get('level')}", f"  分配 {lv.get('level')}（key={lv.get('key')}, parent_key={lv.get('parent_key')}）"
                                              f"权重 = {lv.get('weight')}"))
    for p in draft.get("payout") or []:
        out.append((f"发放 {p.get('id')}", f"  发放 {p.get('id')}[{p.get('scope')}] 暂缓比例 = {p.get('hold_ratio')}"))
    for c in draft.get("checks") or []:
        out.append((f"校验 {c.get('id')}", f"  校验 {c.get('id')}[{c.get('scope', 'global')}]：{c.get('expr')}"))
    for r in draft.get("revisions") or []:
        out.append((f"修订 {r.get('name')}", f"  修订 {r.get('name')}："
                    + "，".join(f"{a.get('node')}={a.get('amount')}" for a in r.get("adjust") or [])))
    return out


def _rules_briefs(drafts: dict, limit: int = 6000) -> dict[str, str]:
    """各方案的规则摘要；和前面某个方案完全相同的行只列名字，不重复全文（每轮都附在状态里，省 tokens）。"""
    seen: dict[str, str] = {}
    out = {}
    for s, d in drafts.items():
        keep, same = [], {}
        for key, line in _rules_lines(d):
            if s != SHARED and seen.get(line, s) != s:
                same.setdefault(seen[line], []).append(key)
            else:
                seen.setdefault(line, s)
                keep.append(line)
        keep += [f"  与「{x}」相同（不重复列出）：{'、'.join(names)}" for x, names in same.items()]
        text = "\n".join(keep)
        out[s] = text if len(text) <= limit else text[:limit] + "……"
    return out


# ---------------- 主循环 ----------------
def system_prompt() -> str:
    dsl = prompt("extract_rules")
    dsl = dsl[dsl.index("# 规则 DSL") + len("# 规则 DSL"):dsl.index("# 数据字典")]
    dsl = dsl.replace("数据字典（datasets）已固定，见下文，你**只能引用其中的字段名**，不能自造字段。你需要输出：",
                      "表达式只能引用你用 define_dataset 定义过的字段名。规则包括：")
    dsl = dsl.replace("{functions}", describe_functions())
    return prompt("agent").replace("{dsl}", dsl.strip())


JSON_PROTOCOL = """
# 回复格式（本模型不使用原生工具调用）
每次只输出一个 JSON 对象，不要输出其他文字：{"thought": "简短的思考", "tool": "工具名", "args": {参数}}
可用工具（JSON Schema）：
"""


@dataclass
class AgentResult:
    ws: Workspace
    steps: list[Step]
    ruleset: RuleSet | None
    result: Result | None
    error: str = ""
    outputs: dict[str, Path] = field(default_factory=dict)
    quote_hits: int = 0
    quote_total: int = 0
    recheck: str = ""
    rulesets: dict[str, RuleSet] = field(default_factory=dict)          # 方案 -> 规则
    pending: list[dict] = field(default_factory=list)                   # 尚未答复的问题
    draft_results: dict[str, dict[str, Result]] = field(default_factory=dict)   # 有问题未答复时的试算结果（不输出）
    round: int = 1
    changes: dict | None = None                                         # 本轮相对上一轮的文件变化
    diff_rows: list[list] = field(default_factory=list)                 # 与上一轮逐个对象的金额对比
    rule_changes: list[str] = field(default_factory=list)
    compared: bool = False                                              # 上一轮有结果，做了金额对比
    results: dict[str, dict[str, Result]] = field(default_factory=dict)  # 方案 -> 修订 -> 结果
    reply: str = ""                                                      # 对话时给财务的回复

    @property
    def single(self) -> bool:
        """只有一个默认方案、没有修订：输出文件和以前一样（rules.yaml、result.xlsx）。"""
        return list(self.results) == [DEFAULT_SCHEME] and list(self.results[DEFAULT_SCHEME]) == ["R0"]

    @property
    def ok(self) -> bool:
        return self.ws.done and self.result is not None

    @property
    def questions(self) -> list[Question]:
        return self.ws.questions

    @property
    def status(self) -> str:
        """OK：已算出，可交财务复核；NEEDS_CONFIRMATION：有问题等财务答复，本轮不出结果；FAILED：未完成。"""
        if self.ws.done and self.pending:
            return "NEEDS_CONFIRMATION"
        return "OK" if self.ok else "FAILED"


def _size(ms: list[dict]) -> int:
    return sum(len(json.dumps(m, ensure_ascii=False)) for m in ms)


def _is_result(m: dict) -> bool:
    return m["role"] == "tool" or (m["role"] == "user" and str(m.get("content", "")).startswith("【") and "的结果】" in
                                   str(m.get("content", "")).split("\n", 1)[0])


def _shrink(history: list[dict], max_chars: int, keep: int = KEEP_FULL_RESULTS) -> bool:
    """历史超过 max_chars 时整理一次（原地修改），返回是否整理过。

    先把较早的工具结果改成摘要、较早的长参数省略；仍超过上限的一半，就从最早的一轮开始丢弃，
    直到不超过一半，留出空间让之后多轮只追加不改动（前缀不变，缓存可以命中）。
    丢弃时保证剩下的历史以 assistant 消息开头（工具结果不落单）。"""
    if _size(history) <= max_chars:
        return False
    results = [i for i, m in enumerate(history) if _is_result(m)]
    cut = results[-keep] if len(results) >= keep else 0
    for i in results[:-keep] if len(results) > keep else []:
        c = history[i]["content"]
        if len(c) > BRIEF_CHARS + 60:
            history[i] = {**history[i], "content": c[:BRIEF_CHARS] + "\n……（较早的结果已省略。关键信息应已记在笔记里，确需原文可重新调用）"}
    for i, m in enumerate(history[:cut]):
        for tc in m.get("tool_calls") or []:
            if len(tc["function"]["arguments"]) > 600:
                tc["function"]["arguments"] = json.dumps({"省略": "较早的参数已省略，当前规则见工作区状态"}, ensure_ascii=False)
    starts = [i for i, m in enumerate(history) if m["role"] == "assistant"]
    for i in starts:
        if _size(history[i:]) <= max_chars // 2:
            del history[:i]
            return True
    if starts:
        del history[:starts[-1]]
    return True


def _answers_text(answers: list[dict] | None) -> str:
    done = [a for a in answers or [] if a.get("answer") not in (None, "")]
    if not done:
        return ""
    lines = ["\n财务已确认的答复（必须按此处理，不要再就这些问题提问）："]
    for a in done:
        ans = a["answer"]
        if isinstance(ans, int) and 0 <= ans < len(a.get("options") or []):
            ans = a["options"][ans]
        lines.append(f"- {a.get('text', '')} → {ans}")
    return "\n".join(lines)


def run_agent(llm: LLM, paths: list[str | Path], out_dir: str | Path | None = None, *,
              answers: list[dict] | None = None, max_steps: int = 80, max_history_chars: int = 40000,
              max_files: int = MAX_FILES, max_tokens: int | None = 2_000_000, exclude: list[str | Path] = (),
              on_step: Callable[[int, Step], None] | None = None, resume: bool = True,
              message: str | None = None, by: str = "", chat: list[dict] | None = None) -> AgentResult:
    """max_steps 和 max_tokens 都是安全上限：模型调用 finish 就结束，正常不会用满。

    out_dir 是案卷目录：里面已有上一轮的 case.json 时（resume=True），从上一轮的工作区接着做，
    财务答复默认读取案卷里的 questions.yaml。"""
    try:
        files = collect_files(paths, exclude=([out_dir] if out_dir else []) + list(exclude), max_files=max_files)
    except (TooManyFiles, FileNotFoundError) as e:
        raise ToolError(str(e)) from None
    if not files:
        raise ToolError("没有找到可处理的文件（支持 .docx .pptx .xlsx .xlsm .png .jpg .txt .md .json .csv）")
    ws = Workspace(llm, build_inventory(files))
    out = Path(out_dir) if out_dir is not None else None
    case = cs.load_case(out) if out is not None else None      # --new 也保留历次记录，只是不接着上一轮的工作区
    resumed = bool(case and resume and case.get("workspace"))
    prints = cs.fingerprint(files)
    rnd = cs.new_round(case, prints)
    ws.round = rnd["n"]
    if answers is None and resumed and (out / "questions.yaml").exists():
        answers = yaml.safe_load((out / "questions.yaml").read_text(encoding="utf-8")) or []
    last = case["rounds"][-1] if case and case.get("rounds") else None
    changes = None
    if resumed:
        cs.restore(ws, case["workspace"])
        changes = cs.file_changes(last["files"] if last else {}, prints)
        gone = set(changes["修改"]) | set(changes["删除"])
        ws.facts = [f for f in ws.facts if f.file not in gone]       # 这些文件的参数要按新内容重新登记
        withdrawn = fx.withdraw_stale(ws.facts, ws.conflicts, ws.questions)
    else:   # 首轮就带了答复（如上一次单轮运行的 questions.yaml）：作为已答复的问题放进工作区
        ws.questions = [Question(id=a.get("id") or f"q{i + 1}", text=a["text"], options=a.get("options") or [])
                        for i, a in enumerate(answers or []) if a.get("text") and a.get("answer") not in (None, "")]
    fresh = cs.apply_answers(ws.questions, answers)
    handlers = {t["name"]: getattr(ws, t["name"]) for t in TOOLS}
    ws.transcribe_all()
    pack = ("\n\n【材料包】以下是本批次材料的内容（程序读取，图片已转写），每轮都附在这里、不会被裁剪。"
            "标了“全文 / 整表”的不要再用工具读；数据表只附了前几行，define_dataset 会读取全部行。\n\n" + ws.briefing())
    if message:             # 对话：财务的一条消息。先按现有规则试算一遍，便于回答和查算式
        ws.chat_by = by or "财务"
        for s in ws._order(ws.schemes):
            try:
                ws._trial(s)
            except Exception:
                pass
        task = _chat_task(ws, message, chat, fresh, (changes if resumed else None))
    elif resumed:
        task = _resume_task(ws, last, changes, fresh, withdrawn)
    else:
        task = ("请完成本批次的奖金计算。文件清单：\n" + ws.list_files() + _answers_text(answers)
                + "\n\n完成后调用 finish。")
    task += pack
    mode = llm.tool_mode
    sys = system_prompt()
    history: list[dict] = []
    steps: list[Step] = []
    error = ""
    nudges = 0
    seen: dict[str, int] = {}          # 只读工具的调用 -> 第几步，用来发现原地打转
    repeats = 0
    last_write: tuple[str, str] | None = None
    fails: dict[str, int] = {}         # 工具 -> 连续失败次数
    stuck = ""
    undo: dict[tuple[str, str], int] = {}
    same_block, prev_block = 0, set()
    for turn in range(max_steps):
        used = llm.usage.prompt_tokens + llm.usage.completion_tokens
        if max_tokens and used >= max_tokens:
            error = f"tokens 用量 {used} 达到上限 {max_tokens}，已停止"
            break
        ws.progress = (len(steps), max_steps)
        ws.budget = f"已用 {len(steps)} 步，上限 {max_steps} 步" + (f"；tokens 已用 {used}，上限 {max_tokens}" if max_tokens else "")
        if _shrink(history, max_history_chars):
            seen.clear()       # 较早的结果已省略，重新读取是合理的
        tail = [{"role": "user", "content": ws.state()}]
        # 缓存标记：系统提示 + 任务固定不变；历史在两次整理之间只追加，标在历史最后一条上
        body = [{"role": "user", "content": task, "_cache": True}] + history[:-1] \
            + ([{**history[-1], "_cache": True}] if history else [])
        if mode == "json":
            msgs = [{"role": "system", "content": sys + JSON_PROTOCOL + json.dumps(TOOLS, ensure_ascii=False)}] + body + tail
            try:
                text = llm.chat(msgs, json_out=True, temperature=0.1)
            except LLMError as e:
                error = str(e)
                break
            history.append({"role": "assistant", "content": text})
            try:
                act = parse_json(text)
                calls = [{"id": "", "name": act["tool"], "arguments": json.dumps(act.get("args") or {}, ensure_ascii=False)}]
            except (LLMError, KeyError, TypeError, ValueError):
                history.append({"role": "user", "content": '格式不对：请只输出 {"thought": ..., "tool": ..., "args": {...}}'})
                continue
        else:
            msgs = [{"role": "system", "content": sys}] + body + tail
            try:
                reply = llm.chat_tools(msgs, openai_tools())
            except LLMError as e:
                if turn == 0:  # 模型或服务不支持原生工具调用，改用 JSON 协议
                    mode = "json"
                    continue
                error = str(e)
                break
            calls = [{**c, "arguments": _fix_arguments(c["arguments"])} for c in reply["tool_calls"]]
            history.append({"role": "assistant", "content": reply["content"],
                            **({"tool_calls": [{"id": c["id"], "type": "function",
                                                "function": {"name": c["name"], "arguments": c["arguments"]}}
                                               for c in calls]} if calls else {})})
            if not calls:
                nudges += 1
                if nudges > 3:
                    error = "模型多次没有调用工具"
                    break
                history.append({"role": "user", "content": "请调用工具继续工作；全部完成后调用 finish。"})
                continue
        turn_start, skipped = len(history) - 1, 0     # 本轮 assistant 消息的位置；被拦下的重复调用数
        for c in calls:
            args = {}
            sig = f"{c['name']}:{c['arguments']}"
            if c["name"] in READ_ONLY and sig in seen:
                repeats += 1
                stuck = f"反复调用 {c['name']}（参数相同，内容没有变化）"
                result = (f"这次调用与第 {seen[sig]} 步完全相同，内容没有变化，已跳过。不要重复调用；"
                          f"需要更多内容请换参数（如 start_row、offset），否则根据已有信息进行下一步。"
                          f"程序判断的下一步：{ws.next_hint()}")
                steps.append(Step(c["name"], {}, result, False))
                if on_step:
                    on_step(len(steps), steps[-1])
                history.append({"role": "user", "content": f"【{c['name']} 的结果】\n{result}"} if mode == "json"
                               else {"role": "tool", "tool_call_id": c["id"], "content": result})
                skipped += 1
                continue
            try:
                args = json.loads(c["arguments"] or "{}")
                fn = handlers.get(c["name"])
                if not isinstance(args, dict):
                    raise ToolError("参数必须是 JSON 对象")
                if "_invalid_json" in args:
                    raise ToolError(args["_invalid_json"])
                args = _coerce_args(c["name"], args)
                if fn is None:
                    raise ToolError(f"没有工具 {c['name']}，可用：{list(handlers)}")
                result, ok = str(fn(**args)), True
            except ToolError as e:
                result, ok = f"错误：{e}", False
            except json.JSONDecodeError as e:
                result, ok = f"错误：参数不是合法 JSON：{e}", False
            except TypeError as e:
                result, ok = f"错误：参数不对：{e}", False
            if len(result) > MAX_RESULT_CHARS:
                result = result[:MAX_RESULT_CHARS] + "……（已截断）"
            step = Step(c["name"], args if isinstance(args, dict) else {}, result, ok)
            if ok and c["name"] in WRITE_TOOLS:
                ws.changed = True
            if c["name"] in READ_ONLY and ok:
                seen[sig] = len(steps) + 1
            elif c["name"] not in READ_ONLY:
                if (sig, result) == last_write:      # 写入类工具原样重复、结果也一样：同样算原地打转
                    repeats += 1
                    stuck = f"反复提交相同的 {c['name']}，结果：{result[:300]}"
                last_write = (sig, result)
                seen.clear()   # 工作区变了，重新读取是合理的
            fails[c["name"]] = 0 if ok else fails.get(c["name"], 0) + 1
            if ok and c["name"] in ("remove_dataset", "remove_facts"):     # 建了又删、删了又建：没有进展的来回改
                what = str((args or {}).get("name", ""))
                undo[(c["name"], what)] = undo.get((c["name"], what), 0) + 1
                if undo[(c["name"], what)] >= MAX_STUCK:
                    repeats = MAX_STUCK
                    stuck = f"「{what}」反复建了又删（{undo[(c['name'], what)]} 次），没有进展"
                elif undo[(c["name"], what)] == MAX_STUCK - 1:
                    result += (f"\n注意：「{what}」已经删了 {MAX_STUCK - 1} 次。先想清楚它该怎么定义再动手，"
                               f"再删一次程序会停下；拿不准就用 ask_question 问财务")
                    step.result = result
            if ok and c["name"] == "set_rules":      # 保存了但阻断问题一直没改掉：同样算没有进展
                now = getattr(ws, "last_blocking", set())
                same_block = same_block + 1 if now and now & prev_block else (1 if now else 0)
                prev_block = now if same_block <= 1 else now & prev_block
                if same_block >= MAX_STUCK:
                    repeats = MAX_STUCK
                    stuck = f"set_rules 连续 {same_block} 次保存后仍有同样的阻断问题，最后一次：{result[:300]}"
                elif same_block == MAX_STUCK - 1:
                    result += ("\n注意：同样的阻断问题已经连续两次没改掉。再出现一次程序会停下；"
                               "如果是规则写法表达不了材料的要求，用 ask_question 问财务或说明卡在哪里")
                    step.result = result
            if fails[c["name"]] >= MAX_STUCK:
                stuck = f"{c['name']} 连续 {fails[c['name']]} 次失败，最后一次：{result[:300]}"
            steps.append(step)
            if on_step:
                on_step(len(steps), step)
            if mode == "json":
                history.append({"role": "user", "content": f"【{c['name']} 的结果】\n{result}"})
            else:
                history.append({"role": "tool", "tool_call_id": c["id"], "content": result})
        if skipped and skipped == len(calls):
            # 整轮都是被拦下的重复调用：不把它留在历史里，否则小模型会照着历史接着重复；只留一句说明
            del history[turn_start:]
            history.append({"role": "user", "content": f"（你刚才又要原样调用 {calls[0]['name']}，和之前完全相同，程序已拦下、没有执行。"
                                                       f"不要再重复；下一步（程序判断）：{ws.next_hint()}）"})
        if ws.done:
            break
        if repeats >= MAX_STUCK or any(n >= MAX_STUCK for n in fails.values()):
            error = (f"智能体卡住了，已停止以免继续消耗 tokens：{stuck}。"
                     f"已用 {len(steps)} 步。请把这条信息反馈给开发，或在对话里告诉智能体怎么处理")
            break
    else:
        error = f"达到最大步数 {max_steps} 仍未完成"

    if message and not ws.changed and not fresh and not any((changes or {}).values()):
        # 对话里只是问答、没有改动：不算新的一轮，不改案卷和输出
        ar = AgentResult(ws, steps, None, None, error, round=(last or {}).get("n", 0))
        ar.reply = ws.reply_text or (f"（没有完成）{failure.to_text(error)}" if error else "")
        return ar
    tested = {s: ws.results[s] for s in ws.schemes if s in ws.results and s not in ws.untested}
    rss = {s: ws._ruleset(s) for s in tested}
    draft = {}
    if ws.pending:          # 还有问题等答复：不出正式结果，试算结果只作草稿留在内存
        draft, tested = tested, {}
    first = next(iter(tested), None)
    ar = AgentResult(ws, steps, rss.get(first), tested[first]["R0"] if first else None, error,
                     rulesets={s: rss[s] for s in tested}, results=tested)
    ar.draft_results, ar.round, ar.changes = draft, rnd["n"], changes
    ar.reply = ws.reply_text
    ar.pending = [q.model_dump() for q in ws.pending]
    if last and last.get("digest") is not None and tested:
        ar.diff_rows, ar.compared = cs.diff_results(last["digest"], tested), True
    if last and last.get("drafts") is not None:
        ar.rule_changes = cs.diff_rules(last["drafts"], ws.drafts)
    for rs in rss.values():
        h, t, _ = check_quotes(rs, ws.inv.docs)
        ar.quote_hits, ar.quote_total = ar.quote_hits + h, ar.quote_total + t
    if out is not None:
        ar.outputs = write_agent_outputs(ar, out, answers)
        rnd["inputs"] = [str(f) for f in ws.data_files]
        rnd["outputs"] = {k: [p.name, cs.sha256(p)] for k, p in ar.outputs.items()
                          if k == "mapping" or k == "rules" or k.startswith(("rules:", "adjust:"))}
        if message:
            rnd["message"] = {"by": ws.chat_by, "text": message}
        rnd.update(status=ar.status, error=error, pending=[q.id for q in ws.pending], summary=ws.summary,
                   answered=[q.id for q in fresh], changes=changes, drafts=json.loads(json.dumps(ws.drafts)),
                   digest=cs.digest(tested) if tested else (last or {}).get("digest"))
        case = case or {"created": rnd["at"], "rounds": []}
        case["rounds"].append(rnd)
        case["workspace"] = cs.snapshot(ws)
        ar.outputs["case"] = cs.save_case(out, case)
        cs.archive(out, rnd["n"], ar.outputs)
    return ar


def _chat_task(ws: Workspace, message: str, chat: list[dict] | None, fresh: list[Question],
               changes: dict | None) -> str:
    lines = ["财务在对话中给你发了一条消息。工作区（数据表、规则、笔记、台账、问题）是上一轮的状态，已按现有规则试算过，"
             "可以直接用 explain / evaluate 查算式和数字。"]
    if chat:
        lines.append("\n之前的对话：")
        lines += [f"- {'财务' if m.get('role') == 'user' else '你'}：{m.get('text', '')[:500]}" for m in chat[-10:]]
    if fresh:
        lines.append("\n财务在问题表里的新答复（必须照做）：" + "；".join(f"{q.id} → {q.answer}" for q in fresh))
    if changes and any(changes.values()):
        lines.append("\n文件变化：" + "；".join(f"{k} {'、'.join(v)}" for k, v in changes.items() if v)
                     + "（新增、修改的规则文档要读并 record_facts 登记）")
    lines += ["", f"财务的消息：{message}", "",
              "请：",
              "1. 是提问：查清楚后用 reply 回答，金额和算式来自 explain / evaluate，不要自己算；",
              "2. 是要求修改：照做（set_rules、define_dataset 等），他的话本身就是依据；要求不明确、和材料冲突、"
              "或缺信息时不要猜，用 reply 反问；",
              "3. 他明确答复了某个待确认问题：用 answer_question 记下原话，再按答复改规则；",
              "4. 最后调用 reply 结束。改过东西时程序会自动试算并检查，不通过会告诉你原因。",
              "", "当前文件清单：", ws.list_files()]
    return "\n".join(lines)


def _resume_task(ws: Workspace, last: dict | None, changes: dict | None, fresh: list[Question],
                 withdrawn: list[Question] = ()) -> str:
    n = ws.round
    lines = [f"这是同一笔奖金的第 {n} 轮处理（上一轮状态：{(last or {}).get('status', '未知')}）。上一轮的数据表、列映射、规则、"
             "笔记和问题都已恢复到工作区（见工作区状态），没有变化的材料不必重读。", "", "本轮变化："]
    if fresh:
        lines.append("- 财务新答复（必须照做）：")
        lines += [f"  · {q.id} {q.text} → {q.answer}" for q in fresh]
    for k, v in (changes or {}).items():
        if v:
            lines.append(f"- {k}文件：{'、'.join(v)}")
    if withdrawn:
        lines.append("- 程序撤回的问题（冲突所依据的材料已修改，请按新内容重新登记参数）：" + "；".join(q.id for q in withdrawn))
    if not fresh and not any((changes or {}).values()):
        lines.append("- 没有新答复，也没有文件变化")
    if ws.pending:
        lines.append("- 仍在等待答复：" + "；".join(f"{q.id} {q.text}" for q in ws.pending))
    lines += ["", "请：",
              "1. 按新答复修改规则（答复只影响它相关的部分，其余保持不变）；",
              "2. 阅读新增、修改的文件，用 record_facts 登记其中的参数（与已有台账同一参数用相同的名字），"
              "取值不一致时程序会自动提问；口径前后矛盾、新增内容不明确时用 ask_question 提问，写明出处，不要自行取舍；",
              "3. 删除的文件如果被数据表用到，改映射或提问；",
              "4. 试算全部方案，确认结果合理后 finish。还有问题等答复时，做完不受影响的部分就 finish 结束本轮。",
              "", "当前文件清单：", ws.list_files()]
    return "\n".join(lines)


# ---------------- 输出 ----------------
def write_agent_outputs(ar: AgentResult, out: Path, answers: list[dict] | None = None) -> dict[str, Path]:
    out.mkdir(parents=True, exist_ok=True)
    for old in [*out.glob("result*.xlsx"), *out.glob("rules*.yaml"), *out.glob("adjust_*.yaml"),
                out / "待确认问题.xlsx", out / "expected_check.txt"]:
        old.unlink(missing_ok=True)       # 上一轮的输出已存档在 history/，不能留在顶层被误用
    ws, paths = ar.ws, {}
    if ws.mapping:
        paths["mapping"] = out / "mapping.yaml"
        paths["mapping"].write_text(yaml.safe_dump({k: v.model_dump(exclude_none=True, exclude_defaults=True)
                                                    for k, v in ws.mapping.items()}, allow_unicode=True, sort_keys=False),
                                    encoding="utf-8")
    for name, rs in ar.rulesets.items():
        key = "rules" if ar.single else f"rules:{name}"
        paths[key] = out / ("rules.yaml" if ar.single else f"rules_{_safe(name)}.yaml")
        paths[key].write_text(rs.dump_yaml(), encoding="utf-8")
        done = []
        for rev in _revisions(ws._merged(name).get("revisions")):
            done += [{"node_id": a["node"], "amount": str(a["amount"]), "reason": a.get("reason", ""), "id": a.get("id", ""),
                      "created_at": f"{len(done) + i:06d}"} for i, a in enumerate(rev["adjust"])]
            paths[f"adjust:{name}:{rev['name']}"] = out / f"adjust_{_safe(name)}_{_safe(rev['name'])}.yaml"
            paths[f"adjust:{name}:{rev['name']}"].write_text(yaml.safe_dump(list(done), allow_unicode=True, sort_keys=False),
                                                             encoding="utf-8")
    qrows = [{"id": q.id, "round": q.round, "text": q.text, "options": q.options, "suggested": q.suggested,
              "related": q.related, "blocking": q.blocking, "answer": q.answer,
              **({"withdrawn": q.withdrawn} if q.withdrawn else {}),
              **({"answered_by": q.answered_by, "answered_at": q.answered_at} if q.answered_by else {})}
             for q in ws.questions]
    paths["questions"] = out / "questions.yaml"
    paths["questions"].write_text(
        "# 财务答复：在 answer 里填选项序号（从 0 开始）或直接写文字；有补充材料放进材料文件夹，\n"
        "# 然后用同一个输出目录再运行一次，智能体会从这一轮接着做（不用 --answers）\n"
        + yaml.safe_dump(qrows, allow_unicode=True, sort_keys=False), encoding="utf-8")
    paths["log"] = out / "agent_log.json"
    paths["log"].write_text(json.dumps({"summary": ws.summary, "error": ar.error, "transcripts": ws.transcripts,
                                        "steps": [vars(s) for s in ar.steps]}, ensure_ascii=False, indent=2),
                            encoding="utf-8")
    if not ar.results:
        if ar.pending or ws.questions:     # 本轮不出结果：把问题、规划和工作记录给财务看
            import openpyxl
            from bonus_core.report.excel import _sheet
            wb = openpyxl.Workbook()
            wb.remove(wb.active)
            for title, headers, data, widths in _sheets(ar, qrows):
                _sheet(wb, title, headers, data, widths=widths)
            paths["pending"] = out / "待确认问题.xlsx"
            wb.save(paths["pending"])
        return paths
    ar.recheck = "；".join(_recheck(ar, paths, name, rev) for name, revs in ar.results.items() for rev in revs)
    if ar.single:
        paths["excel"] = out / "result.xlsx"
        write_excel(ar.ruleset, ar.result, paths["excel"], extra_sheets=_sheets(ar, qrows))
        return paths
    for name, revs in ar.results.items():
        for rev, res in revs.items():
            k = f"excel:{name}:{rev}"
            paths[k] = out / (f"result_{_safe(name)}.xlsx" if rev == "R0" else f"result_{_safe(name)}_{_safe(rev)}.xlsx")
            write_excel(ar.rulesets[name], res, paths[k])
    paths["excel"] = out / "result.xlsx"
    _overview(ar, qrows, paths["excel"])
    return paths


def _safe(s: str) -> str:
    return "".join(ch if ch.isalnum() or ch in "-_" else "_" for ch in s)


def _overview(ar: AgentResult, qrows: list[dict], path: Path):
    """多个方案 / 修订时的总览：每个对象一行，每个方案（修订）一列。修订是同一方案的不同版本，不能相加。"""
    import openpyxl
    from bonus_core.report.excel import MONEY, _sheet
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    cols = [(n, r) for n, revs in ar.results.items() for r in revs]
    heads = [n if r == "R0" else f"{n} {r}" for n, r in cols]
    rows, seen = [], set()
    for n, revs in ar.results.items():
        for node in revs["R0"].nodes.values():
            if node.id in seen:
                continue
            seen.add(node.id)
            rows.append([node.level, node.key or "", node.label]
                        + [float(ar.results[c][r].nodes[node.id].final) if node.id in ar.results[c][r].nodes else None
                           for c, r in cols])
    _sheet(wb, "各方案结果", ["层级", "编号", "名称"] + heads, rows, money_cols=range(4, 4 + len(cols)))
    summary = [[h, float(Decimal(ar.results[c][r].pool)), f"result_{_safe(c)}" + ("" if r == "R0" else f"_{_safe(r)}") + ".xlsx"]
               for h, (c, r) in zip(heads, cols)]
    ws = _sheet(wb, "方案汇总", ["方案", "奖金池", "明细文件（含计算链）"], summary, money_cols=[2])
    ws.append([])
    ws.append(["说明：同一方案的 R1、R2 是人工调整后的修订版本，与 R0 是同一笔奖金的不同版本，不能相加。"])
    for title, headers, data, widths in _sheets(ar, qrows):
        _sheet(wb, title, headers, data, widths=widths)
    for w in wb.worksheets:
        for row in w.iter_rows(min_row=2):
            for c in row:
                if isinstance(c.value, float):
                    c.number_format = MONEY
    wb.save(path)


def _recheck(ar: AgentResult, paths: dict[str, Path], name: str, rev: str) -> str:
    """从写出的规则、映射（和调整）文件重新读入并重算，逐个对象核对金额，证明结果只来自引擎、可复算。"""
    from bonus_core.data.mapping import load_mapping
    label = "" if ar.single else f"{name}{'' if rev == 'R0' else ' ' + rev}："
    want = ar.results[name][rev]
    try:
        rs = RuleSet.load(paths["rules" if ar.single else f"rules:{name}"])
        data, _ = apply_mapping(rs, ar.ws.data_files, load_mapping(paths["mapping"]))
        adj = []
        if rev != "R0":
            adj = [Adjustment(**a) for a in yaml.safe_load(paths[f"adjust:{name}:{rev}"].read_text(encoding="utf-8"))]
        refs = {f"{s}/{r}": res for s, revs in ar.results.items() for r, res in revs.items() if s != name}
        again = compute(rs, data, adj, refs=refs)
    except Exception as e:  # 复算失败要如实写进结果，不能当作通过
        return f"{label}复算失败：{e}"
    diff = [i for i, n in want.nodes.items() if i not in again.nodes or again.nodes[i].final != n.final]
    if again.pool != want.pool or diff:
        return f"{label}复算不一致：奖金池 {again.pool} / {want.pool}，不一致的对象 {diff[:10]}"
    files = "rules.yaml + mapping.yaml" if ar.single else f"规则、映射{'和调整' if rev != 'R0' else ''}文件"
    return f"{label}一致：按 {files} 重新计算，奖金池和 {len(again.nodes)} 个对象的金额逐分相同"


def _sheets(ar: AgentResult, qrows: list[dict]):
    ws = ar.ws
    status = ("待财务答复（NEEDS_CONFIRMATION）：本轮不出结果。请在 questions.yaml 填写答复（或补充材料）后，"
              "用同一个输出目录再运行一次。问题：" + "；".join(f"{r['id']} {r['text']}" for r in ar.pending)
              if ar.status == "NEEDS_CONFIRMATION" else f"未完成：{ar.error}" if not ar.ok else
              "智能体草稿：规则和数据映射由大模型生成，尚未经财务确认。请核对后再发放。")
    plan = [["状态", status], ["轮次", f"第 {ar.round} 轮"],
            ["方案概述", ws.summary],
            ["计算方案", "、".join(f"{n}（{'、'.join(r)}）" for n, r in ar.results.items()) or "无"],
            ["规则来源", "、".join(d.name for d in ws.inv.docs)]]
    for n, d in ws.datasets.items():
        m = ws.mapping[n]
        kind = f"实体表，主键 {d.key}" if d.key else f"明细表，关联 {d.link}" if d.link else "公司级表"
        plan.append([f"数据表：{n}", f"{m.file}/{m.sheet}" + (f"，筛选 {m.filter}" if m.filter else "") + f"（{kind}）"])
        plan += [[f"  {f}", f"← 「{c}」 {d.fields[f].type}" + (f"，{d.fields[f].unit}" if d.fields[f].unit else "")]
                 for f, c in m.columns.items()]
    plan += [["已忽略", f"{i.get('ref')}：{i.get('reason', '')}"] for i in ws.ignored]
    plan.append(["原文引用命中", f"{ar.quote_hits}/{ar.quote_total}"])
    plan.append(["复算核对", ar.recheck])
    if ar.changes is not None:
        plan.append(["本轮文件变化", "；".join(f"{k}：{'、'.join(v)}" for k, v in ar.changes.items() if v) or "无"])
    if ar.round > 1:
        plan.append(["规则改动（与上一轮比）", "；".join(ar.rule_changes) or "无"])
    q = [[r["id"], r.get("round", 1), r["text"], " / ".join(r.get("options") or []),
          (r["options"][r["suggested"]] if isinstance(r.get("suggested"), int)
           and 0 <= r["suggested"] < len(r.get("options") or []) else ""),
          f"（已撤回：{r['withdrawn']}）" if r.get("withdrawn") else "" if r.get("answer") is None else str(r["answer"])]
         for r in qrows]
    log = [[i, s.tool, json.dumps(s.args, ensure_ascii=False)[:300], ("" if s.ok else "失败：") + s.result[:500]]
           for i, s in enumerate(ar.steps, start=1)]
    out = [("待确认问题", ["编号", "提出轮次", "问题", "选项", "智能体建议（未采用）", "财务答复"], q,
            {"问题": 60, "选项": 40, "智能体建议（未采用）": 20})]
    if ws.facts:
        conflicted = {c.name: c.question for c in ws.conflicts.values()}
        out.append(("参数台账", ["参数", "原文取值", "换算", "文件", "原文", "登记轮次", "冲突问题"],
                    [[f.name, fx.label(f), fx.same_value(f) if fx.number(f.value) is not None else "", f.file, f.quote,
                      f.round, conflicted.get(f.name, "以此为准（原文明确修改）" if f.override else "")]
                     for f in sorted(ws.facts, key=lambda f: f.name)],
                    {"参数": 30, "原文": 60}))
    if ar.compared:
        out.append(("与上一轮对比", ["方案", "修订", "对象", "名称", "上一轮应得", "本轮应得", "差额", "上一轮实发", "本轮实发"],
                    ar.diff_rows or [["（所有对象金额与上一轮相同）"]], {"对象": 16, "名称": 14}))
    return out + [
            ("智能体规划", ["项目", "内容"], plan, {"项目": 22, "内容": 100}),
            ("智能体工作记录", ["步骤", "工具", "参数", "结果"], log, {"参数": 60, "结果": 90})]
