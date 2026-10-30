"""固定回归：场景配置、材料排除、提问后用答复接着跑、逐分比对和“算错没拦住”统计。"""
import json
import shutil

import openpyxl

from bonus_agent import regress as rg
from bonus_agent.llm import FakeLLM
from tests.test_syn_acceptance import CASE, EXPECTED_XLSX, _script, run_case


def _suite(tmp_path, **extra):
    inp = tmp_path / "inputs"
    inp.mkdir()
    shutil.copy(run_case.XLSX, inp)
    (inp / "negative_cases.json").write_text("[]", encoding="utf-8")
    shutil.copy(EXPECTED_XLSX, tmp_path / "期望.xlsx")
    spec = {"name": "S1", "inputs": "inputs", "exclude": ["negative_cases*.json"], "expected": "期望.xlsx", **extra}
    import yaml
    (tmp_path / "suite.yaml").write_text(yaml.safe_dump({"scenarios": [spec]}, allow_unicode=True), encoding="utf-8")
    return rg.load_suite(tmp_path / "suite.yaml")


def test_materials_exclude_answers_expected_and_patterns(tmp_path):
    (tmp_path / "答复.txt").write_text("按制度执行", encoding="utf-8")
    sc = _suite(tmp_path, answers="答复.txt")[0]
    sc.inputs.append(tmp_path)                     # 输入里混进了期望结果和答复文件
    names = [f.name for f in sc.materials()]
    assert run_case.XLSX.name in names and "negative_cases.json" not in names
    assert "期望.xlsx" not in names and "答复.txt" not in names


def test_suite_counts_correct_runs(tmp_path):
    sc = _suite(tmp_path)
    s = rg.run_suite(sc, tmp_path / "out", lambda: FakeLLM(_script()), runs=2, log=lambda _: None)
    assert s["总计"]["次数"] == 2 and s["总计"]["金额正确"] == 2 and s["总计"]["算错没拦住"] == 0
    r = s["明细"][0]
    assert r["status"] == "OK" and r["correct"] and r["checked"] > 100 and r["wrong"] == 0 and r["review_rows"] > 0
    assert json.loads((tmp_path / "out" / rg.SUMMARY).read_text(encoding="utf-8"))["场景"]["S1"]["正确率"] == 1


def test_wrong_amount_is_missed(tmp_path):
    sc = _suite(tmp_path)
    wb = openpyxl.load_workbook(tmp_path / "期望.xlsx")
    ws = wb["个人结果"]
    head = next(r for r in ws.iter_rows() if r[0].value == "方案")
    col = next(c.column for c in head if c.value == "应得金额")
    ws.cell(head[0].row + 1, col).value += 1                 # 改坏一个期望金额：相当于结果算错了
    wb.save(tmp_path / "期望.xlsx")
    s = rg.run_suite(sc, tmp_path / "out", lambda: FakeLLM(_script()), runs=1, log=lambda _: None)
    r = s["明细"][0]
    assert r["status"] == "OK" and r["missed"] and r["wrong"] == 1 and s["总计"]["算错没拦住"] == 1
    assert "不一致" in (tmp_path / "out" / "S1" / "第1次" / "expected_check.txt").read_text(encoding="utf-8")


def test_questions_answered_from_file_then_continue(tmp_path):
    (tmp_path / "答复.txt").write_text("年度奖扣减本年度已发的季度奖和过程奖", encoding="utf-8")
    sc = _suite(tmp_path, answers="答复.txt")
    s = rg.run_suite(sc, tmp_path / "out", lambda: FakeLLM(_script(blocking=True) + _script()), runs=1,
                     log=lambda _: None)
    r = s["明细"][0]
    assert r["answered_rounds"] == 1 and r["rounds"] == 2 and r["questions"] == ["跨期抵扣口径尚未确认：年度奖扣减哪些已发金额？"]
    assert r["status"] == "OK" and r["correct"]
    # 没配答复：记为停下提问，不算算错
    sc[0].answers = None
    s = rg.run_suite(sc, tmp_path / "out2", lambda: FakeLLM(_script(blocking=True)), runs=1, log=lambda _: None)
    assert s["明细"][0]["status"] == "NEEDS_CONFIRMATION" and s["场景"]["S1"]["卡住原因"] == {"等财务答复": 1}
    assert s["总计"]["算错没拦住"] == 0 and s["明细"][0]["correct"] is None


def test_cause_and_answer_text(tmp_path):
    assert rg.cause("智能体卡住了：verify 连续 3 次失败，最后一次：对不上：引擎算出 1") == "验算对不上"
    assert rg.cause("达到最大步数 150") == "步数用完" and rg.cause("别的") == "其他"
    p = tmp_path / "a.json"
    p.write_text(json.dumps({"answers": [{"question_id": "Q1", "answer": "6%"}]}, ensure_ascii=False), encoding="utf-8")
    assert rg.answer_text(p) == "(Q1) 6%"


def test_default_suite_files_exist():
    for sc in rg.load_suite(CASE.parent.parent / "regress.yaml"):
        assert sc.expected.exists() and sc.materials(), sc.name
        assert not any(f.name.startswith("negative_cases") for f in sc.materials())
