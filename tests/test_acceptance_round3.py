"""Q1–Q3 验收发现的问题（2026-10-30）：

- Q3：分档表的分档点按分钟写，传入的却是比值，所有工单都落在第一档；规则和验算用同一张表，验算发现不了；
- Q1：算净利润只减了成本、漏了费用，整批奖金包偏大，验算照着同样的理解写，也发现不了；
- Q2：表名、说明是跨整表的合并单元格时把说明行当成表头；在人员上汇总参与记录的层级变量（按人汇总的实发验算）算不出来；
  evaluate 认不出验算时保存的辅助变量；
- results.json 的已发抵扣是实际抵扣额，和 Excel 发放表、财务口径（已发全额，超发另列）不一致。
"""
import importlib.util
from decimal import Decimal
from pathlib import Path

import openpyxl

from bonus_agent import lint
from bonus_agent.agent import run_agent
from bonus_agent.llm import FakeLLM
from bonus_core.data.table_detect import detect_sheet

_spec = importlib.util.spec_from_file_location("delivery_make", Path(__file__).parent.parent / "evals/cases/delivery_h2/make.py")
mk = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(mk)
_t = importlib.util.spec_from_file_location("delivery_test", Path(__file__).parent / "test_delivery_h2.py")
dt = importlib.util.module_from_spec(_t)
_t.loader.exec_module(dt)


def test_merged_title_rows_are_not_headers():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["人员名单"])
    ws.append(["来源：样例数据，第 4 行为表头。"])
    ws.append([])
    ws.append(["编号", "姓名", "部门", "性质"])
    for i in range(5):
        ws.append([f"S{i}", f"人员{i}", "一部", "正式"])
    ws.merge_cells("A1:D1")
    ws.merge_cells("A2:D2")
    t = detect_sheet(ws, "x.xlsx")
    assert t.header_row == 4 and t.headers == ["编号", "姓名", "部门", "性质"] and len(t.rows) == 5


def _delivery(tmp_path):
    case = mk.make(tmp_path / "case")
    steps = [s for s in dt.script() if s["tool"] != "finish"]
    ar = run_agent(FakeLLM(steps + [{"tool": "finish", "args": {"summary": "x"}}]), [case / "materials"], tmp_path / "out")
    return ar.ws


def test_entity_node_sums_level_variables_and_evaluate_sees_helpers(tmp_path):
    ws = _delivery(tmp_path)
    out = ws.evaluate('sum(raw_weight, of: "参与分配")', "人员:U001", "V1")
    total = Decimal(out.rsplit("=", 1)[1].strip().replace(",", ""))
    assert total > 0
    ws.check_vars["v_double"] = {"scope": "参与记录", "expr": "raw_weight * 2"}
    out2 = ws.evaluate('sum(v_double, of: "参与分配")', "人员:U001", "V1")
    assert Decimal(out2.rsplit("=", 1)[1].strip().replace(",", "")) == total * 2


def test_unused_money_field_and_tier_scale_are_gated(tmp_path):
    ws = _delivery(tmp_path)
    rs = ws._ruleset("V1")
    data, _ = ws._data(rs)
    clean = lint.check(rs, data, "")
    assert not [g for g in clean.gates if "金额字段" in g or "分档表" in g]
    v = dict(rs.variables)
    v["cost"] = v["cost"].model_copy(update={"expr": 'count(of: "成本费用", where: approved == "是") * 0'})
    bad = lint.check(rs.model_copy(update={"variables": v}), data, "")
    assert any("成本费用" in g and "expense_amount" in g for g in bad.gates)
    assert "金额 expense_amount" in bad.keys
    # 拦下时告诉模型放过要写的 ref；照写就放过（Q1：模型自己拼的 ref 对不上，同一条拦了 6 次）
    shown = ws._open_gates(bad)
    assert any("〔放过时 ignored 的 ref 写「金额 expense_amount」〕" in g for g in shown)
    ws.ignored = [{"ref": "金额 expense_amount", "reason": "费用不计入本奖金"}]
    assert not [g for g in ws._open_gates(bad) if "expense_amount" in g]
    # 其他方案用到了就不算漏用
    assert not [g for g in lint.check(rs.model_copy(update={"variables": v}), data, "", used_elsewhere={"expense_amount"}).gates
                if "expense_amount" in g]
    # 分档点按十万写、传入的是比值：全部落在第一档，差几个数量级
    stats = {"contrib_tier": {"n": 6, "min": Decimal("0.2"), "max": Decimal("3.5"), "bands": {0: 6}}}
    rep = lint.check(rs, data, "", tiers=stats)
    assert any("contrib_tier" in g and "量纲" in g for g in rep.gates) and "分档 contrib_tier" in rep.keys
    # 都落在同一档但量级相近：只提示
    near = {"contrib_tier": {"n": 6, "min": Decimal("20000"), "max": Decimal("90000"), "bands": {0: 6}}}
    rep = lint.check(rs, data, "", tiers=near)
    assert not [g for g in rep.gates if "分档表" in g] and any("分档表 contrib_tier" in w for w in rep.warns)
    # 传入很多个值、表有三档以上却全部落在同一档（Q1：毛利率写成成本率，216 个值全落在最高档）：要处理
    many = {"contrib_tier": {"n": 30, "min": Decimal("20000"), "max": Decimal("90000"), "bands": {0: 30}}}
    if len(rs.tables["contrib_tier"].steps) >= 3:
        rep = lint.check(rs, data, "", tiers=many)
        assert any("contrib_tier" in g and "传错了数" in g for g in rep.gates) and "分档 contrib_tier" in rep.keys


def test_engine_records_tier_inputs(tmp_path):
    ws = _delivery(tmp_path)
    st = ws.results["V1"]["R0"].tier_stats["contrib_tier"]
    assert st["n"] > 0 and st["min"] <= st["max"] and sum(st["bands"].values()) == st["n"]
    assert "tier_stats" not in ws.results["V1"]["R0"].to_dict()


def test_unfiltered_status_and_period_are_flagged(tmp_path):
    """读了带审批列、日期列的明细表却没按它们筛选：审批列必须处理，期间只提示；同名字段在别的表上筛选过不算。"""
    ws = _delivery(tmp_path)
    rs = ws._ruleset("V1")
    data, _ = ws._data(rs)
    clean = lint.check(rs, data, "")
    assert not [t for t in clean.gates + clean.warns if "筛选" in t and ("审批" in t or "期间" in t)]
    v = dict(rs.variables)
    v["cost"] = v["cost"].model_copy(update={"expr": 'sum(expense_amount, of: "成本费用")'})   # 收款流水、工时还按 approved 筛
    bad = lint.check(rs.model_copy(update={"variables": v}), data, "")
    assert any("「成本费用」" in g and "approved" in g and "是、否" in g for g in bad.gates) and "状态 approved" in bad.keys
    assert any("「成本费用」" in w and "expense_date" in w and "2026-05" in w for w in bad.warns)
    assert not [g for g in bad.gates if "「补充费用」" in g or "「补充收款」" in g]       # 没读到的补充表不管
