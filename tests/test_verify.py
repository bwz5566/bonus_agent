"""思考-行动-观察 + 独立验算：程序指定对象，模型按材料原文从原始字段另写算式，引擎结果要对到分，否则不能出结果。"""
import openpyxl
import pytest

from bonus_agent.agent import ToolError, Workspace
from bonus_agent.auto import build_inventory
from bonus_agent.llm import FakeLLM
from test_param_tables import case, rules

pytestmark = pytest.mark.verify


def test_targets_are_chosen_by_program_and_finish_waits(tmp_path):
    w = case(tmp_path / "m")
    rules(w, "amount * rate")
    gaps = w._verify_gaps()
    assert [g.split("（")[0] for g in gaps] == ["方案「月度」个人:002", "方案「月度」个人:001"]
    assert "应得金额最高，引擎算出 52 元" in gaps[0] and "应得金额最低的非零" in gaps[1]
    assert "待独立验算（程序指定，全部通过才能 finish）：方案「月度」个人:002" in w.state()
    with pytest.raises(ToolError, match="独立验算还没全部通过"):
        w.finish("x", ignored=[{"ref": "01_规则.md", "reason": "测试"}], review=False)


def test_verify_must_be_independent_and_match_to_the_cent(tmp_path):
    w = case(tmp_path / "m")
    rules(w, "amount * rate")
    with pytest.raises(ToolError, match="引用了规则变量 oc"):
        w.verify("个人:002", 'round(sum(oc, of: "订单") * coef, 2)', "个人提成 = 订单提成合计 × 计奖系数")
    with pytest.raises(ToolError, match="对不上：引擎算出 52"):          # 漏了计奖系数
        w.verify("个人:002", 'sum(round(amount * rate, 2), of: "订单")', "个人提成 = 订单提成合计 × 计奖系数")
    out = w.verify("个人:002", 'round(sum(round(amount * rate, 2), of: "订单") * coef, 2)', "个人提成 = 订单提成合计 × 计奖系数")
    assert "一致" in out and "还要验算" in out
    w.verify("个人:001", 'round(sum(round(amount * rate, 2), of: "订单") * coef, 2)', "个人提成 = 订单提成合计 × 计奖系数")
    assert w._verify_gaps() == []
    # 规则改了（漏乘系数）：之前的验算式按新结果重新比对，不一致就拦下
    w.set_rules(scheme="月度", variables={"pay": {"scope": "个人", "expr": 'round(sum(oc, of: "订单"), 2)'}})
    w.trial_run()
    assert any("重新比对不一致" in g for g in w._verify_gaps())


def test_share_allocation_rounding_remainder_is_accepted(tmp_path):
    """按份额分配 100 元给 3 人：33.34 / 33.33 / 33.33，独立算式 100 / 3 与引擎差的 0.01 是尾差。"""
    d = tmp_path / "m"
    d.mkdir()
    (d / "01_规则.md").write_text("奖金 100 元由三人平分。", encoding="utf-8")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "人员"
    for r in (["工号"], ["E1"], ["E2"], ["E3"]):
        ws.append(r)
    wb.save(d / "02_数据.xlsx")
    w = Workspace(FakeLLM([]), build_inventory(sorted(d.iterdir())))
    w.define_dataset("人员", "02_数据.xlsx", "人员", {"emp_id": {"column": "工号", "type": "id"}}, header_rows=[1], key="emp_id")
    w.set_rules(scheme="S", pool={"expr": "100"}, allocation=[{"level": "个人", "key": "emp_id", "weight": "1"}])
    w.trial_run()
    top, low = [g.split("（")[0].split("」")[1] for g in w._verify_gaps()]
    assert "尾差" in w.verify(top, '100 / count(of: "人员")', "三人平分", at="全部")
    assert "一致" in w.verify(low, '100 / count(of: "人员")', "三人平分", at="全部")
    with pytest.raises(ToolError, match="对不上"):
        w.verify(top, '100 / (count(of: "人员") - 1)', "错的算式", at="全部")


def test_thought_is_recorded_and_not_passed_to_tools():
    from bonus_agent.agent import run_agent
    from test_agent import ANNUAL
    calls = [{"thought": "先看有哪些文件", "tool": "list_files"},
             {"tool": "read_sheet", "args": {"thought": "看花名册表头", "file": "年度数据.xlsx", "sheet": "员工花名册", "rows": 3}},
             {"tool": "read_sheet", "args": {"thought": "换个说法再读一次", "file": "年度数据.xlsx", "sheet": "员工花名册", "rows": 3}},
             {"tool": "note", "args": {"topic": "总包", "text": "净利润 8%"}}]
    ar = run_agent(FakeLLM(calls), [ANNUAL], max_steps=4)
    assert [s.thought for s in ar.steps] == ["先看有哪些文件", "看花名册表头", "换个说法再读一次", ""]
    assert ar.steps[1].ok and "thought" not in ar.steps[1].args
    assert "已跳过" in ar.steps[2].result                    # 只是思考不同、参数相同，仍算重复
    assert "提醒：每次调用都先在 thought 里写思考" in ar.steps[3].result


def test_verify_tolerance_is_only_the_rounding_remainder(tmp_path):
    """尾差只放过按份额分配产生的 0.01；每人按公式直接算的对象差 0.01 也要拦下，份额对象差 0.02 以上也要拦下。"""
    w = case(tmp_path / "m")
    rules(w, "amount * rate")
    with pytest.raises(ToolError, match="对不上"):              # 引擎 52.00；差 0.01（无份额分配）
        w.verify("个人:002", 'round(sum(round(amount * rate, 2), of: "订单") * coef, 2) + 0.01', "x")
    d = tmp_path / "s"
    d.mkdir()
    (d / "01_规则.md").write_text("奖金 100 元由三人平分，0.02 元。", encoding="utf-8")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "人员"
    for r in (["工号"], ["E1"], ["E2"], ["E3"]):
        ws.append(r)
    wb.save(d / "02_数据.xlsx")
    s = Workspace(FakeLLM([]), build_inventory(sorted(d.iterdir())))
    s.define_dataset("人员", "02_数据.xlsx", "人员", {"emp_id": {"column": "工号", "type": "id"}}, header_rows=[1], key="emp_id")
    s.set_rules(scheme="S", pool={"expr": "100"}, allocation=[{"level": "个人", "key": "emp_id", "weight": "1"}])
    s.trial_run()
    top = s._verify_gaps()[0].split("（")[0].split("」")[1]
    with pytest.raises(ToolError, match="对不上"):
        s.verify(top, '100 / count(of: "人员") + 0.02', "x", at="全部")


def test_helpers_reach_across_tables_like_q1(tmp_path):
    """Q1：个人业绩 = Σ(项目销售额 × 参与份额)，项目销售额要先在项目层汇总。不能借用规则变量时，
    用 helpers 自己在项目层定义辅助变量；辅助变量不能和规则变量同名，也不能引用规则变量。"""
    d = tmp_path / "m"
    d.mkdir()
    (d / "01_规则.md").write_text("个人奖金 = Σ(项目销售额 × 参与份额) × 1%。", encoding="utf-8")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "员工"
    for r in (["工号"], ["E1"], ["E2"]):
        ws.append(r)
    for name, rows in (("项目", [["项目号"], ["P1"], ["P2"]]),
                       ("销售明细", [["单号", "项目号", "销售额"], ["S1", "P1", 1000], ["S2", "P1", 3000], ["S3", "P2", 2000]]),
                       ("项目参与", [["参与号", "项目号", "工号", "份额"], ["A1", "P1", "E1", 0.5], ["A2", "P1", "E2", 0.5],
                                     ["A3", "P2", "E1", 1]])):
        t = wb.create_sheet(name)
        for r in rows:
            t.append(r)
    wb.save(d / "02_数据.xlsx")
    w = Workspace(FakeLLM([]), build_inventory(sorted(d.iterdir())))
    f = lambda c, t: {"column": c, "type": t}   # noqa: E731
    w.define_dataset("员工", "02_数据.xlsx", "员工", {"emp_id": f("工号", "id")}, header_rows=[1], key="emp_id")
    w.define_dataset("项目", "02_数据.xlsx", "项目", {"project_id": f("项目号", "id")}, header_rows=[1], key="project_id")
    w.define_dataset("销售明细", "02_数据.xlsx", "销售明细", {"sale_id": f("单号", "id"), "project_id": f("项目号", "id"),
                                                        "sales_amount": f("销售额", "money")}, header_rows=[1], link="project_id")
    w.define_dataset("项目参与", "02_数据.xlsx", "项目参与", {"part_id": f("参与号", "id"), "project_id": f("项目号", "id"),
                                                        "emp_id": f("工号", "id"), "share": f("份额", "ratio")},
                     header_rows=[1], link="emp_id")
    w.set_rules(scheme="S", variables={"proj_sales": {"scope": "项目", "expr": 'sum(sales_amount, of: "销售明细")'},
                                       "emp_sales": {"scope": "个人", "expr": 'sum(proj_sales * share, of: "项目参与")'}},
                pool={"expr": "round(emp_sales * 1%, 2)", "scope": "个人"}, allocation=[{"level": "个人", "key": "emp_id"}])
    w.trial_run()
    assert w.results["S"]["R0"].nodes["个人:E1"].final == "40.00"          # (4000×0.5 + 2000×1) × 1%
    helpers = {"v_proj": {"scope": "项目", "expr": 'sum(sales_amount, of: "销售明细")'}}
    assert "一致" in w.verify("个人:E1", 'round(sum(v_proj * share, of: "项目参与") * 1%, 2)', "个人奖金 = Σ(项目销售额 × 参与份额) × 1%",
                              helpers=helpers)
    with pytest.raises(ToolError, match="同名"):
        w.verify("个人:E1", 'round(sum(proj_sales * share, of: "项目参与") * 1%, 2)', "x", helpers={"proj_sales": helpers["v_proj"]})
    with pytest.raises(ToolError, match="引用了规则变量 proj_sales"):
        w.verify("个人:E1", 'round(sum(v_proj * share, of: "项目参与") * 1%, 2)', "x",
                 helpers={"v_proj": {"scope": "项目", "expr": "proj_sales"}})
