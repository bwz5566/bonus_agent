import json
from pathlib import Path

import pytest
import yaml

from bonus_agent.extract_rules import check_quotes, extract
from bonus_agent.llm import FakeLLM, LLMError, parse_json
from bonus_agent.parsing.docs import parse_any
from bonus_core.dsl.models import RuleSet
from evals.run_eval import fake_llm_for, load_case, main as eval_main, run_one

ROOT = Path(__file__).parent.parent
CASE = ROOT / "evals" / "cases" / "sales_q3"
RULES = RuleSet.load(ROOT / "examples" / "sales_q3" / "rules.yaml")


def test_parse_json_tolerant():
    assert parse_json('```json\n{"a": 1}\n```') == {"a": 1}
    assert parse_json('好的，结果如下：{"a": {"b": "}"}} 以上') == {"a": {"b": "}"}}
    with pytest.raises(LLMError):
        parse_json("没有 JSON")


def test_parse_case_documents():
    blocks = [b for f in ("激励方案.docx", "系数表.pptx") for b in parse_any(CASE / f).blocks]
    kinds = [b.kind for b in blocks]
    assert kinds.count("image") == 2 and "table" in kinds
    assert any("1千万" in b.text for b in blocks)


def _answer(**override):
    base = {
        "tables": {"t": {"kind": "tier", "mode": "whole", "steps": [[0, 0], [10000000, 0.1]],
                         "sources": [{"doc": "激励方案.docx", "quote": "以部门本季度合同销售额合计为基数"}]}},
        "variables": {"total": {"scope": "global", "expr": "sum(sales_amount)"}},
        "pool": {"expr": "tier(t, total)"},
        "allocation": [{"level": "小组", "key": "group_id", "weight": "group_score"},
                       {"level": "个人", "key": "employee_id", "parent_key": "group_id", "weight": "share(seniority)"}],
        "payout": [],
        "questions": [{"text": "净利润是全额还是超额？", "options": ["全额", "超额"]}],
    }
    base.update(override)
    return json.dumps(base, ensure_ascii=False)


def test_extract_repairs_invalid_output():
    bad = _answer(pool={"expr": "tier(t, total) + unknown_name"})
    llm = FakeLLM(["图片转写", "图片转写", bad, _answer()])
    res = extract(llm, [CASE / "激励方案.docx", CASE / "系数表.pptx"], RULES.datasets)
    assert res.ok and res.attempts == 2
    assert "unknown_name" in json.dumps(llm.calls[-1][-1]["content"], ensure_ascii=False)
    assert res.quote_hits == 1 and res.quote_total == 1
    assert len(res.transcripts) == 2


def test_hallucinated_quote_is_flagged():
    ans = _answer(tables={"t": {"kind": "tier", "mode": "whole", "steps": [[0, 0], [10000000, 0.1]],
                                "sources": [{"doc": "激励方案.docx", "quote": "文档里根本没有这句话"}]}})
    res = extract(FakeLLM(["转写", "转写", ans]), [CASE / "激励方案.docx", CASE / "系数表.pptx"], RULES.datasets)
    assert res.quote_hits == 0 and any(i.code == "quote_not_found" for i in res.issues)


def test_eval_pipeline_with_fake_model(tmp_path):
    case = load_case(CASE)
    row = run_one(fake_llm_for(case), case, "two_step")
    assert row["ok"] and row["check_score"] == 1.0 and row["leaf_match_rate"] == 1.0
    assert eval_main(["--fake", "--out", str(tmp_path)]) == 0
    assert list(tmp_path.glob("*.md"))


def test_secrets_file_and_model_groups(tmp_path, monkeypatch):
    from bonus_agent.llm import LLM, load_config, resolve_models
    monkeypatch.delenv("DASHSCOPE_API_KEY", raising=False)
    (tmp_path / "models.yaml").write_text((ROOT / "bonus_agent" / "config" / "models.yaml").read_text(encoding="utf-8"), encoding="utf-8")
    (tmp_path / "secrets.yaml").write_text("providers:\n  dashscope:\n    api_key: sk-from-file\n", encoding="utf-8")
    cfg = load_config(tmp_path / "models.yaml")
    llm = LLM.from_config("qwen3.6-35b-a3b", cfg)
    assert llm.api_key == "sk-from-file" and llm.extra_body == {"enable_thinking": False}
    monkeypatch.setenv("DASHSCOPE_API_KEY", "sk-from-env")
    assert LLM.from_config("qwen3.5-27b", cfg).api_key == "sk-from-env"
    assert resolve_models(None, cfg) == resolve_models("candidates", cfg) == cfg["candidates"]
    assert resolve_models("candidates,open", cfg) == cfg["groups"]["open"]
    assert resolve_models("reference,qwen3.6-27b", cfg) == ["qwen3-vl-plus", "qwen3.6-flash", "qwen3.6-27b"]
    with pytest.raises(LLMError):
        resolve_models("no-such-model", cfg)


ANNUAL = ROOT / "evals" / "cases" / "annual_2026"


def test_annual_reference_golden():
    from bonus_core.golden import run_golden
    rs = RuleSet.load(ANNUAL / "rules.yaml")
    assert all(r.passed for r in run_golden(rs))


def test_auto_end_to_end_and_rerun(tmp_path):
    import openpyxl
    import yaml as _yaml
    from bonus_agent.auto import run_auto
    from bonus_core.cli import main as bonus_main

    case = load_case(ANNUAL)
    ar = run_auto(fake_llm_for(case), [ANNUAL], tmp_path)   # 传文件夹，自动区分规则和数据
    assert ar.ok, ar.error
    assert ar.result.pool == case.ref_result.pool
    wb = openpyxl.load_workbook(ar.outputs["excel"])
    assert {"待确认问题", "智能体规划", "汇总"} <= set(wb.sheetnames)
    qs = _yaml.safe_load(ar.outputs["questions"].read_text(encoding="utf-8"))
    assert qs and qs[0]["answer"] is None
    # 输出的规则和映射可以直接用计算引擎复算
    out = tmp_path / "rerun.xlsx"
    assert bonus_main(["run", "--rules", str(ar.outputs["rules"]), "--mapping", str(ar.outputs["mapping"]),
                       "--out", str(out), str(ANNUAL / "年度数据.xlsx")]) == 0
    # 带答复重跑：答复进入提示词
    qs[0]["answer"] = 1
    llm = fake_llm_for(case)
    run_auto(llm, [ANNUAL], tmp_path / "2", answers=qs)
    assert any("财务已确认" in json.dumps(m, ensure_ascii=False) for c in llm.calls for m in c)


def test_plan_errors_are_repaired():
    from bonus_agent.auto import build_inventory, collect_files, make_plan
    from evals.run_eval import fake_plan
    case = load_case(ANNUAL)
    good = fake_plan(case)
    bad = json.loads(json.dumps(good, ensure_ascii=False))
    bad["datasets"]["员工"]["fields"]["monthly_salary"]["column"] = "工资"          # 不存在的列
    bad["datasets"]["员工"]["fields"]["name"]["type"] = "money"                    # 类型与数据不符
    bad["datasets"]["员工"]["fields"]["months"]["column"] = "本年在职 月数"         # 空格差异，应自动更正
    inv = build_inventory(collect_files([ANNUAL]))
    llm = FakeLLM([json.dumps(bad, ensure_ascii=False), json.dumps(good, ensure_ascii=False)])
    po = make_plan(llm, inv)
    assert po.attempts == 2 and not po.errors
    feedback = llm.calls[1][-1]["content"]
    assert "工资" in feedback and "name" in feedback and "本年在职 月数" not in feedback


def test_collect_files_skips_program_and_output_dirs(tmp_path):
    from bonus_agent.auto import TooManyFiles, collect_files
    (tmp_path / ".venv" / "lib").mkdir(parents=True)
    (tmp_path / ".venv" / "lib" / "x.xlsx").write_bytes(b"")
    (tmp_path / "out").mkdir()
    (tmp_path / "out" / "result.xlsx").write_bytes(b"")
    (tmp_path / "~$temp.docx").write_bytes(b"")
    (tmp_path / "方案.docx").write_bytes(b"")
    (tmp_path / "说明.exe").write_text("x")
    assert [f.name for f in collect_files([tmp_path], exclude=[tmp_path / "out"])] == ["方案.docx"]
    for i in range(5):
        (tmp_path / f"{i}.png").write_bytes(b"")
    with pytest.raises(TooManyFiles):
        collect_files([tmp_path], max_files=3)


def test_plan_accepts_month_labels_as_dates():
    from bonus_agent.auto import build_inventory, collect_files, make_plan
    from evals.run_eval import fake_plan
    plan = fake_plan(load_case(ANNUAL))
    plan["datasets"]["月度利润"]["fields"]["month"]["type"] = "date"   # 数据是 "2026-01" 这样的年月
    po = make_plan(FakeLLM([json.dumps(plan, ensure_ascii=False)]), build_inventory(collect_files([ANNUAL])))
    assert not po.errors, po.errors


def _agent_case():
    from evals.run_eval import fake_agent_script
    case = load_case(ANNUAL)
    return case, fake_agent_script(case)


def test_agent_native_end_to_end(tmp_path):
    import openpyxl
    from bonus_agent.agent import run_agent
    case, script = _agent_case()
    # 模型先犯错再自己修正：列名写错、没试算就 finish
    bad = json.loads(json.dumps(next(s for s in script if s["tool"] == "define_dataset"), ensure_ascii=False))
    first_field = next(iter(bad["args"]["fields"]))
    bad["args"]["fields"][first_field]["column"] = "不存在的列"
    i = script.index(next(s for s in script if s["tool"] == "define_dataset"))
    script = script[:i] + [bad] + script[i:]
    j = next(k for k, s in enumerate(script) if s["tool"] == "trial_run")
    script = script[:j] + [{"tool": "finish", "args": {"summary": "x"}}] + script[j:]
    script.insert(j + 2, {"tool": "evaluate", "args": {"expr": "sum(net_profit)"}})
    script.insert(j + 3, {"tool": "explain", "args": {"id": "个人:E2001"}})
    steps = []
    ar = run_agent(FakeLLM(script), [ANNUAL], tmp_path, on_step=lambda n, s: steps.append(s))
    # 第 1 轮：对全额 / 超额提了问题，等财务答复，不出结果
    assert ar.status == "NEEDS_CONFIRMATION" and not ar.results and ar.draft_results, ar.error
    assert not (tmp_path / "result.xlsx").exists() and (tmp_path / "待确认问题.xlsx").exists()
    failed = [s for s in steps if not s.ok]
    assert [s.tool for s in failed] == ["define_dataset", "finish"]
    # 财务在 questions.yaml 填答复后，用同一个目录再跑：从上一轮接着做，只需试算和完成
    q = yaml.safe_load((tmp_path / "questions.yaml").read_text(encoding="utf-8"))
    q[0]["answer"] = 0
    (tmp_path / "questions.yaml").write_text(yaml.safe_dump(q, allow_unicode=True), encoding="utf-8")
    llm2 = FakeLLM([{"tool": "trial_run"}, {"tool": "finish", "args": {"summary": "x"}}])
    ar = run_agent(llm2, [ANNUAL], tmp_path)
    assert ar.ok and ar.status == "OK" and ar.round == 2, ar.error
    task = llm2.calls[0][1]["content"]
    assert "第 2 轮" in task and "全额累进" in task and "没有变化的材料不必重读" in task
    assert ar.result.pool == case.ref_result.pool
    assert "不存在的列" in failed[0].result and "trial_run" in failed[1].result
    ev = next(s for s in steps if s.tool == "evaluate")
    assert "35,877,700.00" in ev.result or "35877700" in ev.result.replace(",", "")
    assert "权重" in next(s for s in steps if s.tool == "explain").result
    wb = openpyxl.load_workbook(ar.outputs["excel"])
    assert {"待确认问题", "智能体规划", "智能体工作记录"} <= set(wb.sheetnames)
    assert ar.recheck.startswith("一致"), ar.recheck            # 从写出的文件复算，逐分一致


def test_agent_json_mode_and_state_in_every_turn():
    from bonus_agent.agent import run_agent
    case, script = _agent_case()
    texts = [json.dumps({"thought": "", "tool": s["tool"], "args": s.get("args", {})}, ensure_ascii=False) for s in script]
    llm = FakeLLM(["这不是 JSON"] + texts)
    llm.tool_mode, llm.images = "json", "（图片转写）"
    q = next(s["args"] for s in script if s["tool"] == "ask_question")      # 财务已答复过同一个问题
    ar = run_agent(llm, [ANNUAL], answers=[{"text": q["text"], "options": q["options"], "answer": 0}])
    assert ar.ok, ar.error
    agent_calls = [c for c in llm.calls if "材料包" in json.dumps(c, ensure_ascii=False)]
    assert all("【当前工作区状态】" in c[-1]["content"] for c in agent_calls)
    assert "格式不对" in json.dumps(agent_calls[1], ensure_ascii=False)


def test_agent_history_shrink_keeps_prefix_until_needed():
    from bonus_agent.agent import _shrink, _size
    h = []
    for i in range(30):
        h += [{"role": "assistant", "content": "", "tool_calls": [
                  {"id": str(i), "type": "function", "function": {"name": "set_rules", "arguments": "{" + "a" * 800 + "}"}}]},
              {"role": "tool", "tool_call_id": str(i), "content": f"结果{i}" + "x" * 1000}]
    small = [dict(m) for m in h[:4]]
    assert not _shrink(small, 100000) and small == h[:4]     # 未超限：不动，前缀保持不变
    assert _shrink(h, 40000)
    assert h[0]["role"] == "assistant" and _size(h) <= 20000
    assert h[-1]["content"].endswith("x" * 100)              # 最近的结果保留全文
    assert sum("已省略" in m.get("content", "") for m in h) > 0


def test_explicit_cache_markers():
    from bonus_agent.llm import LLM
    llm = LLM("m", "m", "", "", cache="explicit")
    out = llm._prepare([{"role": "user", "content": "任务", "_cache": True}, {"role": "user", "content": "状态"}])
    assert out[0]["content"][0]["cache_control"] == {"type": "ephemeral"} and out[1]["content"] == "状态"
    plain = LLM("m", "m", "", "")._prepare([{"role": "user", "content": "任务", "_cache": True}])
    assert plain == [{"role": "user", "content": "任务"}]


def test_agent_stops_looping_and_paginates(monkeypatch):
    from bonus_agent.agent import run_agent
    same = {"tool": "read_sheet", "args": {"file": "年度数据.xlsx", "sheet": "员工花名册", "rows": 5}}
    import bonus_agent.agent as agent_mod
    monkeypatch.setattr(agent_mod, "PIN_SHEET_ROWS", 10)       # 花名册不整表附进材料包，才需要翻页
    llm = FakeLLM([{"tool": "note", "args": {"topic": "总包", "text": "净利润 8%"}}] + [same] * 10)
    ar = run_agent(llm, [ANNUAL], max_steps=30)
    assert not ar.ok and "卡住" in ar.error and "反复调用 read_sheet" in ar.error and len(ar.steps) == 5
    first = ar.steps[1].result
    assert "start_row=6" in first and "共 28 行" in first
    assert "已跳过" in ar.steps[2].result
    assert "净利润 8%" in llm.calls[-1][-1]["content"]       # 笔记每轮附在状态里


def test_agent_token_budget():
    from bonus_agent.agent import run_agent
    llm = FakeLLM([{"tool": "list_files"}])
    llm.usage.prompt_tokens = 2_000_000
    ar = run_agent(llm, [ANNUAL], max_tokens=1_000_000)
    assert not ar.ok and "tokens" in ar.error and not ar.steps


def test_briefing_pins_short_docs_and_small_sheets():
    from bonus_agent.agent import Workspace
    from bonus_agent.auto import build_inventory, collect_files
    inp = ROOT / "evals" / "cases" / "syn_acceptance" / "inputs"
    ws = Workspace(FakeLLM(["（图片转写）"]), build_inventory(collect_files([inp])))
    ws.transcribe_all()
    pack = ws.briefing()
    # 规则文档、确认单、参数表整份附上；大的 JSON 和长数据表只附开头
    assert {"01_奖金制度.docx", "02_规则补充.pptx", "finance_answers.json", "毛利率系数图.png",
            "04_计算参数.xlsx/阶梯与系数"} <= ws.pinned
    assert "business_data.json" not in ws.pinned and "03_业务输入.xlsx/销售明细" not in ws.pinned
    assert "Q_OFFSET" in pack and "R_Q_TIER" in pack and "（图片转写）" in pack and len(pack) < 30000
    assert "材料包" in ws.read_document("01_奖金制度.docx")
    assert "材料包" in ws.read_sheet("04_计算参数.xlsx", "数值参数")
    assert "第7行" in ws.read_sheet("03_业务输入.xlsx", "销售明细", start_row=7, rows=3)
