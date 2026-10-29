"""结果文件里的「过程复核」表：取数列是引擎值，公式列是 Excel 公式，最后用公式复算应得 / 权重并和引擎结果比对。"""
import importlib.util
import shutil
import subprocess
from pathlib import Path

import openpyxl
import pytest

from bonus_agent.agent import run_agent
from bonus_agent.llm import FakeLLM

HERE = Path(__file__).parent


def _mod(name):
    spec = importlib.util.spec_from_file_location(name, HERE / f"{name}.py")
    m = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(m)
    return m


def _recalc(path: Path, outdir: Path) -> openpyxl.Workbook:
    subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", str(outdir), str(path)],
                   check=True, capture_output=True, timeout=120)
    return openpyxl.load_workbook(outdir / path.name, data_only=True)


def test_per_person_review_sheet(tmp_path):
    p = _mod("test_per_person")
    ar = run_agent(FakeLLM(p.script()), [p.make(tmp_path / "m")], tmp_path / "out")
    assert ar.status == "OK", ar.error
    ws = openpyxl.load_workbook(tmp_path / "out" / "result.xlsx")["个人过程复核"]
    heads = [c.value for c in ws[2]]
    assert heads[:2] == ["编号", "名称"] and heads[-3:] == ["应得（Excel 公式复算）", "应得（引擎结果）", "核对"]
    amount = heads.index("amount")
    assert ws.cell(3, amount + 1).value == "=ROUND(MAX(((C3+(D3*12.5))-E3),0),2)"
    assert ws.cell(1, amount + 1).value == "公式" and ws.cell(1, 3).value == "取数（引擎）"
    if shutil.which("soffice"):
        calc = _recalc(tmp_path / "out" / "result.xlsx", tmp_path / "calc")["个人过程复核"]
        assert [r[-1] for r in calc.iter_rows(min_row=3, max_row=5, values_only=True)] == ["一致"] * 3


@pytest.mark.skipif(not shutil.which("soffice"), reason="需要 LibreOffice 重算公式")
def test_tier_lookup_formulas_match_engine(tmp_path):
    t = _mod("test_delivery_h2")
    case = t.mk.make(tmp_path / "case")
    ar = run_agent(FakeLLM(t.script()), [case / "materials"], tmp_path / "out")
    assert ar.status == "OK", ar.error
    for f in sorted((tmp_path / "out").glob("result_*.xlsx")):
        wb = _recalc(f, tmp_path / "calc")
        for ws in [w for w in wb.worksheets if w.title.endswith("过程复核")]:
            heads = [c.value for c in ws[2]]
            rows = [r for r in ws.iter_rows(min_row=3, values_only=True) if r[0] and r[0] != "不一致的个数"]
            assert rows and all(r[heads.index("核对")] in ("一致", "人工调整") for r in rows), (f.name, ws.title)
