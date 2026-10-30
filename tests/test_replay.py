"""用真实模型的调用记录回放：模型怎么写就怎么喂给程序，看防错关能不能拦住真实的错误写法，结果和财务的期望文件逐分比对。

fixtures/q55：Q55 第 1 轮真实模型的 17 步调用。它把“实发 = max(应得 − 预发, 0)”写进了奖金池，当时程序放行了，
导出的「个人发放」表应得、已发抵扣、超预发三列都是错的（实发碰巧对）。
"""
import copy
import importlib.util
import json
from pathlib import Path

import openpyxl
import pytest

from bonus_agent.agent import run_agent
from bonus_agent.llm import FakeLLM

pytestmark = pytest.mark.verify
FIX = Path(__file__).parent / "fixtures" / "q55"
spec = importlib.util.spec_from_file_location("q55_make", FIX / "make.py")
q55 = importlib.util.module_from_spec(spec)
spec.loader.exec_module(q55)


def real_steps() -> list[dict]:
    return [{"thought": s["thought"], "tool": s["tool"], "args": s["args"]}
            for s in json.loads((FIX / "steps.json").read_text(encoding="utf-8"))["steps"]]


def test_real_q55_run_is_now_blocked_at_finish(tmp_path):
    """原样回放：程序要在 finish 时拦下（预发扣在奖金池里），不能再出结果。"""
    ar = run_agent(FakeLLM(real_steps()), [q55.make(tmp_path / "m")], tmp_path / "out", max_steps=20)
    assert ar.status != "OK"
    fin = [s for s in ar.steps if s.tool == "finish"]
    assert fin and not fin[0].ok and "直接扣减了已发" in fin[0].result, [s.result[:200] for s in fin]
    paid = {k.split(":")[1]: float(p.paid) for k, p in ar.ws.results["默认"]["R0"].payouts.items()}
    assert paid == {k: v["实发"] for k, v in q55.expected().items()}       # 实发碰巧对：所以只看实发发现不了


def corrected_steps() -> list[dict]:
    """同一次运行，只把规则改成正确的写法：应得进奖金池，预发写成已发抵扣；再按程序指定的对象独立验算。"""
    steps = real_steps()
    rules = next(s for s in steps if s["tool"] == "set_rules")
    rules["args"] = copy.deepcopy(rules["args"])
    rules["args"]["pool"] = {"expr": "earned_commission", "scope": "个人"}
    rules["args"]["offsets"] = [{"id": "预发", "scope": "个人", "expr": 'sum(prepay_amount, of: "预发记录")'}]
    earned = ('round(sum(round(min(contract_amount, max(sum(receipt_amount, of: "回款流水", where: approval_status == "已审批" '
              'and between(receipt_date, "2026-09-01", "2026-09-30")), 0)) * rate, 2), of: "订单") * bonus_coef, 2)')
    net = f'max({earned} - sum(prepay_amount, of: "预发记录"), 0)'
    basis = "应得提成 = ROUND(订单提成合计 × 人员表计奖系数, 2)；本次实发 = max(应得提成 − 预发合计, 0)"
    end = next(i for i, s in enumerate(steps) if s["tool"] == "verify")
    checks = [("个人:001", earned, "应得"), ("个人:003", earned, "应得"), ("个人:004", earned, "应得"),
              ("个人:001", net, "实发"), ("个人:002", net, "实发")]
    steps = steps[:end] + [{"thought": f"独立验算 {n} 的{a}", "tool": "verify",
                            "args": {"node": n, "expr": e, "basis": basis, "amount": a}} for n, e, a in checks]
    return steps + [{"thought": "全部验算通过", "tool": "finish", "args": {"summary": "2026 年 9 月销售提成"}}]


def test_corrected_q55_matches_expected_file_to_the_cent(tmp_path):
    ar = run_agent(FakeLLM(corrected_steps()), [q55.make(tmp_path / "m")], tmp_path / "out", max_steps=30)
    assert ar.status == "OK", (ar.error, [(s.tool, s.result[:300]) for s in ar.steps if not s.ok])
    ws = openpyxl.load_workbook(tmp_path / "out" / "result.xlsx", data_only=True)["个人发放"]
    head = [c.value for c in ws[1]]
    got = {}
    for r in ws.iter_rows(min_row=2, values_only=True):
        row = dict(zip(head, r))
        got[row["节点"].split(":")[1]] = {"应得": row["应得"], "预发": row["已发抵扣"], "实发": row["本次实发"], "超预发": row["超发待核对"]}
    assert got == q55.expected()


@pytest.mark.verify
def test_q1_rerun_verify_calls_now_agree_with_engine():
    """Q1 新跑（最新代码三轮都停下）：用模型当时写的规则和全部 verify 调用回放。按财务理解写的同级合计、份额
    （share(小组权重)、sum(小组权重, of: "小组", where: ...)）现在和引擎对到分；真错的两次仍拦下：
    把中间结果直接写成数字的（第 2 轮第 16 步），以及个人权重里贡献评分漏了 /100 的（第 2 轮第 19 步）。"""
    import json as _json
    from bonus_agent import case as cs
    from bonus_agent.agent import ToolError, Workspace
    from bonus_agent.auto import build_inventory
    from bonus_agent.llm import FakeLLM
    fx = Path(__file__).parent / "fixtures" / "q1_rerun"
    inp = Path(__file__).parent.parent / "evals" / "cases" / "syn_acceptance" / "inputs"
    ws = Workspace(FakeLLM([]), build_inventory(sorted(inp.iterdir())))
    cs.restore(ws, _json.loads((fx / "workspace.json").read_text(encoding="utf-8")))
    ws.untested = set(ws.schemes)
    ws.trial_run()
    got = {}
    for c in _json.loads((fx / "verify_calls.json").read_text(encoding="utf-8")):
        try:
            ws.verify(**c["args"])
            got[(c["round"], c["step"])] = "通过"
        except ToolError as e:
            got[(c["round"], c["step"])] = str(e)[:60]
    bad = {k: v for k, v in got.items() if v != "通过"}
    assert set(bad) == {(2, 16), (2, 19)}, got
    assert "在材料里找不到" in bad[(2, 16)] and "对不上" in bad[(2, 19)]
