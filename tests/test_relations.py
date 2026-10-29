"""多表关联：关联实体（项目、销售单）、编号引用、日期期间、按部门算奖金包、行筛选、数据校验。"""
import json
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest
import yaml

from bonus_core.data.mapping import apply_mapping, load_mapping
from bonus_core.dsl.models import RuleSet
from bonus_core.dsl.validator import validate
from bonus_core.engine.issues import EngineError
from bonus_core.engine.run import Adjustment, compute

CASE = Path(__file__).parent.parent / "evals" / "cases" / "acceptance_q3"
XLSX = CASE / "业务输入.xlsx"
EXPECTED = json.loads((CASE / "expected.json").read_text(encoding="utf-8"))


def load(scheme: str) -> RuleSet:
    y = lambda n: yaml.safe_load((CASE / n).read_text(encoding="utf-8"))  # noqa: E731
    shared, own = y("weights.yaml"), y(f"rules_{scheme}.yaml")
    own["datasets"] = y("datasets.yaml")
    own["tables"] = {**shared["tables"], **own.get("tables", {})}
    own["variables"] = {**shared["variables"], **own.get("variables", {})}
    own["allocation"], own["checks"] = shared["allocation"], shared["checks"]
    return RuleSet.model_validate(own)


def run(scheme: str, xlsx: Path = XLSX, adjustments=()):
    rs = load(scheme)
    data, issues = apply_mapping(rs, [xlsx], load_mapping(CASE / "mapping.yaml"))
    assert not [i for i in issues if i.severity == "blocking"], issues
    return rs, compute(rs, data, adjustments)


@pytest.mark.parametrize("scheme", ["quarter", "process"])
def test_matches_independent_reference(scheme):
    rs, res = run(scheme)
    assert not [i for i in validate(rs) if i.severity == "blocking"]
    want = EXPECTED[scheme.upper()]
    got = {k: n.final for k, n in res.nodes.items() if k in want}
    assert got == want
    for k, w in EXPECTED["weights"].items():
        assert Decimal(res.nodes[k].weight) == Decimal(w), k
    # 部门 = 小组合计 = 个人合计
    for d in ("D01", "D02"):
        groups = [n for n in res.nodes.values() if n.parent == f"部门:{d}"]
        people = [n for n in res.nodes.values() if n.parent in {g.id for g in groups}]
        assert sum(Decimal(g.final) for g in groups) == sum(Decimal(p.final) for p in people) == Decimal(want[f"部门:{d}"])


def test_trace_reaches_project_and_rows():
    _, res = run("quarter")
    t = res.traces["项目:P01/q_eff"]
    assert "项目:P01/q_sales" in t.inputs and "lookup" not in t.value
    assert any("销售明细" in r for r in res.traces["项目:P01/q_sales"].rows)
    assert "部门:D01/奖金池" in res.traces


def test_department_pool_is_fixed_and_adjustments_stay_inside():
    _, base = run("quarter")
    _, r1 = run("quarter", adjustments=[Adjustment("小组:G01", Decimal("300000"))])
    assert r1.nodes["小组:G01"].final == "300000.00"
    assert r1.nodes["部门:D01"].final == base.nodes["部门:D01"].final
    assert r1.nodes["部门:D02"].final == base.nodes["部门:D02"].final
    assert all(r1.nodes[k].final == base.nodes[k].final for k in ("个人:E005", "个人:E007"))   # 其他部门不受影响
    from bonus_core.engine.allocation import AdjustConflict
    with pytest.raises(AdjustConflict, match="规则直接算出"):
        run("quarter", adjustments=[Adjustment("部门:D01", Decimal("1"))])


def test_checks_block_bad_shares(tmp_path):
    wb = openpyxl.load_workbook(XLSX)
    ws = wb["项目参与"]
    ws.cell(2, 3).value = 0.45            # P01 的参与比例合计变成 0.95
    bad = tmp_path / "业务输入.xlsx"
    wb.save(bad)
    with pytest.raises(EngineError, match="share_sum.*项目:P01"):
        run("quarter", bad)


def test_unapproved_scores_and_filters_are_respected():
    rs = load("quarter")
    data, _ = apply_mapping(rs, [XLSX], load_mapping(CASE / "mapping.yaml"))
    assert {r.values["dept_id"] for r in data["部门目标"].rows} == {"D01", "D02"}
    assert len(data["个人目标"].rows) == 8
    assert any(r.values["approved"] is False for r in data["绩效评分"].rows)


def test_scope_errors_for_project_fields():
    rs = load("quarter")
    bad = rs.model_copy(deep=True)
    bad.variables["d_eff"].expr = "q_eff"                      # 部门层直接用项目变量
    codes = {i.code for i in validate(bad) if i.severity == "blocking"}
    assert "scope_aux" in codes
    bad = rs.model_copy(deep=True)
    bad.variables["q_sales"].expr = "sum(share)"               # 项目层聚合参与表可以
    assert not [i for i in validate(bad) if i.severity == "blocking"]
    bad.variables["q_sales"].expr = 'sum(sales_target, of: "个人目标")'   # 个人目标和项目没有关联
    assert "agg_unrelated" in {i.code for i in validate(bad)}


def test_date_functions():
    from bonus_core.dsl.parser import parse
    from bonus_core.engine.evaluator import Ctx, Evaluator
    from bonus_core.engine.tree import Node
    rs = load("quarter")
    root = Node("全部", "全部", None)
    ev = Evaluator(rs, root, {})
    val = lambda s: ev.ev(parse(s), Ctx(root)).v  # noqa: E731
    assert val('years_between("2017-06-30", "2026-06-29")') == 8
    assert val('years_between("2017-06-30", "2026-06-30")') == 9
    assert val('between(date("2026-09-30"), "2026-07-01", "2026-09-30")') is True
    assert val('quarter("2026-08-15") == 3 and month("2026-08-15") == 8') is True
    assert val('days_between("2026-01-01", "2026-03-01")') == 59


def _agent_script():
    """按 datasets.yaml / mapping.yaml / 规则文件生成一段智能体操作：共用部分 + 两个方案 + 修订。"""
    y = lambda n: yaml.safe_load((CASE / n).read_text(encoding="utf-8"))  # noqa: E731
    dsets, maps, shared = y("datasets.yaml"), y("mapping.yaml"), y("weights.yaml")
    script = []
    for name, d in dsets.items():
        m = maps[name]
        fields = {f: {"column": m["columns"][f], **fd} for f, fd in d["fields"].items()}
        args = {"name": name, "file": m["file"], "sheet": m["sheet"], "fields": fields,
                **{k: d[k] for k in ("key", "link") if k in d}}
        if "filter" in m:
            args["filter"] = m["filter"]
        script.append({"tool": "define_dataset", "args": args})
    script.append({"tool": "set_rules", "args": {"scheme": "共用", **shared}})
    for scheme in ("quarter", "process"):
        own = y(f"rules_{scheme}.yaml")
        args = {"scheme": scheme.upper(), **{k: own[k] for k in ("tables", "variables", "pool") if k in own}}
        if scheme == "quarter":
            args["revisions"] = [{"name": "R1", "adjust": [{"node": "小组:G01", "amount": 300000}]},
                                 {"name": "R2", "adjust": [{"node": "个人:E001", "amount": 150000}]}]
        script.append({"tool": "set_rules", "args": args})
    script += [{"tool": "trial_run"}, {"tool": "explain", "args": {"id": "项目:P01/q_eff", "scheme": "QUARTER"}},
               {"tool": "evaluate", "args": {"expr": "q_sales", "node": "项目:P01", "scheme": "QUARTER"}},
               {"tool": "finish", "args": {"summary": "季度奖和过程奖"}}]
    return script


def test_agent_multi_scheme_with_revisions(tmp_path):
    from bonus_agent.agent import run_agent
    from bonus_agent.llm import FakeLLM
    ar = run_agent(FakeLLM(_agent_script()), [XLSX], tmp_path)
    assert ar.ok, (ar.error, [s.result for s in ar.steps if not s.ok])
    assert set(ar.results) == {"QUARTER", "PROCESS"} and list(ar.results["QUARTER"]) == ["R0", "R1", "R2"]
    for scheme in ("QUARTER", "PROCESS"):
        r0 = ar.results[scheme]["R0"]
        assert {k: r0.nodes[k].final for k in EXPECTED[scheme]} == EXPECTED[scheme]
    r2 = ar.results["QUARTER"]["R2"]
    assert r2.nodes["小组:G01"].final == "300000.00" and r2.nodes["个人:E001"].final == "150000.00"
    assert r2.nodes["部门:D01"].final == EXPECTED["QUARTER"]["部门:D01"]
    assert ar.recheck.count("一致") == 4 and "不一致" not in ar.recheck, ar.recheck
    names = {p.name for p in tmp_path.iterdir()}
    assert {"result.xlsx", "result_QUARTER.xlsx", "result_QUARTER_R2.xlsx", "result_PROCESS.xlsx",
            "rules_QUARTER.yaml", "adjust_QUARTER_R2.yaml", "mapping.yaml"} <= names
    wb = openpyxl.load_workbook(tmp_path / "result.xlsx")
    assert {"各方案结果", "方案汇总", "待确认问题", "智能体规划"} <= set(wb.sheetnames)
    assert [c.value for c in wb["各方案结果"][1]][3:] == ["QUARTER", "QUARTER R1", "QUARTER R2", "PROCESS"]
    ev = next(s for s in ar.steps if s.tool == "evaluate")
    assert ev.ok and "项目:P01/q_sales" not in ev.result


def test_compare_with_expected_workbook(tmp_path):
    from bonus_agent.agent import run_agent
    from bonus_agent.compare import compare_expected, report
    from bonus_agent.llm import FakeLLM
    # 期望结果：一张长表，方案列 + 编号列 + 金额列；故意改错一个人
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "期望结果"
    ws.append(["scheme", "level", "entity_id", "expected_bonus"])
    for scheme in ("QUARTER", "PROCESS"):
        for k, v in EXPECTED[scheme].items():
            amount = float(v) + (1 if (scheme, k) == ("PROCESS", "个人:E003") else 0)
            ws.append([scheme, k.split(":")[0], k.split(":")[1], amount])
    exp = tmp_path / "期望.xlsx"
    wb.save(exp)
    inputs = tmp_path / "in"
    inputs.mkdir()
    (inputs / "业务输入.xlsx").write_bytes(XLSX.read_bytes())
    (inputs / "期望.xlsx").write_bytes(exp.read_bytes())
    ar = run_agent(FakeLLM(_agent_script()), [inputs], tmp_path / "out", exclude=[inputs / "期望.xlsx"])
    assert ar.ok and all(f.name != "期望.xlsx" for f in ar.ws.inv.files)      # 期望结果不给智能体
    ms = {(m.scheme, m.revision): m for m in compare_expected(ar.results, exp)}
    assert ms[("QUARTER", "R0")].ok and ms[("QUARTER", "R0")].column == "expected_bonus"
    p = ms[("PROCESS", "R0")]
    assert p.matched == p.total - 1 and p.diffs[0][0] == "E003"
    text = "\n".join(report(ms.values()))
    assert "E003 期望" in text and "14/14" in text


def test_global_variable_used_per_person_is_flagged():
    rs = load("quarter")
    assert "global_in_level" not in {i.code for i in validate(rs)}
    bad = rs.model_copy(deep=True)
    bad.variables["contrib"].scope = "global"                  # 漏写 scope：变成全公司评分合计
    bad.variables["contrib"].expr = 'sum(contribution_score, of: "绩效评分", where: period == "Q3" and approved)'
    w = [i for i in validate(bad) if i.code == "global_in_level"]
    assert w and w[0].severity == "warning" and "scope: 个人" in w[0].message


def test_engine_error_hints():
    rs = load("quarter")
    data, _ = apply_mapping(rs, [XLSX], load_mapping(CASE / "mapping.yaml"))
    bad = rs.model_copy(deep=True)
    bad.variables["contrib"].expr = 'sum(contribution_score, where: period)'     # 文字字段当条件
    with pytest.raises(EngineError, match='period == "Q3"'):
        compute(bad, data)


def test_offset_rolled_up_to_person_across_projects():
    """一人参与多个项目：按参与记录分配，再按工号汇总可发上限，在人员上抵扣预发。"""
    from bonus_core.data.table import from_records
    rs = RuleSet.model_validate(yaml.safe_load('''
name: 交付结算
version: v1
datasets:
  项目: {key: project_id, fields: {project_id: {type: id}, budget: {type: money}, visit_ok: {type: bool}}}
  参与: {key: part_id, fields: {part_id: {type: id}, project_id: {type: id}, emp_id: {type: id}, share: {type: ratio}}}
  人员: {key: emp_id, fields: {emp_id: {type: id}, emp_name: {type: text}, employ_type: {type: text}}}
  预发: {link: emp_id, fields: {emp_id: {type: id}, amount: {type: money}}}
pool: {scope: global, expr: "10000"}
allocation:
  - {level: 项目, key: project_id, weight: budget}
  - {level: 参与记录, key: part_id, parent_key: project_id, weight: 'if(employ_type == "正式", share, 0)'}
payout:
  - {id: 回访, scope: 参与记录, pay_ratio: 'if(visit_ok, 1, 0.8)'}
offsets:
  - {id: 预发, scope: 人员, expr: 'sum(amount, of: "预发")'}
'''))
    assert not [i for i in validate(rs) if i.severity == "blocking"], validate(rs)
    data = from_records(rs, {
        "项目": [{"project_id": "J1", "budget": 3, "visit_ok": True}, {"project_id": "J2", "budget": 1, "visit_ok": False}],
        "参与": [{"part_id": "A1", "project_id": "J1", "emp_id": "U1", "share": 0.5},
                 {"part_id": "A2", "project_id": "J1", "emp_id": "U2", "share": 0.5},
                 {"part_id": "A3", "project_id": "J2", "emp_id": "U1", "share": 1},
                 {"part_id": "A4", "project_id": "J2", "emp_id": "U3", "share": 1}],     # 实习：权重 0
        "人员": [{"emp_id": "U1", "emp_name": "甲", "employ_type": "正式"}, {"emp_id": "U2", "emp_name": "乙", "employ_type": "正式"},
                 {"emp_id": "U3", "emp_name": "丙", "employ_type": "实习"}],
        "预发": [{"emp_id": "U1", "amount": 1000}, {"emp_id": "U2", "amount": 4000}],
    })
    res = compute(rs, data)
    u1, u2 = res.rollups["人员:U1"], res.rollups["人员:U2"]
    # U1：J1 一半 3750 全发 + J2 2500 发 80% = 2000 → 上限 5750，扣 1000
    assert (u1.final, u1.held, u1.cap, u1.deducted, u1.paid, u1.overpaid) == ("6250.00", "500.00", "5750.00", "1000.00", "4750.00", "0")
    # U2：上限 3750，预发 4000 → 实发 0，超发 250
    assert (u2.cap, u2.deducted, u2.paid, u2.overpaid) == ("3750.00", "3750.00", "0.00", "250.00")
    assert sorted(u1.members) == ["参与记录:A1", "参与记录:A3"] and res.nodes["参与记录:A4"].final == "0.00"
    assert res.traces["参与记录:A4/权重"].rendered.startswith("if(employ_type[U3]")
    from bonus_agent.compare import _figures
    f = _figures(res)
    assert f["本次可发"] == {"U1": Decimal("4750.00"), "U2": Decimal("0.00"), "U3": Decimal("0.00")} and f["应得"]["U1"] == Decimal("6250.00")
    # 最末层直接写个人层、但个人表没有 project_id：提示改用参与记录 + 按人汇总
    bad = rs.model_copy(deep=True)
    bad.allocation[1].key, bad.allocation[1].level = "emp_id", "个人"
    msgs = [i.message for i in validate(bad) if i.code == "parent_key_missing"]
    assert msgs and "offsets 的 scope" in msgs[0]


def test_data_patches_append_and_correct():
    """数据修订：补充表并入原表（按单号去重、取值对照）；长表 / 宽表更正（核对旧值）；有歧义直接报错。"""
    from bonus_core.data.table import from_records
    base = yaml.safe_load('''
name: 修订
version: v2
datasets:
  项目: {key: project_id, fields: {project_id: {type: id}, grade: {type: text}, budget: {type: money}}}
  参与: {key: part_id, fields: {part_id: {type: id}, project_id: {type: id}, score: {type: number}}}
  收款: {link: project_id, fields: {receipt_id: {type: id}, project_id: {type: id}, amount: {type: money}, approved: {type: text}}}
  补充收款: {link: project_id, fields: {receipt_id: {type: id}, project_id: {type: id}, amount: {type: money}, approved: {type: text}}}
  项目更正: {fields: {project_id: {type: id}, field_name: {type: text}, new_value: {type: text}}}
  评分修订: {fields: {part_id: {type: id}, old_score: {type: number}, new_score: {type: number}}}
pool: {scope: global, expr: 'sum(amount, of: "收款", where: approved == "是")'}
allocation:
  - {level: 项目, key: project_id, weight: 'if(grade == "优", 2, 1)'}
  - {level: 参与记录, key: part_id, parent_key: project_id, weight: score}
patches:
  - {kind: append, source: 补充收款, into: 收款, unique: receipt_id, values: {approved: {通过: 是, 驳回: 否}}}
  - {kind: correct, source: 项目更正, target: 项目, key: project_id, field: field_name, value: new_value, labels: {交付评级: grade}}
  - {kind: correct, source: 评分修订, target: 参与, key: part_id, set: {score: new_score}, expect: {score: old_score}}
''')
    recs = {
        "项目": [{"project_id": "J1", "grade": "合格", "budget": 1}, {"project_id": "J2", "grade": "合格", "budget": 1}],
        "参与": [{"part_id": "A1", "project_id": "J1", "score": 80}, {"part_id": "A2", "project_id": "J1", "score": 20},
                 {"part_id": "A3", "project_id": "J2", "score": 50}],
        "收款": [{"receipt_id": "R1", "project_id": "J1", "amount": 1000, "approved": "是"}],
        "补充收款": [{"receipt_id": "R1", "project_id": "J1", "amount": 1000, "approved": "通过"},     # 重传：跳过
                     {"receipt_id": "R2", "project_id": "J2", "amount": 500, "approved": "通过"},
                     {"receipt_id": "R3", "project_id": "J2", "amount": 700, "approved": "驳回"}],
        "项目更正": [{"project_id": "J1", "field_name": "交付评级", "new_value": "优"},
                     {"project_id": "J2", "field_name": "交付评级", "new_value": "合格"}],                # 同值：不算变化
        "评分修订": [{"part_id": "A2", "old_score": 20, "new_score": 30}],
    }
    rs = RuleSet.model_validate(base)
    assert not [i for i in validate(rs) if i.severity == "blocking"], validate(rs)
    res = compute(rs, from_records(rs, recs))
    assert res.pool == "1500.00"                                      # 1000 + 500，驳回的 700 不计，R1 不重复
    assert res.nodes["项目:J1"].final == "1000.00"                    # J1 改成“优”权重 2
    assert res.nodes["参与记录:A1"].final == "727.27" and res.nodes["参与记录:A2"].final == "272.73"   # 80 : 30
    notes = [i.message for i in res.issues if i.code == "data_patch"]
    assert "2 行并入「收款」，1 行重复已跳过" in notes[0], notes
    assert any("J1 grade：合格 → 优" in n for n in notes) and any("A2 score：20 → 30" in n for n in notes)
    # 没有修订的方案（V1）不受影响
    v1 = RuleSet.model_validate({**base, "patches": []})
    assert compute(v1, from_records(v1, recs)).pool == "1000.00"

    def fails(**change):
        r = {**recs, **change}
        with pytest.raises(EngineError) as ei:
            compute(rs, from_records(rs, r))
        return str(ei.value.issues[0].message)
    assert "内容不同" in fails(补充收款=[{"receipt_id": "R1", "project_id": "J1", "amount": 999, "approved": "通过"}])
    assert "对不上" in fails(评分修订=[{"part_id": "A2", "old_score": 25, "new_score": 30}])
    assert "重复修订" in fails(评分修订=[{"part_id": "A2", "old_score": 20, "new_score": 30}] * 2)
    assert "找不到" in fails(评分修订=[{"part_id": "A9", "old_score": 20, "new_score": 30}])
    assert "新值不同" in fails(项目更正=[{"project_id": "J1", "field_name": "交付评级", "new_value": "优"},
                                      {"project_id": "J1", "field_name": "交付评级", "new_value": "合格"}])
    assert "没有字段" in fails(项目更正=[{"project_id": "J1", "field_name": "项目经理", "new_value": "x"}])
