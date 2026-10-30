"""冲突自动检测：同一参数在不同材料里取值不一致时，由程序判定并提问，不交给模型取舍。"""
import json
from decimal import Decimal

import docx
import openpyxl
import yaml

from bonus_agent import facts as fx
import pytest

from bonus_agent.agent import ToolError, Workspace, run_agent
from bonus_agent.auto import build_inventory
from bonus_agent.llm import FakeLLM

NAME = "季度奖-计提比例"


def _doc(path, *paras):
    d = docx.Document()
    for t in paras:
        d.add_paragraph(t)
    d.save(path)
    return path


def _materials(d):
    d.mkdir(exist_ok=True)
    _doc(d / "方案.docx", "季度奖方案", "季度奖按部门销售额的 5% 计提奖金池，按个人销售额占比分配。")
    _doc(d / "补充说明.docx", "补充说明", "季度奖计提比例为 6%。")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "人员"
    ws.append(["员工编号", "销售额"])
    for e, s in (("E1", 100000), ("E2", 300000)):
        ws.append([e, s])
    wb.save(d / "数据.xlsx")
    return d


def fact(value, file, quote, **kw):
    return {"name": NAME, "value": value, "file": file, "quote": quote, **kw}


F5 = fact("5%", "方案.docx", "季度奖按部门销售额的 5% 计提奖金池")
F6 = fact("6%", "补充说明.docx", "季度奖计提比例为 6%。")
DATASET = {"tool": "define_dataset", "args": {"name": "人员", "file": "数据.xlsx", "sheet": "人员", "key": "eid", "fields": {
    "eid": {"column": "员工编号", "type": "id"}, "sales": {"column": "销售额", "type": "money", "unit": "元"}}}}


CLAUSES = [{"tool": "record_clauses", "args": {"clauses": [
    {"point": "奖金池按部门销售额 5% 计提", "file": "方案.docx", "quote": F5["quote"]},
    {"point": "计提比例 6%", "file": "补充说明.docx", "quote": F6["quote"]}]}},
    {"tool": "cover_clauses", "args": {"items": {"C1": "引擎自带：测试", "C2": "引擎自带：测试"}}}]


def rules(rate):
    return {"tool": "set_rules", "args": {"pool": {"expr": f"sum(sales) * {rate}"},
                                          "allocation": [{"level": "个人", "key": "eid", "weight": "sales"}]}}


def test_number_and_quote_checks(tmp_path):
    assert fx.number("5%") == Decimal("0.05") and fx.number("30万元") == 300000 and fx.number("1,200") == 1200
    assert fx.number("含税") is None and fx.same_value(fx.Fact(**F5)) == fx.same_value(fx.Fact(**{**F5, "value": "5 %"}))
    ws = Workspace(FakeLLM([]), build_inventory(sorted(_materials(tmp_path / "m").iterdir())))
    out = ws.record_facts([fact("5%", "方案.docx", "按销售额的 5% 计提"),                 # 原文没有这句话
                           fact("8%", "方案.docx", "季度奖按部门销售额的 5% 计提奖金池"),   # 取值不在引用里
                           fact("6%", "补充说明.docx", "季度奖计提比例为 6%。", override=True),   # 没写明修改
                           fact("5%", "不存在.docx", "x"), F5])
    assert "已登记 1 条" in out and out.count("未登记") == 4 and "找不到" in out and "不在 quote 里" in out
    assert "override" in out and "没有文件" in out
    assert ws.record_facts([fact("100000", "数据.xlsx", "E1 | 100000")]).startswith("已登记 1 条")   # Excel 单元格也能引用


def test_conflict_asked_by_program(tmp_path):
    ws = Workspace(FakeLLM([]), build_inventory(sorted(_materials(tmp_path / "m").iterdir())))
    ws.record_facts([F5])
    out = ws.record_facts([F6])
    assert "自动向财务提问" in out and len(ws.pending) == 1
    q = ws.pending[0]
    assert q.related == f"参数:{NAME}" and q.blocking and q.options[:2] == ["5%（方案.docx）", "6%（补充说明.docx）"]
    assert "补充说明.docx“季度奖计提比例为 6%。”" in q.text
    with pytest.raises(ToolError, match="之前已登记过"):         # 再登记同样的取值：不重复提问，作为失败返回
        ws.record_facts([F5, F6, {**F5, "value": "5 %"}])
    assert len(ws.questions) == 1 and len(ws.facts) == 2
    # 有冲突未答复：可以结束本轮（不出结果）
    assert "等待财务答复" in ws.finish("x")


def test_explicit_override_is_not_a_conflict(tmp_path):
    m = _materials(tmp_path / "m")
    _doc(m / "补充说明.docx", "补充说明", "季度奖计提比例修改为 6%，以本通知为准。")
    ws = Workspace(FakeLLM([]), build_inventory(sorted(m.iterdir())))
    ws.record_facts([F5, fact("6%", "补充说明.docx", "季度奖计提比例修改为 6%，以本通知为准", override=True)])
    assert not ws.questions and fx.chosen(NAME, ws.conflicts, ws.questions, ws.facts) == (True, "0.06")
    # 第三份材料又写了别的值：再明确的修改也要问
    _doc(m / "通知.docx", "季度奖计提比例 7%。")
    ws = Workspace(FakeLLM([]), build_inventory(sorted(m.iterdir())))
    ws.record_facts([F5, fact("6%", "补充说明.docx", "季度奖计提比例修改为 6%，以本通知为准", override=True),
                     fact("7%", "通知.docx", "季度奖计提比例 7%")])
    assert len(ws.pending) == 1 and len(ws.pending[0].options) == 4


def test_finish_requires_facts_and_uses_them(tmp_path):
    ws = Workspace(FakeLLM([]), build_inventory(sorted(_materials(tmp_path / "m").iterdir())))
    ws.define_dataset(**DATASET["args"])
    ws.set_rules(**rules("4%")["args"])
    ws.trial_run()
    try:
        ws.finish("x")
        raise AssertionError("没登记参数也能完成")
    except Exception as e:
        assert "方案.docx" in str(e) and "补充说明.docx" in str(e)
    ws.record_facts([F5])
    with pytest.raises(ToolError, match="方案.docx」还没有摘出条款"):
        ws.finish("x", ignored=[{"ref": "补充说明.docx", "reason": "只是重复"}])
    ws.record_clauses([{"point": "按销售额 5% 计提", "file": "方案.docx", "quote": F5["quote"]}])
    with pytest.raises(ToolError, match="C1「按销售额 5% 计提」：没有写落在哪里"):
        ws.finish("x")
    with pytest.raises(ToolError, match="variables.rate 不存在"):
        ws.cover_clauses({"C1": "variables.rate"})
    ws.cover_clauses({"C1": "pool"})
    assert "C1 [方案.docx] 按销售额 5% 计提 → pool" in ws.state()
    try:
        ws.finish("x")
        raise AssertionError("登记的 5% 没用到也能完成")
    except ToolError as e:
        assert f"{NAME}（5%）" in str(e)
    ws.set_rules(**rules("5%")["args"])
    ws.trial_run()
    with pytest.raises(ToolError) as ei:               # 核对过的规则又改了：并排列出原文和现在的规则，自查一次
        ws.finish("x")
    assert "结束前自查" in str(ei.value) and "原文（方案.docx）：季度奖按部门销售额的 5% 计提奖金池" in str(ei.value)
    assert "奖金池[global] = sum(sales) * 5%" in str(ei.value)
    assert ws.finish("x") == "已完成"          # 上一次 finish 写的忽略项仍然有效


def test_conflict_rounds(tmp_path):
    m, out = _materials(tmp_path / "m"), tmp_path / "案卷"
    # 第 1 轮：两份材料比例不同，程序自动提问，不出结果
    r1 = [{"tool": "record_facts", "args": {"facts": [F5, F6]}}, *CLAUSES, DATASET, {"tool": "finish", "args": {"summary": "季度奖"}}]
    ar = run_agent(FakeLLM(r1), [m], out)
    assert ar.status == "NEEDS_CONFIRMATION" and not (out / "result.xlsx").exists()
    q = yaml.safe_load((out / "questions.yaml").read_text(encoding="utf-8"))
    assert q[0]["related"] == f"参数:{NAME}"
    wb = openpyxl.load_workbook(out / "待确认问题.xlsx")
    assert "参数台账" in wb.sheetnames
    # 第 2 轮：财务选 6%，模型先写错成 5%，完成时程序拦下
    q[0]["answer"] = 1
    (out / "questions.yaml").write_text(yaml.safe_dump(q, allow_unicode=True), encoding="utf-8")
    llm = FakeLLM([rules("5%"), {"tool": "trial_run"}, {"tool": "finish", "args": {"summary": "季度奖"}},
                   rules("6%"), {"tool": "trial_run"}, {"tool": "finish", "args": {"summary": "季度奖"}}])
    ar = run_agent(llm, [m], out)
    assert ar.status == "OK", ar.error
    bad = next(s for s in ar.steps if s.tool == "finish" and not s.ok)
    assert "财务选定 0.06" in bad.result
    assert ar.result.pool == "24000.00"
    # 第 3 轮：新材料写了第三个比例，再次提问；选项里保留前面的取值
    _doc(m / "通知.docx", "季度奖计提比例调整 7%。")
    r3 = [{"tool": "record_facts", "args": {"facts": [fact("7%", "通知.docx", "季度奖计提比例调整 7%")]}},
          {"tool": "finish", "args": {"summary": "季度奖"}}]
    ar = run_agent(FakeLLM(r3), [m], out)
    assert ar.status == "NEEDS_CONFIRMATION" and [p["id"] for p in ar.pending] == ["q2"]
    assert ar.pending[0]["options"][:3] == ["5%（方案.docx）", "6%（补充说明.docx）", "7%（通知.docx）"]
    # 第 4 轮：通知改了内容（修改的文件，台账按新内容重新登记）
    _doc(m / "通知.docx", "季度奖计提比例修改为 6%，以本通知为准。")
    r4 = [{"tool": "record_facts", "args": {"facts": [
        fact("6%", "通知.docx", "季度奖计提比例修改为 6%，以本通知为准", override=True)]}},
        {"tool": "record_clauses", "args": {"clauses": [{"point": "改为 6%", "file": "通知.docx", "quote": "以本通知为准"}]}},
        {"tool": "cover_clauses", "args": {"items": {"C3": "pool"}}},
        {"tool": "trial_run"}, {"tool": "finish", "args": {"summary": "季度奖"}}]
    llm = FakeLLM(r4)
    ar = run_agent(llm, [m], out)
    assert "修改文件：通知.docx" in llm.calls[0][1]["content"]
    assert not any(f.value == "7%" for f in ar.ws.facts)
    # q2（5% / 6% / 7%）依据的 7% 已随文件修改作废：程序撤回；通知明确修改为 6%，不算冲突，出结果
    assert "程序撤回的问题" in llm.calls[0][1]["content"] and "q2" in llm.calls[0][1]["content"]
    assert ar.status == "OK" and ar.ws.questions[1].withdrawn, ar.error
    q = yaml.safe_load((out / "questions.yaml").read_text(encoding="utf-8"))
    assert q[1]["withdrawn"] and ar.result.pool == "24000.00"


def test_quote_miss_suggests_original_and_hints_next_step(tmp_path):
    ws = Workspace(FakeLLM([]), build_inventory(sorted(_materials(tmp_path / "m").iterdir())))
    with pytest.raises(ToolError) as ei:
        ws.record_facts([fact("5%", "方案.docx", "季度奖按销售额 5% 计提")])
    out = str(ei.value)
    assert "最接近的一句：季度奖按部门销售额的 5% 计提奖金池" in out
    ws.record_facts([F5])
    with pytest.raises(ToolError, match="之前已登记过"):
        ws.record_facts([F5])
    assert "define_dataset" in ws.next_hint() and "补充说明.docx" in ws.next_hint()
    ws.progress = (30, 80)
    assert "下一步（程序判断）" in ws.state()


def test_chat_answers_and_changes(tmp_path):
    m, out = _materials(tmp_path / "m"), tmp_path / "案卷"
    r1 = [{"tool": "record_facts", "args": {"facts": [F5, F6]}}, *CLAUSES, DATASET, {"tool": "finish", "args": {"summary": "季度奖"}}]
    assert run_agent(FakeLLM(r1), [m], out).status == "NEEDS_CONFIRMATION"
    # 对话里财务明确答复了问题：记下原话、按答复改规则，程序自动试算并检查，生成第 2 轮
    llm = FakeLLM([{"tool": "answer_question", "args": {"id": "q1", "answer": "以补充说明为准，6%"}}, rules("6%"),
                   {"tool": "reply", "args": {"text": "已按 6% 计提"}}])
    ar = run_agent(llm, [m], out, message="计提比例以补充说明为准，6%", by="张会计")
    assert "财务的消息：计提比例以补充说明为准，6%" in llm.calls[0][1]["content"]
    assert ar.status == "OK" and ar.round == 2 and ar.reply == "已按 6% 计提" and ar.result.pool == "24000.00"
    q = yaml.safe_load((out / "questions.yaml").read_text(encoding="utf-8"))[0]
    assert q["answer"] == "以补充说明为准，6%" and q["answered_by"] == "张会计（对话）"
    case = json.loads((out / "case.json").read_text(encoding="utf-8"))
    assert case["rounds"][-1]["message"] == {"by": "张会计", "text": "计提比例以补充说明为准，6%"}
    # 只是提问：可以查算式，不生成新的一轮
    llm = FakeLLM([{"tool": "explain", "args": {"id": "个人:E1"}}, {"tool": "reply", "args": {"text": "E1 按销售额占比分得 6000"}}])
    ar = run_agent(llm, [m], out, message="E1 怎么算的？", by="张会计", chat=[{"role": "user", "text": "之前的话"}])
    assert ar.reply.startswith("E1") and ar.steps[0].ok and "最终金额 6,000" in ar.steps[0].result
    assert len(json.loads((out / "case.json").read_text(encoding="utf-8"))["rounds"]) == 2
    assert "之前的话" in llm.calls[0][1]["content"]
    # 改了规则但试算 / 检查不通过：reply 被拦下，模型修正后才结束
    llm = FakeLLM([rules("4%"), {"tool": "reply", "args": {"text": "改好了"}},
                   rules("6%"), {"tool": "reply", "args": {"text": "改回 6%"}}])
    ar = run_agent(llm, [m], out, message="试试 4%", by="张会计")
    bad = next(s for s in ar.steps if s.tool == "reply" and not s.ok)
    assert "财务选定 0.06" in bad.result and ar.reply == "改回 6%" and ar.round == 3
    # 不是对话时不能用 reply / answer_question
    ws = Workspace(FakeLLM([]), build_inventory(sorted(m.iterdir())))
    for call in (lambda: ws.reply("x"), lambda: ws.answer_question("q1", "x")):
        try:
            call()
            raise AssertionError("非对话也能用")
        except Exception as e:
            assert "对话" in str(e)


def test_stops_after_three_identical_failures(tmp_path):
    m = _materials(tmp_path / "m")
    bad = {"tool": "record_facts", "args": {"facts": [fact("5%", "方案.docx", "原文没有这句 5%")]}}
    ar = run_agent(FakeLLM([bad] * 20), [m])
    assert "卡住" in ar.error and len(ar.steps) == 3 and "record_facts" in ar.error


def test_tiers_under_one_name_are_not_a_conflict(tmp_path):
    ws = Workspace(FakeLLM([]), build_inventory(sorted(_materials(tmp_path / "m").iterdir())))
    out = ws.record_facts([F5, {**F6, "file": "方案.docx", "quote": F5["quote"], "value": "5%"},
                           {**F5, "quote": "季度奖按部门销售额的 5% 计提奖金池", "value": "5%"}])
    assert not ws.questions
    tier = [fact("5%", "方案.docx", "季度奖按部门销售额的 5% 计提奖金池"), fact("6%", "补充说明.docx", "季度奖计提比例为 6%。")]
    ws2 = Workspace(FakeLLM([]), build_inventory(sorted(_materials(tmp_path / "n").iterdir())))
    ws2.record_facts(tier)
    assert len(ws2.pending) == 1                           # 两个文件各给一个不同的值：仍是冲突
    assert "撤回问题 q1" in ws2.remove_facts(tier[0]["name"]) and not ws2.pending


def test_define_remove_oscillation_stops(tmp_path):
    """建了又删、删了又建（每次都成功，不算失败也不算原样重复）：同一张表删到第 3 次就停下，不再耗 tokens。"""
    from bonus_agent.agent import run_agent
    from bonus_agent.failure import explain
    from bonus_agent.llm import FakeLLM
    rm = {"tool": "remove_dataset", "args": {"name": "人员"}}
    ar = run_agent(FakeLLM([DATASET, rm] * 4 + [{"tool": "finish", "args": {"summary": "x"}}]),
                   [_materials(tmp_path / "m")], tmp_path / "out")
    assert "反复建了又删" in ar.error and len(ar.steps) == 6, ar.error
    assert "删了 2 次" in ar.steps[3].result
    assert "「人员」" in explain(ar.error)["title"]


def test_date_spellings_are_not_conflicts():
    """同一个日期的不同写法（Word 里的“2026 年 7 月 1 日”、Excel 里的日期时间、没写年份的“9 月 30 日”）不算冲突。"""
    mk = lambda v, f: fx.Fact(name="季度奖统计期间-结束", value=v, file=f, quote=v)    # noqa: E731
    fs = [mk("2026 年 9 月 30 日", "a.docx"), mk("2026-09-30 00:00:00", "b.xlsx"), mk("9 月 30 日", "c.pptx")]
    assert fx.detect(fs, {}, [], 1) == []
    fs.append(mk("2026-10-31", "d.xlsx"))
    q = fx.detect(fs, {}, [], 1)
    assert len(q) == 1 and len(q[0].options) == 3      # 两个不同日期 + 其他


def test_saving_rules_with_same_blocking_issue_stops(tmp_path):
    """set_rules 每次都“保存成功”，但同一个阻断问题一直没改掉：第 3 次停下（内容不同，不算原样重复）。"""
    from bonus_agent.agent import run_agent
    from bonus_agent.failure import explain
    from bonus_agent.llm import FakeLLM

    def bad(i):
        return {"tool": "set_rules", "args": {"pool": {"expr": f"sum(sales) * 0.0{i}"},
                                              "allocation": [{"level": "个人", "key": "eid", "weight": "no_such_var"}]}}
    ar = run_agent(FakeLLM([DATASET, bad(1), bad(2), bad(3), bad(4), {"tool": "finish", "args": {"summary": "x"}}]),
                   [_materials(tmp_path / "m")], tmp_path / "out")
    assert "同样的阻断问题" in ar.error and len(ar.steps) == 4, ar.error
    assert "连续几次没改掉" in ar.steps[2].result
    assert "同一个问题一直没改掉" in explain(ar.error)["title"]


def test_blocking_issues_shrinking_is_progress(tmp_path):
    """每次保存都改掉了几个阻断问题、只有一个一直挂着：在推进，不算卡住（同一问题挂满两倍次数才停）。"""
    from bonus_agent.agent import run_agent
    from bonus_agent.llm import FakeLLM

    def save(*bad):
        return {"tool": "set_rules", "args": {"pool": {"expr": "sum(sales) * 0.05"},
                                              "variables": {v: {"expr": f"{v}_undefined + 1" if v in bad else "1"}
                                                            for v in "abcdef"},
                                              "allocation": [{"level": "个人", "key": "eid", "weight": "no_such_var"}]}}
    script = [DATASET, save("a", "b", "c"), save("a", "b"), save("a"), save(), {"tool": "finish", "args": {"summary": "x"}}]
    ar = run_agent(FakeLLM(script), [_materials(tmp_path / "m")], tmp_path / "out")
    assert "同样的阻断问题" not in ar.error and len(ar.steps) >= 5, ar.error
    assert all("需要修正" in s.result for s in ar.steps[1:5])          # 每次都还有阻断（权重一直没改），但没被停下
    ar = run_agent(FakeLLM([DATASET] + [save(*"abcdef"[:6 - i]) for i in range(7)]
                           + [{"tool": "finish", "args": {"summary": "x"}}]), [_materials(tmp_path / "m2")], tmp_path / "out2")
    assert "同样的阻断问题" in ar.error and len(ar.steps) == 7, ar.error     # 权重问题连挂 6 次：停下


def test_manual_stop_keeps_steps(tmp_path):
    """页面上点「停止」：下一步开始前停下，已做的步骤照常保存。"""
    from bonus_agent.agent import run_agent
    from bonus_agent.failure import explain
    from bonus_agent.llm import FakeLLM
    calls = []
    ar = run_agent(FakeLLM([DATASET, rules(0.1), {"tool": "finish", "args": {"summary": "x"}}]),
                   [_materials(tmp_path / "m")], tmp_path / "out", should_stop=lambda: bool(calls.append(1)) or len(calls) > 1)
    assert ar.error.startswith("已手动停止") and len(ar.steps) == 1
    assert "手动停止" in explain(ar.error)["title"] and (tmp_path / "out" / "agent_log.json").exists()


def test_blocking_issues_come_with_fixes(tmp_path):
    """阻断问题附上改法；改法表里的每一类都是程序真的会报的问题。"""
    import re as _re
    from pathlib import Path as _P
    from bonus_agent.fixes import FIXES
    src = "".join(p.read_text(encoding="utf-8") for p in _P("bonus_core").rglob("*.py"))
    codes = set(_re.findall(r'(?:block|_block|warn)\("([a-z_]+)"', src)) | {"check_failed", "eval_error"}
    assert set(FIXES) <= codes, set(FIXES) - codes
    ws = Workspace(FakeLLM([]), build_inventory(sorted(_materials(tmp_path / "m").iterdir())))
    ws.define_dataset(**DATASET["args"])
    out = ws.set_rules(pool={"expr": "sum(sales) * 5%"}, allocation=[{"level": "个人", "key": "eid", "weight": "salse"}])
    assert "name_unknown" in out and "改法：" in out and "不是 Excel 列名" in out
