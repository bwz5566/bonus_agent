"""测试反馈的 7 个问题的回归测试。"""
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest

from bonus_core.data.mapping import DatasetMapping, apply_mapping
from bonus_core.data.table import from_records
from bonus_core.dsl.models import RuleSet
from bonus_core.dsl.validator import validate
from bonus_core.engine.issues import EngineError
from bonus_core.engine.ledger import ledger_entries, release
from bonus_core.engine.run import Adjustment, Engine, compute

ROOT = Path(__file__).parent.parent


def simple(pool="100", payout=None, people=(("E1", 1), ("E2", 1))):
    rs = RuleSet.model_validate({
        "name": "t", "version": "v1",
        "datasets": {"人员": {"key": "eid", "fields": {"eid": {"type": "id"}, "score": {"type": "number"},
                                                       "ratio": {"type": "number"}}}},
        "pool": {"expr": pool},
        "allocation": [{"level": "个人", "key": "eid", "weight": "score"}],
        "payout": payout or [],
    })
    rows = [{"eid": e, "score": s, "ratio": 0.5} for e, s in people]
    return rs, from_records(rs, {"人员": rows})


# 1. 指定的文件 / 工作表不存在时报错，不能静默改读别的表
def test_missing_sheet_does_not_fall_back(tmp_path):
    rs, _ = simple()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "旧数据"
    ws.append(["编号", "得分", "比例"])
    ws.append(["E1", 1, 0.5])
    f = tmp_path / "数据.xlsx"
    wb.save(f)
    cols = {"eid": "编号", "score": "得分", "ratio": "比例"}
    for m in (DatasetMapping(file="数据.xlsx", sheet="本期数据", columns=cols),
              DatasetMapping(file="不存在.xlsx", sheet="本期数据", columns=cols),
              DatasetMapping(file="不存在.xlsx", sheet="旧数据", columns=cols)):
        data, issues = apply_mapping(rs, [f], {"人员": m})
        assert "人员" not in data and any(i.code in ("sheet_missing", "file_missing") for i in issues), m
    data, issues = apply_mapping(rs, [f], {"人员": DatasetMapping(file="数据.xlsx", sheet="旧数据", columns=cols)})
    assert len(data["人员"].rows) == 1


# 2. 挂起条件编号重复会互相覆盖台账，静态校验要拦下
def test_duplicate_payout_ids_blocked():
    rs, data = simple(payout=[{"id": "h", "scope": "个人", "hold_ratio": "0.5"},
                              {"id": "h", "scope": "个人", "hold_ratio": "0.5"}])
    assert any(i.code == "id_duplicate" for i in validate(rs))
    with pytest.raises(EngineError, match="id_duplicate"):
        compute(rs, data)
    rs, data = simple(payout=[{"id": "h1", "scope": "个人", "hold_ratio": "0.5"},
                              {"id": "h2", "scope": "个人", "hold_ratio": "0.5"}])
    res = compute(rs, data)
    held = sum(Decimal(e["held"]) for e in ledger_entries(res) if e["node_id"] == "个人:E1")
    assert held == Decimal(res.payouts["个人:E1"].held) == Decimal("50.00")      # 台账和挂起金额一致


# 3. 期望结果核对：不能因为旁边有一列“正好一致”的数字就判通过
def test_expected_compare_no_false_pass(tmp_path):
    from bonus_agent.compare import compare_expected, report
    rs, data = simple(people=[(f"E{i}", i) for i in range(1, 10)])
    res = compute(rs, data)
    ours = {n.key: n.final for n in res.nodes.values() if n.key}
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["员工编号", "期望奖金", "参考列"])
    for k, v in ours.items():
        ws.append([k, float(v) + 1, float(v)])          # 期望列全部改错，另一列与结果相同
    p = tmp_path / "期望.xlsx"
    wb.save(p)
    m = compare_expected({"S": {"R0": res}}, p)[0]
    assert not m.ok and m.column == "期望奖金" and m.matched == 0 and m.total == 9
    # 两列都认不出是哪一列：判为无法确定，不通过
    ws.cell(1, 2).value, ws.cell(1, 3).value = "A", "B"
    wb.save(p)
    m = compare_expected({"S": {"R0": res}}, p)[0]
    assert not m.ok and m.ambiguous and "无法确定" in report([m])[0]
    m = compare_expected({"S": {"R0": res}}, p, column="B")[0]   # 指定列名
    assert m.ok and m.total == 9


# 4. 补发时比例同样要在 0~1 之间，剩余挂起不能为负
def test_release_checks_ratio_range():
    rs, data = simple(payout=[{"id": "h", "scope": "个人", "hold_ratio": "ratio"}])
    entries = ledger_entries(compute(rs, data))
    bad = from_records(rs, {"人员": [{"eid": "E1", "score": 1, "ratio": -1}, {"eid": "E2", "score": 1, "ratio": 0.2}]})
    with pytest.raises(EngineError, match="hold_ratio_range"):
        release(rs, entries, bad)
    good = from_records(rs, {"人员": [{"eid": "E1", "score": 1, "ratio": 0}, {"eid": "E2", "score": 1, "ratio": 0.2}]})
    out = {e["node_id"]: e for e in release(rs, entries, good)}
    # 各应发 50、当初挂起 25；E1 比例降到 0 全部补发，E2 降到 0.2 仍挂起 10
    assert out["个人:E1"]["release_now"] == "25.00" and out["个人:E1"]["still_held"] == "0.00"
    assert out["个人:E2"]["release_now"] == "15.00" and out["个人:E2"]["still_held"] == "10.00"
    assert all(Decimal(e["still_held"]) >= 0 for e in out.values())


# 5. 输入哈希要包含人工调整
def test_input_hash_covers_adjustments():
    rs, data = simple(pool="11.12")
    eng = Engine(rs, data)
    a, b = eng.compute(), eng.compute([Adjustment("个人:E1", Decimal("1"))])
    assert a.nodes["个人:E1"].final == "5.56" and b.nodes["个人:E1"].final == "1.00"
    assert a.input_hash != b.input_hash and a.data_hash == b.data_hash
    c = eng.compute([Adjustment("个人:E1", Decimal("2"))])
    assert c.input_hash not in (a.input_hash, b.input_hash)


# 6. 安装包自带默认模型配置，在任意目录都能读到
def test_default_config_is_packaged(tmp_path, monkeypatch):
    from bonus_agent import llm
    assert llm.DEFAULT_CONFIG.exists() and llm.DEFAULT_CONFIG.is_relative_to(Path(llm.__file__).parent)
    monkeypatch.chdir(tmp_path)
    monkeypatch.delenv("BONUS_MODELS_CONFIG", raising=False)
    monkeypatch.setattr(llm, "REPO_CONFIG", tmp_path / "no-such-dir")
    assert llm.load_config()["models"]
    import tomllib
    pkg = tomllib.loads((ROOT / "pyproject.toml").read_text(encoding="utf-8"))["tool"]["setuptools"]["package-data"]
    assert "config/*.yaml" in pkg["bonus_agent"]


# 7. 奖金池精确值 0.005 显示为 0.01，把唯一的人锁定为 0.01 不能报超额
def test_sub_cent_pool_lock():
    rs, data = simple(pool="0.005", people=[("E1", 1)])
    eng = Engine(rs, data)
    assert eng.compute().nodes["个人:E1"].final == "0.01"
    assert eng.compute([Adjustment("个人:E1", Decimal("0.01"))]).nodes["个人:E1"].final == "0.01"
    rs, data = simple(pool="0.005")
    r = Engine(rs, data).compute([Adjustment("个人:E1", Decimal("0.01"))])
    assert r.nodes["个人:E1"].final == "0.01" and r.nodes["个人:E2"].final == "0.00"
