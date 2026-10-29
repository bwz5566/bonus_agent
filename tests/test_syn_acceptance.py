"""财务验收样例（evals/cases/syn_acceptance）：季度 R0/R1/R2、过程奖、年度奖（按回款比例发放、扣减预付和季度 R2）。

- 手写的标准规则经引擎计算，与独立参考实现（reference.py）逐分一致；
- 与财务给的 05_期望奖金结果.xlsx 自动核对一致；
- 10 个异常用例中 8 个数据类被拦截（另 2 个是政策未确认，由智能体的阻断问题处理）；
- 智能体按同样的规则操作时，自动先算被引用的季度 R2，结果一致、可复算。
"""
import json
import sys
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest
import yaml

CASE = Path(__file__).parent.parent / "evals" / "cases" / "syn_acceptance"
sys.path.insert(0, str(CASE))
import reference  # noqa: E402
import run_case  # noqa: E402

from bonus_core.engine.issues import EngineError  # noqa: E402

EXPECTED_XLSX = CASE / "expected" / "05_期望奖金结果.xlsx"
SHEET = {"sales": "销售明细", "projects": "项目", "participation": "项目参与", "receipts": "回款明细", "scores": "绩效评分",
         "targets": "目标值"}


@pytest.fixture(scope="module")
def results():
    return run_case.run()


def _reference():
    t = reference.load()
    reference.GROUP.update({e["employee_id"]: e["group_id"] for e in t["员工"]})
    reference.DEPT.update({g["group_id"]: g["dept_id"] for g in t["小组"]})
    return reference.compute(t)


def test_reference_matches_expected_workbook():
    """参考实现与 05 期望结果的部门、小组、个人、项目各列全部一致（独立验证 05 本身）。"""
    lines = reference.check_expected(_reference(), EXPECTED_XLSX)
    assert lines[0].endswith("不一致 0 个"), lines[:10]


def test_engine_matches_reference(results):
    ref = json.loads((CASE / "reference.json").read_text(encoding="utf-8"))
    assert run_case.diff_reference(results, ref) == []
    for s, revs in results.items():
        for rev, r in revs.items():
            for d in ("D01", "D02"):        # 部门 = 小组合计 = 个人合计
                groups = [n for n in r.nodes.values() if n.parent == f"部门:{d}"]
                people = [n for n in r.nodes.values() if n.parent in {g.id for g in groups}]
                assert sum(Decimal(g.final) for g in groups) == sum(Decimal(p.final) for p in people) \
                    == Decimal(r.nodes[f"部门:{d}"].final), (s, rev, d)
            for p in r.payouts.values():    # 应得 = 本次实发 + 暂缓 + 已发抵扣 − 超发
                assert Decimal(p.final) == Decimal(p.paid) + Decimal(p.held) + Decimal(p.offset) - Decimal(p.overpaid)


def test_annual_offsets_quarter_r2(results):
    a, q = results["ANNUAL"]["R0"].payouts["个人:E001"], results["QUARTER"]["R2"].payouts["个人:E001"]
    assert Decimal(a.offset) == Decimal("300000") + Decimal(q.paid)
    assert a.paid == "0" and Decimal(a.overpaid) == Decimal(a.offset) - Decimal(a.cap)
    assert "QUARTER/R2:个人:E001" in results["ANNUAL"]["R0"].traces["个人:E001/已发抵扣/prior"].inputs


def test_compare_finds_all_sheets(results, tmp_path):
    from bonus_agent.compare import compare_expected, report
    ms = compare_expected(results, EXPECTED_XLSX)
    assert all(m.ok for m in ms), report(ms)
    got = {(m.scheme, m.revision, m.metric): m.total for m in ms}
    assert got[("QUARTER", "R2", "应得")] == 32 and got[("ANNUAL", "R0", "超发待核对")] == 26
    # 改错一个人的期望金额，能定位出来
    wb = openpyxl.load_workbook(EXPECTED_XLSX)
    ws = wb["个人结果"]
    head = [c.value for c in ws[5]]
    ws.cell(6, head.index("应得金额") + 1).value = 1
    wb.save(tmp_path / "05.xlsx")
    m = next(m for m in compare_expected(results, tmp_path / "05.xlsx") if (m.scheme, m.revision, m.metric) == ("QUARTER", "R0", "应得"))
    assert m.matched == 31 and m.diffs[0][0] == "E001"


def _patched(case, tmp_path) -> Path:
    from bonus_agent.negcheck import patch_workbook
    out = tmp_path / run_case.XLSX.name      # 文件名要和映射里指定的一致
    patch_workbook(run_case.XLSX, out, case["patch"], SHEET)
    return out


NEG = [c for c in json.loads((CASE / "negative_cases.json").read_text(encoding="utf-8")) if c["patch"]["target"] == "data"]
CODE = {"DUPLICATE_KEY": "key_dup", "UNKNOWN_REGION": "R_REGION", "SHARE_SUM_NOT_ONE": "share_sum",
        "UNKNOWN_SALES": "link_unmatched", "SCORE_NOT_APPROVED": "score_approved", "NONPOSITIVE_TARGET": "target_dept",
        "RECEIPT_EXCEEDS_SALES": "receipt_le_sales", "SCORE_OUT_OF_RANGE": "score_range"}


@pytest.mark.parametrize("case", NEG, ids=[c["case_id"] for c in NEG])
def test_negative_cases_blocked(case, tmp_path):
    with pytest.raises(EngineError) as e:
        run_case.run(_patched(case, tmp_path))
    assert CODE[case["expected_error_code"]] in " ".join(str(i) for i in e.value.issues)


def _script(order=("ANNUAL", "QUARTER", "PROCESS"), blocking=False):
    y = lambda p: yaml.safe_load(p.read_text(encoding="utf-8"))  # noqa: E731
    dsets, maps, shared = y(CASE / "datasets.yaml"), y(CASE / "mapping.yaml"), y(CASE / "rules" / "shared.yaml")
    script = []
    for name, d in dsets.items():
        m = maps[name]
        fields = {f: {"column": m["columns"][f], **fd} for f, fd in d["fields"].items()}
        script.append({"tool": "define_dataset", "args": {"name": name, "file": m["file"], "sheet": m["sheet"],
                                                          "fields": fields, **{k: d[k] for k in ("key", "link") if k in d},
                                                          **({"filter": m["filter"]} if "filter" in m else {})}})
    script.append({"tool": "set_rules", "args": {"scheme": "共用", **shared}})
    for s in order:     # 故意先写年度：试算时引擎要先算被引用的季度
        own = y(CASE / "rules" / f"rules_{s}.yaml")
        if blocking and s == "ANNUAL":
            own.pop("offsets")
        script.append({"tool": "set_rules", "args": {"scheme": s, **{k: v for k, v in own.items()
                                                                    if k not in ("name", "version")}}})
    if blocking:
        script.append({"tool": "ask_question", "args": {"text": "跨期抵扣口径尚未确认：年度奖扣减哪些已发金额？",
                                                        "options": ["待财务确认"], "blocking": True}})
    script += [{"tool": "trial_run", "args": {"scheme": "ANNUAL"}}, {"tool": "trial_run"},
               {"tool": "finish", "args": {"summary": "季度奖、过程奖、年度奖"}}]
    return script


def test_agent_orders_schemes_by_reference(tmp_path):
    from bonus_agent.agent import run_agent
    from bonus_agent.llm import FakeLLM
    ar = run_agent(FakeLLM(_script()), [run_case.XLSX], tmp_path)
    assert ar.ok and ar.status == "OK", (ar.error, [s.result for s in ar.steps if not s.ok])
    first = next(s for s in ar.steps if s.tool == "trial_run")
    assert first.ok and "引用：QUARTER/R2" in first.result and "已发抵扣" in first.result
    ref = json.loads((CASE / "reference.json").read_text(encoding="utf-8"))
    assert run_case.diff_reference(ar.results, ref) == []
    assert ar.recheck.count("一致") == 5 and "不一致" not in ar.recheck and "失败" not in ar.recheck, ar.recheck


def test_agent_changing_quarter_retests_annual(tmp_path):
    from bonus_agent.agent import Workspace
    from bonus_agent.auto import build_inventory
    from bonus_agent.llm import FakeLLM
    ws = Workspace(FakeLLM([]), build_inventory([run_case.XLSX]))
    for step in _script()[:-3]:
        getattr(ws, step["tool"])(**step["args"])
    ws.trial_run()
    assert not ws.untested
    ws.trial_run("QUARTER")
    assert ws.untested == {"ANNUAL"}          # 季度重算后，引用它的年度要重算
    ws.set_rules(scheme="ANNUAL", offsets=[{"id": "x", "scope": "个人", "expr": 'scheme_paid("QUARTER", "R9")'}])
    from bonus_agent.agent import ToolError
    with pytest.raises(ToolError, match="R9"):
        ws.trial_run("ANNUAL")


def test_agent_blocking_question_then_next_round(tmp_path):
    """第 1 轮缺跨期抵扣口径：提问后结束本轮，不出结果；财务答复并补上规则后，第 2 轮接着做并与第 1 轮的草稿对比。"""
    from bonus_agent.agent import run_agent
    from bonus_agent.llm import FakeLLM
    out = tmp_path / "案卷"
    ar = run_agent(FakeLLM(_script(blocking=True)), [run_case.XLSX], out)
    assert ar.status == "NEEDS_CONFIRMATION" and not ar.results and set(ar.draft_results) == {"QUARTER", "PROCESS", "ANNUAL"}
    assert not list(out.glob("result*.xlsx"))
    wb = openpyxl.load_workbook(out / "待确认问题.xlsx")
    plan = {r[0]: r[1] for r in wb["智能体规划"].iter_rows(min_row=2, values_only=True)}
    assert plan["状态"].startswith("待财务答复（NEEDS_CONFIRMATION）") and "q1" in plan["状态"]
    case = json.loads((out / "case.json").read_text(encoding="utf-8"))
    assert case["rounds"][0]["status"] == "NEEDS_CONFIRMATION" and case["rounds"][0]["pending"] == ["q1"]

    q = yaml.safe_load((out / "questions.yaml").read_text(encoding="utf-8"))
    q[0]["answer"] = "扣本年度预付和季度 R2 实发，过程奖不扣"
    (out / "questions.yaml").write_text(yaml.safe_dump(q, allow_unicode=True), encoding="utf-8")
    offsets = yaml.safe_load((CASE / "rules" / "rules_ANNUAL.yaml").read_text(encoding="utf-8"))["offsets"]
    llm = FakeLLM([{"tool": "ask_question", "args": {"text": q[0]["text"], "options": ["待财务确认"]}},
                   {"tool": "set_rules", "args": {"scheme": "ANNUAL", "offsets": offsets}},
                   {"tool": "trial_run"}, {"tool": "finish", "args": {"summary": "三个方案"}}])
    ar2 = run_agent(llm, [run_case.XLSX], out)
    assert ar2.status == "OK" and ar2.round == 2, ar2.error
    assert "财务已答复" in ar2.steps[0].result               # 重复提问直接拿到答复
    task = llm.calls[0][1]["content"]
    assert "第 2 轮" in task and "扣本年度预付" in task
    ref = json.loads((CASE / "reference.json").read_text(encoding="utf-8"))
    assert run_case.diff_reference(ar2.results, ref) == []
    assert ar2.rule_changes == ["「ANNUAL」：已发抵扣修改"]
    assert (out / "result.xlsx").exists() and (out / "history" / "第1轮" / "待确认问题.xlsx").exists()
    assert not (out / "待确认问题.xlsx").exists()
    # 第 3 轮：修改数据文件（E001 的贡献评分），与第 2 轮逐个对象对比
    wb = openpyxl.load_workbook(run_case.XLSX)
    ws = wb["绩效评分"]
    ws.cell(2, 3).value = 60
    data = tmp_path / "补充材料" / run_case.XLSX.name
    data.parent.mkdir()
    wb.save(data)
    llm = FakeLLM([{"tool": "trial_run"}, {"tool": "finish", "args": {"summary": "三个方案"}}])
    ar3 = run_agent(llm, [data], out)
    assert ar3.status == "OK" and ar3.changes["修改"] == [run_case.XLSX.name]
    assert "修改文件：03_业务输入.xlsx" in llm.calls[0][1]["content"]
    changed = {(r[0], r[1], r[2]) for r in ar3.diff_rows}
    assert ("QUARTER", "R0", "个人:E001") in changed and not any(r[2].startswith("部门") for r in ar3.diff_rows)
    wb = openpyxl.load_workbook(out / "result.xlsx")
    assert "与上一轮对比" in wb.sheetnames


def test_confirm_freezes_template_and_runs_without_llm(tmp_path):
    """财务确认后冻结为模板：复算一致才冻结；模板按顺序算三个方案和修订；模板被改过拒绝计算。"""
    from bonus_agent.agent import run_agent
    from bonus_agent.case import ConfirmError, confirm
    from bonus_agent.llm import FakeLLM
    from bonus_core.cli import main as bonus_main
    from bonus_core.template import Template, TemplateError
    out = tmp_path / "案卷"
    pending = run_agent(FakeLLM(_script(blocking=True)), [run_case.XLSX], out)
    assert pending.status == "NEEDS_CONFIRMATION"
    with pytest.raises(ConfirmError, match="等财务答复"):
        confirm(out, "财务甲")
    ar = run_agent(FakeLLM(_script()), [run_case.XLSX], out, resume=False)
    assert ar.status == "OK"
    rules = out / "rules_ANNUAL.yaml"
    text = rules.read_text(encoding="utf-8")
    rules.write_text(text + "\n# 手改\n", encoding="utf-8")
    with pytest.raises(ConfirmError, match="rules_ANNUAL.yaml"):
        confirm(out, "财务甲")
    rules.write_text(text, encoding="utf-8")
    tpl = confirm(out, "财务甲", "Q3 验收")
    m = yaml.safe_load((tpl / "manifest.yaml").read_text(encoding="utf-8"))
    order = [s["name"] for s in m["schemes"]]
    assert order.index("QUARTER") < order.index("ANNUAL") and m["confirmed"]["by"] == "财务甲"   # 被引用的先算
    assert [r["name"] for r in next(s for s in m["schemes"] if s["name"] == "QUARTER")["revisions"]] == ["R1", "R2"]
    case = json.loads((out / "case.json").read_text(encoding="utf-8"))
    assert case["confirmed"]["round"] == 2 and case["rounds"][-1]["confirmed"]
    # 按模板计算（不调用大模型），与参考实现一致
    results, _ = Template.load(tpl).run([run_case.XLSX])
    ref = json.loads((CASE / "reference.json").read_text(encoding="utf-8"))
    assert run_case.diff_reference(results, ref) == []
    assert bonus_main(["run-template", "--template", str(tpl), "--out", str(tmp_path / "下期"), str(run_case.XLSX)]) == 0
    wb = openpyxl.load_workbook(tmp_path / "下期" / "result_QUARTER_R2.xlsx")
    assert dict(wb["确认信息"].iter_rows(min_row=2, values_only=True))["确认人"] == "财务甲"
    (tpl / "rules_QUARTER.yaml").write_text("# 改过\n" + (tpl / "rules_QUARTER.yaml").read_text(encoding="utf-8"),
                                            encoding="utf-8")
    with pytest.raises(TemplateError, match="被修改过"):
        Template.load(tpl)


def test_negcheck_replays_agent_rules_on_bad_data(tmp_path):
    """坏数据回归：用智能体案卷里的规则（不调用大模型）在改坏的数据上重算，8 种数据问题都要拦住。"""
    from bonus_agent import case as cs
    from bonus_agent.agent import run_agent
    from bonus_agent.llm import FakeLLM
    from bonus_agent.negcheck import replay, report
    ar = run_agent(FakeLLM(_script()), [run_case.XLSX], tmp_path)
    res = replay(cs.snapshot(ar.ws), [run_case.XLSX], NEG, SHEET)
    assert len(res) == 8 and all(r.blocked for r in res), report(res)


def test_trial_spot_check_and_missing_scope_guard(tmp_path):
    """试算结果附带每层一个对象的权重代入值；按人引用却漏写 scope 的变量，试算前必须先改。"""
    from bonus_agent.agent import ToolError, run_agent
    from bonus_agent.llm import FakeLLM
    ar = run_agent(FakeLLM(_script()), [run_case.XLSX], tmp_path)
    out = ar.ws._trial("QUARTER")
    assert "抽查「个人」" in out and "0.2 * " in out
    ar.ws.drafts["共用"]["variables"]["n_proj"].pop("scope")     # 漏写 scope：变成全公司项目数
    with pytest.raises(ToolError, match="n_proj 没写 scope"):
        ar.ws._trial("QUARTER")
    ar.ws.drafts["共用"]["variables"]["n_proj"]["scope"] = "global"   # 明确写全局：放行
    ar.ws._trial("QUARTER")
