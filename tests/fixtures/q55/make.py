"""按 Q55 的材料重建回放用的输入文件（原件在财务电脑上加密，不进仓库）：
- 01_计算规则说明.docx：模型在 Q55 里引用过的原文段落（回放时登记参数、摘条款要能在原文里找到）；
- 02_订单与回款.xlsx、03_人员与费率.xlsx：取自期望结果文件的「原始数据」页，和 Q55 原件一致。"""
import json
from pathlib import Path

import docx
import openpyxl

HERE = Path(__file__).parent
EXPECTED = HERE / "04_最终期望结果.xlsx"


def _blocks(ws):
    """「原始数据」页按小标题分块：{标题: (表头, 行)}。"""
    out, title, rows = {}, None, []
    for r in ws.iter_rows(values_only=True):
        vals = [v for v in r if v is not None]
        if not vals:
            continue
        if len(vals) == 1 and isinstance(vals[0], str) and title != vals[0] and (not rows or len(rows) > 1):
            if title:
                out[title] = (rows[0], rows[1:])
            title, rows = vals[0], []
            continue
        rows.append(list(r[:len(vals)]))
    out[title] = (rows[0], rows[1:])
    return out


def make(d: Path) -> Path:
    d.mkdir(parents=True, exist_ok=True)
    doc = docx.Document()
    for p in json.loads((HERE / "rules_text.json").read_text(encoding="utf-8")):
        doc.add_paragraph(p)
    doc.save(d / "01_计算规则说明.docx")
    blocks = _blocks(openpyxl.load_workbook(EXPECTED, data_only=True)["原始数据"])
    for name, sheets in (("02_订单与回款.xlsx", ("订单", "回款流水")), ("03_人员与费率.xlsx", ("人员", "品类费率", "预发记录"))):
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        for s in sheets:
            ws = wb.create_sheet(s)
            head, rows = blocks[s]
            ws.append(head)
            for r in rows:
                ws.append(r)
        wb.save(d / name)
    return d


def expected() -> dict[str, dict]:
    """员工编号 -> {应得, 预发, 实发, 超预发}（期望文件里保存的计算值）。"""
    ws = openpyxl.load_workbook(EXPECTED, data_only=True)["人员结果"]
    out = {}
    for r in ws.iter_rows(min_row=5, values_only=True):
        if r[0] and str(r[0]).isdigit():
            out[str(r[0])] = {"应得": r[4], "预发": r[5], "实发": r[6], "超预发": r[7]}
    return out
