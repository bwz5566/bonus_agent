"""第二轮代码审查发现的问题的回归测试：嵌套聚合、上级编号不存在、Excel 复核的短路 / 文本键 / 错误计数。"""
from decimal import Decimal

import pytest

from bonus_core.data.table import Row, StdTable
from bonus_core.dsl.models import RuleSet
from bonus_core.engine.issues import EngineError
from bonus_core.engine.run import compute


def ds(key=None, link=None, **fields):
    return {"key": key, "link": link, "fields": {k: {"type": t} for k, t in fields.items()}}


def table(name, *rows):
    return StdTable(name, [Row(dict(r)) for r in rows])


def rules(**kw):
    base = {"name": "t", "version": "1", "datasets": {}, "pool": {"expr": "0"}, "allocation": []}
    return RuleSet.model_validate({**base, **kw})


def test_nested_aggregate_uses_each_rows_own_node():
    """项目层对参与记录求和，内层再对工时求和：内层只取这条参与记录自己的工时（10×1 + 30×2 = 70，不是 40×1 + 40×2）。"""
    rs = rules(datasets={"项目": ds("project_id", project_id="id"),
                         "参与记录": ds("part_id", part_id="id", project_id="id", coef="number"),
                         "工时": ds(link="part_id", part_id="id", hours="number")},
               pool={"expr": 'sum(sum(hours, of: "工时") * coef, of: "参与记录")'},
               allocation=[{"level": "项目", "key": "project_id", "weight": "1"},
                           {"level": "参与记录", "key": "part_id", "parent_key": "project_id", "weight": "coef"}])
    data = {"项目": table("项目", {"project_id": "P1"}),
            "参与记录": table("参与记录", {"part_id": "A", "project_id": "P1", "coef": Decimal(1)},
                          {"part_id": "B", "project_id": "P1", "coef": Decimal(2)}),
            "工时": table("工时", {"part_id": "A", "hours": Decimal(10)}, {"part_id": "B", "hours": Decimal(30)})}
    assert Decimal(compute(rs, data).pool) == Decimal("70.00")


def test_parent_not_in_entity_table_is_blocked():
    """有小组实体表时，员工引用不存在的小组要报错，不能凭空建一个小组再分钱。"""
    rs = rules(datasets={"小组": ds("group_id", group_id="id"), "员工": ds("emp_id", emp_id="id", group_id="id", w="number")},
               pool={"expr": "100"},
               allocation=[{"level": "小组", "key": "group_id", "weight": "1"},
                           {"level": "员工", "key": "emp_id", "parent_key": "group_id", "weight": "w"}])
    data = {"小组": table("小组", {"group_id": "G1"}),
            "员工": table("员工", {"emp_id": "E1", "group_id": "G1", "w": Decimal(1)},
                        {"emp_id": "E2", "group_id": "MISSING", "w": Decimal(1)})}
    with pytest.raises(EngineError) as ei:
        compute(rs, data)
    assert "MISSING" in str(ei.value.issues[0])
    # 没有小组实体表时，小组由员工的 group_id 推导出来（原有行为）
    rs2 = rules(datasets={"员工": rs.datasets["员工"]}, pool={"expr": "100"}, allocation=rs.allocation)
    res = compute(rs2, {"员工": data["员工"]})
    assert "小组:MISSING" in res.nodes


def test_negcheck_uses_saved_header_row(tmp_path):
    """表头在第 4 行：改字段按列名找到列，复制的是第一条数据行（插在数据后面），不是标题行。"""
    import openpyxl

    from bonus_agent.negcheck import patch_workbook
    from bonus_core.data.table_detect import table_at
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "工单明细"
    ws.append(["客服工单明细"])
    ws.append(["2026年9月"])
    ws.append([])
    ws.append(["工单编号", "工号", "耗时（分钟）"])
    ws.append(["T1", "E1", 30])
    ws.append(["T2", "E2", 45])
    ws.append(["合计", None, 75])
    src, out = tmp_path / "a.xlsx", tmp_path / "b.xlsx"
    wb.save(src)
    patch_workbook(src, out, {"op": "set", "table": "工单明细", "index": 1, "field": "耗时（分钟）", "value": "-5"}, header_rows=[4])
    assert [v for _, v in table_at(out, "工单明细", [4]).rows][1] == ["T2", "E2", -5]
    patch_workbook(src, out, {"op": "set", "table": "工单明细", "index": 0, "field": "handle_minutes", "value": 0},
                   columns={"handle_minutes": "耗时（分钟）"})               # 标准字段名 + 自动识别表头
    assert table_at(out, "工单明细", [4]).rows[0][1] == ["T1", "E1", 0]
    patch_workbook(src, out, {"op": "append_copy", "table": "工单明细", "index": 0}, header_rows=[4])
    rows = [v for _, v in table_at(out, "工单明细", [4]).rows]
    assert rows[:3] == [["T1", "E1", 30], ["T2", "E2", 45], ["T1", "E1", 30]]


def _review_book(tmp_path):
    from bonus_core.report.excel import write_excel
    rs = rules(datasets={"人员": ds("eid", eid="id", n="number", code="text")},
               tables={"code_coef": {"kind": "lookup", "mapping": {"001": 10, "1": 20}}},
               variables={"a": {"scope": "个人", "expr": "if(n == 0 or 100 / n > 10, 100, 0)"},
                          "b": {"scope": "个人", "expr": "lookup(code_coef, code)"}},
               pool={"expr": "a + b", "scope": "个人"}, allocation=[{"level": "个人", "key": "eid"}])
    data = {"人员": table("人员", {"eid": "E1", "n": Decimal(0), "code": "001"}, {"eid": "E2", "n": Decimal(20), "code": "1"})}
    res = compute(rs, data)
    assert {k: v.final for k, v in res.nodes.items() if k.startswith("个人")} == {"个人:E1": "110.00", "个人:E2": "20.00"}
    return write_excel(rs, res, tmp_path / "r.xlsx")


def _recalc(path, outdir):
    import subprocess

    import openpyxl
    subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", str(outdir), str(path)],
                   check=True, capture_output=True, timeout=120)
    return openpyxl.load_workbook(outdir / path.name, data_only=True)


def test_excel_review_keeps_short_circuit_and_text_keys(tmp_path):
    """n=0 时 or 的右边（100/n）不能算；查找表键 001 和 1 是两个文本键，不能都变成数字 1。"""
    import shutil

    import openpyxl
    path = _review_book(tmp_path)
    keys = [r[1] for r in openpyxl.load_workbook(path)["参数表"].iter_rows(values_only=True) if r and r[1] not in (None, "键")]
    assert keys == ["001", "1"]
    if not shutil.which("soffice"):
        pytest.skip("需要 LibreOffice 重算公式")
    ws = _recalc(path, tmp_path / "calc")["个人过程复核"]
    rows = {r[0]: r for r in ws.iter_rows(min_row=3, values_only=True) if r[0] in ("E1", "E2")}
    assert rows["E1"][-1] == "一致" and rows["E2"][-1] == "一致", rows
    summary = {r[0]: r[1] for r in ws.iter_rows(values_only=True) if r[0] in ("复核未通过的个数", "公式出错的单元格")}
    assert summary == {"复核未通过的个数": 0, "公式出错的单元格": 0}


def test_excel_review_counts_formula_errors(tmp_path):
    """复核公式报错（#DIV/0!、#N/A）要算作复核未通过，不能统计成 0。"""
    import shutil

    import openpyxl
    if not shutil.which("soffice"):
        pytest.skip("需要 LibreOffice 重算公式")
    path = _review_book(tmp_path)
    wb = openpyxl.load_workbook(path)
    ws = wb["个人过程复核"]
    heads = [c.value for c in ws[2]]
    ws.cell(3, heads.index("a") + 1).value = "=1/0"          # 模拟公式出错
    wb.save(path)
    calc = _recalc(path, tmp_path / "calc")["个人过程复核"]
    e1 = next(r for r in calc.iter_rows(min_row=3, values_only=True) if r[0] == "E1")
    assert e1[-1] == "公式出错"
    summary = {r[0]: r[1] for r in calc.iter_rows(values_only=True) if r[0] in ("复核未通过的个数", "公式出错的单元格")}
    assert summary["复核未通过的个数"] == 1 and summary["公式出错的单元格"] >= 2


def test_money_cannot_stay_in_empty_middle_level():
    """部门分到了钱、下面却没有员工：拦下，不能让这笔钱留在部门层没人领。"""
    rs = rules(datasets={"部门": ds("dept_id", dept_id="id", w="number"), "员工": ds("emp_id", emp_id="id", dept_id="id")},
               pool={"expr": "100"},
               allocation=[{"level": "部门", "key": "dept_id", "weight": "w"},
                           {"level": "员工", "key": "emp_id", "parent_key": "dept_id", "weight": "1"}])
    data = {"部门": table("部门", {"dept_id": "D1", "w": Decimal(1)}, {"dept_id": "D2", "w": Decimal(1)}),
            "员工": table("员工", {"emp_id": "E1", "dept_id": "D1"})}
    with pytest.raises(EngineError) as ei:
        compute(rs, data)
    assert ei.value.issues[0].code == "branch_empty" and "部门:D2" in str(ei.value.issues[0])
    data["部门"] = table("部门", {"dept_id": "D1", "w": Decimal(1)}, {"dept_id": "D2", "w": Decimal(0)})   # 权重为 0：分不到钱，不拦
    assert compute(rs, data).nodes["员工:E1"].final == "100.00"
