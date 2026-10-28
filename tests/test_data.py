from decimal import Decimal
from pathlib import Path

import openpyxl

from bonus_core.data.mapping import apply_mapping, load_mapping, suggest_heuristic
from bonus_core.data.table_detect import detect
from bonus_core.dsl.models import RuleSet
from bonus_core.engine.ledger import ledger_entries, release
from bonus_core.engine.run import compute
from bonus_core.report.excel import write_excel

EX = Path(__file__).parent.parent / "examples" / "sales_q3"


def test_detect_messy_headers():
    sheets = {s.sheet: s for s in detect(EX / "data.xlsx")}
    assert sheets["人员名单"].header_row == 3
    assert "合同金额" in sheets["销售明细"].headers
    assert len(sheets["销售明细"].totals) == 1


def test_heuristic_mapping_matches_confirmed():
    rs = RuleSet.load(EX / "rules.yaml")
    sug = suggest_heuristic(rs, detect(EX / "data.xlsx"))
    confirmed = load_mapping(EX / "mapping.yaml")
    for ds, m in confirmed.items():
        assert sug[ds].columns == m.columns, ds


def test_end_to_end_with_units_and_report(tmp_path):
    rs = RuleSet.load(EX / "rules.yaml")
    data, issues = apply_mapping(rs, [EX / "data.xlsx"], load_mapping(EX / "mapping.yaml"))
    assert not [i for i in issues if i.severity == "blocking"], issues
    # 万元已换算为元
    assert all(r.values["sales_amount"] > 100000 for r in data["销售明细"].rows)
    res = compute(rs, data)
    total = sum(r.values["sales_amount"] for r in data["销售明细"].rows)
    assert Decimal(res.pool) == (total * Decimal("0.12")).quantize(Decimal("0.01"))
    out = write_excel(rs, res, tmp_path / "r.xlsx", ledger_entries(res))
    assert "计算链" in openpyxl.load_workbook(out).sheetnames


def test_total_row_mismatch_is_flagged(tmp_path):
    wb = openpyxl.load_workbook(EX / "data.xlsx")
    ws = wb["销售明细"]
    ws.cell(ws.max_row, 5).value = 1.0
    p = tmp_path / "bad.xlsx"
    wb.save(p)
    rs = RuleSet.load(EX / "rules.yaml")
    _, issues = apply_mapping(rs, [p], load_mapping(EX / "mapping.yaml"))
    assert any(i.code == "total_mismatch" for i in issues)


def test_ledger_release_uses_original_rules():
    rs = RuleSet.load(EX / "rules.yaml")
    case = rs.tests[0].datasets
    from bonus_core.data.table import from_records
    res = compute(rs, from_records(rs, case))
    entries = ledger_entries(res)
    assert [e["node_id"] for e in entries] == ["个人:E2"]
    # 下一期 E2 的 P2 回款到 400 万（共 500 万）
    later = {**case, "销售明细": [dict(r, collected=400) if r["project_id"] == "P2" else r for r in case["销售明细"]]}
    out = release(rs, entries, from_records(rs, later))
    due = Decimal(entries[0]["due"])
    still = (due * Decimal("0.2")).quantize(Decimal("0.01"))
    assert Decimal(out[0]["release_now"]) == Decimal(entries[0]["held"]) - still
