"""参数在数据表里（品类费率、人员计奖系数）：规则要直接读表，不能把表照抄成常数，否则下一期费率变了模板悄悄算错。"""
from decimal import Decimal

import openpyxl
import pytest

from bonus_agent import facts as fx
from bonus_agent.agent import ToolError, Workspace
from bonus_agent.auto import build_inventory
from bonus_agent.llm import FakeLLM


def case(d):
    d.mkdir()
    (d / "01_规则.md").write_text("订单提成 = 合同额 × 品类提成率，提成率直接读取品类费率表；个人提成 = 订单提成合计 × 计奖系数（读人员表）。",
                                  encoding="utf-8")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "人员"
    for r in (["工号", "计奖系数"], ["001", 1], ["002", 0.8]):
        ws.append(r)
    t = wb.create_sheet("订单")
    for r in (["订单号", "工号", "品类", "合同额"], ["O1", "001", "01", 1000], ["O2", "002", "02", 1000], ["O3", "002", "01", 500]):
        t.append(r)
    t = wb.create_sheet("品类费率")
    for r in (["品类", "提成率"], ["01", 0.03], ["02", 0.05]):
        t.append(r)
    wb.save(d / "02_数据.xlsx")
    w = Workspace(FakeLLM([]), build_inventory(sorted(d.iterdir())))
    f = lambda c, t: {"column": c, "type": t}   # noqa: E731
    w.define_dataset("人员", "02_数据.xlsx", "人员", {"emp_id": f("工号", "id"), "coef": f("计奖系数", "number")},
                     header_rows=[1], key="emp_id")
    w.define_dataset("订单", "02_数据.xlsx", "订单", {"order_id": f("订单号", "id"), "emp_id": f("工号", "id"),
                                                    "cat": f("品类", "id"), "amount": f("合同额", "money")},
                     header_rows=[1], key="order_id")
    w.define_dataset("品类费率", "02_数据.xlsx", "品类费率", {"cat": f("品类", "id"), "rate": f("提成率", "number")},
                     header_rows=[1], key="cat")
    return w


def rules(w, order_expr, tables=None):
    w.set_rules(scheme="月度", tables=tables or {},
                variables={"oc": {"scope": "订单", "expr": f"round({order_expr}, 2)"},
                           "pay": {"scope": "个人", "expr": 'round(sum(oc, of: "订单") * coef, 2)'}},
                pool={"expr": "pay", "scope": "个人"}, allocation=[{"level": "个人", "key": "emp_id"}])
    return w.trial_run()


def test_order_reads_rate_from_rate_table(tmp_path):
    """订单（关联实体）按自己行上的品类编号直接读品类费率表的提成率。"""
    w = case(tmp_path / "m")
    out = rules(w, "amount * rate")
    assert "照抄" not in out
    nodes = w.results["月度"]["R0"].nodes
    assert nodes["个人:001"].final == "30.00" and nodes["个人:002"].final == "52.00"     # (50 + 15) × 0.8


def test_copied_rate_table_is_blocked_even_if_ignored(tmp_path):
    w = case(tmp_path / "m")
    out = rules(w, "amount * lookup(cat_rate, cat)",
                tables={"cat_rate": {"kind": "lookup", "mapping": {"01": 0.03, "02": 0.05}}})
    assert "查找表 cat_rate 照抄了数据表「品类费率」的 rate" in out
    assert "常数 0.03" not in out                                    # 不再另报“找不到出处”
    w.ignored = [{"ref": "查找表 cat_rate", "reason": "出处为品类费率表"}]
    assert w._open_gates(w.lint["月度"])                              # 写进 ignored 也不算处理


def test_param_read_from_data_counts_as_used(tmp_path):
    """从数据文件登记的参数（002 的计奖系数 0.8），规则直接读表时算用上了；规则文档里的参数仍要在规则里出现。"""
    w = case(tmp_path / "m")
    rules(w, "amount * rate")
    w.record_facts([{"name": "计奖系数-002", "value": "0.8", "file": "02_数据.xlsx", "quote": "002 | 0.8"}])
    read = set().union(*w.read_nums.values())
    assert fx.unused(w.facts, w.drafts, w.conflicts, w.questions, read) == []
    doc = fx.Fact(name="品类提成率", value="0.03", file="01_规则.md", quote="0.03")
    assert fx.unused([doc], w.drafts, w.conflicts, w.questions, read) == ["品类提成率（0.03）"]
