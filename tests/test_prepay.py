"""预发 / 已发抵扣要写成 offsets：写进奖金池时「发放」表把实发当成应得、已发抵扣和超发都为 0（Q55）。"""
import openpyxl
import pytest

from bonus_agent.agent import ToolError, Workspace
from bonus_agent.auto import build_inventory
from bonus_agent.llm import FakeLLM

pytestmark = pytest.mark.verify

OFFSET = [{"id": "预发", "scope": "个人", "expr": 'sum(prepay_amount, of: "预发记录")'}]


def case(d):
    d.mkdir()
    (d / "01_规则.md").write_text("应得提成 = 订单金额合计 × 10%；本次实发 = max(应得提成 − 预发合计, 0)，超预发登记 = max(预发合计 − 应得提成, 0)，只登记不追回。",
                                  encoding="utf-8")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "人员"
    for r in (["员工编号"], ["001"], ["002"], ["003"]):
        ws.append(r)
    t = wb.create_sheet("订单")
    for r in (["订单号", "员工编号", "金额"], ["O1", "001", 4000], ["O2", "002", 1360], ["O3", "003", 500]):
        t.append(r)
    t = wb.create_sheet("预发记录")
    for r in (["预发单号", "员工编号", "预发金额"], ["A1", "001", 100], ["A2", "001", 50], ["A3", "002", 150]):
        t.append(r)
    wb.save(d / "02_数据.xlsx")
    w = Workspace(FakeLLM([]), build_inventory(sorted(d.iterdir())))
    f = lambda c, t: {"column": c, "type": t}   # noqa: E731
    w.define_dataset("人员", "02_数据.xlsx", "人员", {"emp_id": f("员工编号", "id")}, header_rows=[1], key="emp_id")
    w.define_dataset("订单", "02_数据.xlsx", "订单", {"order_id": f("订单号", "id"), "emp_id": f("员工编号", "id"),
                                                    "amount": f("金额", "money")}, header_rows=[1], key="order_id")
    w.define_dataset("预发记录", "02_数据.xlsx", "预发记录", {"prepay_id": f("预发单号", "id"), "emp_id": f("员工编号", "id"),
                                                        "prepay_amount": {**f("预发金额", "money"), "desc": "预发金额"}},
                     header_rows=[1], link="emp_id")
    return w


EARNED = {"earned": {"scope": "个人", "expr": 'round(sum(amount, of: "订单") * 10%, 2)'}}


def test_prepay_deducted_inside_pool_is_blocked(tmp_path):
    w = case(tmp_path / "m")
    w.set_rules(scheme="S", variables={**EARNED, "pay": {"scope": "个人", "expr": 'max(earned - sum(prepay_amount, of: "预发记录"), 0)'}},
                pool={"expr": "pay", "scope": "个人"}, allocation=[{"level": "个人", "key": "emp_id"}])
    out = w.trial_run()
    assert "奖金池用 max(… − …, 0) 直接扣减了已发 / 预发金额（prepay_amount）" in out and "offsets" in out


def test_prepay_as_offset_fills_payout_columns_and_verifies_net_pay(tmp_path):
    w = case(tmp_path / "m")
    w.set_rules(scheme="S", variables=EARNED, pool={"expr": "earned", "scope": "个人"},
                allocation=[{"level": "个人", "key": "emp_id"}], offsets=OFFSET)
    out = w.trial_run()
    assert "直接扣减" not in out
    pay = w.results["S"]["R0"].payouts
    assert (pay["个人:001"].final, pay["个人:001"].offset, pay["个人:001"].paid) == ("400.00", "150.00", "250.00")
    assert (pay["个人:002"].final, pay["个人:002"].paid, pay["个人:002"].overpaid) == ("136.00", "0.00", "14.00")
    gaps = w._verify_gaps()
    assert any("个人:001（实发最高" in g for g in gaps) and any("个人:002（有超发登记" in g for g in gaps)
    net = 'max(round(sum(amount, of: "订单") * 10%, 2) - sum(prepay_amount, of: "预发记录"), 0)'
    with pytest.raises(ToolError, match="实发对不上"):
        w.verify("个人:001", 'round(sum(amount, of: "订单") * 10%, 2)', "本次实发 = max(应得提成 − 预发合计, 0)", amount="实发")
    assert "一致" in w.verify("个人:001", net, "本次实发 = max(应得提成 − 预发合计, 0)", amount="实发")
    assert "一致" in w.verify("个人:002", net, "本次实发 = max(应得提成 − 预发合计, 0)", amount="实发")
    assert not any("实发" in g for g in w._verify_gaps())
