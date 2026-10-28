"""生成评测样例的规则文档：Word 正文 + 阶梯图片，PPT 里有原生表格和一张系数图片。

规则与 examples/sales_q3/rules.yaml 一致，用它作为标准答案。
净利润一条故意写得有歧义（没说全额还是超额），用来检验模型会不会提问。

运行：python evals/cases/sales_q3/make_case.py
"""
from io import BytesIO
from pathlib import Path

import docx
from docx.shared import Cm
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Cm as PCm

HERE = Path(__file__).parent
FONT = "/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc"


def font(size):
    try:
        return ImageFont.truetype(FONT, size)
    except OSError:
        return ImageFont.load_default()


def step_chart(title, steps, xlabel, ylabel, fmt_x, fmt_y) -> bytes:
    """画一张阶梯柱状图：每档一根柱子，柱顶标注比例，柱下标注门槛。"""
    W, H = 900, 560
    im = Image.new("RGB", (W, H), "white")
    d = ImageDraw.Draw(im)
    d.text((40, 24), title, fill="black", font=font(30))
    x0, y0, x1, y1 = 110, 110, W - 60, H - 110
    d.line([(x0, y1), (x1, y1)], fill="black", width=2)
    d.line([(x0, y0), (x0, y1)], fill="black", width=2)
    top = max(v for _, v in steps) * 1.25
    bw = (x1 - x0) / len(steps)
    for i, (lo, v) in enumerate(steps):
        bx0 = x0 + i * bw + bw * 0.2
        bx1 = x0 + (i + 1) * bw - bw * 0.2
        by = y1 - (y1 - y0) * (v / top)
        d.rectangle([bx0, by, bx1, y1], fill=(68, 114, 196))
        d.text((bx0 + 8, by - 38), fmt_y(v), fill="black", font=font(26))
        d.text((bx0, y1 + 12), fmt_x(lo), fill="black", font=font(22))
    d.text((x1 - 260, y1 + 56), xlabel, fill=(80, 80, 80), font=font(22))
    d.text((20, y0 - 40), ylabel, fill=(80, 80, 80), font=font(22))
    buf = BytesIO()
    im.save(buf, "PNG")
    return buf.getvalue()


def main():
    sales_png = step_chart(
        "销售额阶梯提成比例（全额累进）", [(0, 0.0), (10_000_000, 0.10), (20_000_000, 0.12)],
        "部门季度销售额（达到即适用）", "提成比例",
        lambda x: "1000万以下" if x == 0 else f"≥{x // 10000}万", lambda v: f"{v:.0%}")
    margin_png = step_chart(
        "毛利率系数", [(0, 0.8), (0.2, 1.0), (0.35, 1.2)], "单个合同毛利率（达到即适用）", "系数",
        lambda x: "20%以下" if x == 0 else f"≥{x:.0%}", lambda v: f"×{v:.1f}")

    d = docx.Document()
    d.add_heading("销售部 2026 年第三季度激励方案", 0)
    d.add_heading("一、激励总包", 1)
    d.add_paragraph("部门激励总包由销售额提成和净利润提成两部分相加构成。")
    d.add_paragraph("1. 销售额提成：以部门本季度合同销售额合计为基数，按下图阶梯比例计提。")
    d.add_picture(BytesIO(sales_png), width=Cm(14))
    d.add_paragraph("2. 净利润提成：部门本季度净利润达到1千万元的，按5%计提。")
    d.add_heading("二、分配到小组", 1)
    d.add_paragraph("激励总包按各小组的部门绩效评分占比分配到小组。")
    d.add_heading("三、分配到个人", 1)
    d.add_paragraph("小组奖金在组内按以下四项指标的组内占比加权分配：工龄占20%，付出程度评分占20%，"
                    "折算销售额占40%，处理的项目数量占20%。")
    d.add_paragraph("折算销售额 = 各合同销售额 × 区域系数 × 新项目系数 × 毛利率系数，区域系数和毛利率系数见附件 PPT。"
                    "新项目系数为1.2，老项目为1。")
    d.add_heading("四、发放", 1)
    d.add_paragraph("未收款不发放：个人奖金按其负责合同的未回款比例暂缓发放，回款后按本方案补发。")
    d.save(HERE / "激励方案.docx")

    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "附件一：区域系数"
    rows = [("区域", "系数"), ("华东", "1.0"), ("华南", "0.9"), ("西北", "1.2"), ("其他区域", "1.0")]
    tbl = s.shapes.add_table(len(rows), 2, PCm(4), PCm(5), PCm(14), PCm(8)).table
    for i, r in enumerate(rows):
        for j, v in enumerate(r):
            tbl.cell(i, j).text = v
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "附件二：毛利率系数"
    s.shapes.add_picture(BytesIO(margin_png), PCm(3), PCm(4.5), width=PCm(19))
    prs.save(HERE / "系数表.pptx")
    print("已生成 激励方案.docx、系数表.pptx")


if __name__ == "__main__":
    main()
