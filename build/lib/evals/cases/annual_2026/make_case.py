"""生成第二个评测样例：公司年终奖。结构与 sales_q3 完全不同，用来检验全自动模式（auto）的适应能力。

- 年终奖方案.pptx：总包规则（净利润阶梯，故意没写全额还是超额）、分配办法、绩效系数表（图片）、分期发放；
- 补充通知.docx：研发中心 KPI 按 1.1 倍计入（规则分散在两个文件里）；
- 年度数据.xlsx：填表说明（应忽略）、部门 KPI、员工花名册（带标题行和无关列）、月度利润（公司级数据，万元，带合计行）。

标准答案见同目录 rules.yaml。运行：python evals/cases/annual_2026/make_case.py
"""
import random
import sys
from io import BytesIO
from pathlib import Path

import docx
import openpyxl
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Cm as PCm, Pt

HERE = Path(__file__).parent
sys.path.insert(0, str(HERE.parent / "sales_q3"))
from make_case import font  # noqa: E402

rng = random.Random(20261231)


def grade_table_png() -> bytes:
    rows = [("绩效等级", "绩效系数"), ("A（优秀）", "1.5"), ("B（良好）", "1.2"), ("C（合格）", "1.0"), ("D（待改进）", "0.5")]
    W, H, rh = 760, 60 + 70 * len(rows), 70
    im = Image.new("RGB", (W, H), "white")
    d = ImageDraw.Draw(im)
    for i, (a, b) in enumerate(rows):
        y = 30 + i * rh
        fill = (217, 226, 243) if i == 0 else "white"
        d.rectangle([40, y, W - 40, y + rh], outline="black", width=2, fill=fill)
        d.line([(W // 2, y), (W // 2, y + rh)], fill="black", width=2)
        d.text((70, y + 20), a, fill="black", font=font(28))
        d.text((W // 2 + 30, y + 20), b, fill="black", font=font(28))
    buf = BytesIO()
    im.save(buf, "PNG")
    return buf.getvalue()


def make_ppt():
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[0])
    s.shapes.title.text = "2026 年度年终奖方案"
    s.placeholders[1].text = "人力资源部 · 财务部"

    def bullets(title, lines):
        sl = prs.slides.add_slide(prs.slide_layouts[1])
        sl.shapes.title.text = title
        tf = sl.placeholders[1].text_frame
        tf.text = lines[0]
        for ln in lines[1:]:
            tf.add_paragraph().text = ln
        for p in tf.paragraphs:
            for r in p.runs:
                r.font.size = Pt(20)

    bullets("一、年终奖总包", [
        "以公司全年净利润（各月净利润合计）为基数计提年终奖总包。",
        "全年净利润达到 3000 万元的，按 8% 计提。",
        "未达到 3000 万元的，按 5% 计提。",
    ])
    bullets("二、分配到部门", ["年终奖总包按各部门年度 KPI 得分的占比分配到部门。"])
    sl = prs.slides.add_slide(prs.slide_layouts[5])
    sl.shapes.title.text = "三、分配到个人"
    tb = sl.shapes.add_textbox(PCm(2), PCm(4), PCm(21), PCm(3)).text_frame
    tb.word_wrap = True
    tb.text = ("部门奖金在部门内按「月薪 × 绩效系数 × 本年在职月数 ÷ 12」的占比分配到个人；"
               "试用期员工不参与年终奖分配。绩效系数如下：")
    sl.shapes.add_picture(BytesIO(grade_table_png()), PCm(5), PCm(8), width=PCm(14))
    bullets("四、发放", ["个人年终奖的 70% 随春节前工资发放；其余 30% 暂缓，于次年 6 月 30 日后发放。"])
    prs.save(HERE / "年终奖方案.pptx")


def make_notice():
    d = docx.Document()
    d.add_heading("关于年终奖部门分配的补充通知", 1)
    d.add_paragraph("考虑到研发中心承担的新产品研发任务，经公司研究决定：在年终奖分配到部门时，"
                    "研发中心的 KPI 得分按 1.1 倍计入，其余部门按实际得分计入。")
    d.add_paragraph("本通知自发布之日起执行。")
    d.save(HERE / "补充通知.docx")


def make_data():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "填表说明"
    for line in ["1. 本表由人力资源部和财务部共同维护。", "2. 金额单位见各表表头。", "3. 如有疑问请联系人力资源部。"]:
        ws.append([line])

    depts = [("D01", "销售中心", 92), ("D02", "研发中心", 88), ("D03", "运营中心", 81), ("D04", "职能中心", 75)]
    ws = wb.create_sheet("部门KPI")
    ws.append(["部门编码", "部门名称", "KPI得分", "负责人"])
    for code, name, score in depts:
        ws.append([code, name, score, rng.choice("张王李赵刘陈") + "总"])

    ws = wb.create_sheet("员工花名册")
    ws["A1"] = "2026 年度员工花名册"
    ws.merge_cells("A1:H1")
    ws.append(["员工编号", "姓名", "所在部门", "月薪（元）", "绩效等级", "是否试用期", "本年在职月数", "手机号"])
    names = "赵钱孙李周吴郑王冯陈褚卫蒋沈韩杨朱秦尤许何吕施张孔曹"
    for i, ch in enumerate(names):
        code = depts[i % 4][0]
        probation = "是" if i in (5, 14, 23) else "否"
        months = 12 if i % 7 else rng.randint(3, 11)
        ws.append([f"E{2001 + i}", ch + "某", code, rng.randrange(8000, 40001, 500),
                   rng.choice("AABBBBCCCD"), probation, months, f"138{rng.randint(10_000_000, 99_999_999)}"])

    ws = wb.create_sheet("月度利润")
    ws.append(["月份", "营业收入（万元）", "净利润（万元）"])
    tot_rev = tot_np = 0
    for m in range(1, 13):
        rev, np_ = round(rng.uniform(1800, 2600), 2), round(rng.uniform(230, 340), 2)
        tot_rev, tot_np = tot_rev + rev, tot_np + np_
        ws.append([f"2026-{m:02d}", rev, np_])
    ws.append(["合计", round(tot_rev, 2), round(tot_np, 2)])
    wb.save(HERE / "年度数据.xlsx")
    print(f"已生成 年度数据.xlsx（全年净利润 {tot_np:.2f} 万元）")


if __name__ == "__main__":
    make_ppt()
    make_notice()
    make_data()
    print("已生成 年终奖方案.pptx、补充通知.docx")
