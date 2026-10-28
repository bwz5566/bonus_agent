"""按 01 制度、02 补充、04 参数和财务确认单（finance_answers.json）独立实现的参考计算，不经过引擎。

用途有两个：
1. 核对财务给的 05_期望奖金结果.xlsx 是否正确（逐个对象、逐列比对，打印差异）；
2. 给测试当对照，核对引擎按 rules_*.yaml 算出的结果。

运行：python evals/cases/syn_acceptance/reference.py [05_期望奖金结果.xlsx]
"""
from __future__ import annotations

import json
import sys
from collections import defaultdict
from datetime import date, datetime
from decimal import ROUND_DOWN, ROUND_HALF_UP, Decimal as D
from pathlib import Path

import openpyxl

HERE = Path(__file__).parent
INPUT = HERE / "inputs" / "03_业务输入.xlsx"

PERIODS = {"Q3": (date(2026, 7, 1), date(2026, 9, 30)), "FY": (date(2026, 1, 1), date(2026, 12, 31))}
REGION = {"华东": D("1.10"), "华南": D("1.00"), "华北": D("0.90"), "西部": D("1.20")}
DIFF = {1: D("1.00"), 2: D("1.10"), 3: D("1.25")}
NEW_SINCE = date(2026, 1, 1)
W = {"tenure": D("0.1"), "contribution": D("0.2"), "difficulty": D("0.15"), "sales": D("0.3"), "count": D("0.1"),
     "output": D("0.15")}
ADJ = {"R1": [("小组", "G01", D("900000"))], "R2": [("小组", "G01", D("900000")), ("个人", "E001", D("150000"))]}


def money(x) -> D:
    return D(x).quantize(D("0.01"), ROUND_HALF_UP)


def r6(x) -> D:
    return D(x).quantize(D("0.000001"), ROUND_HALF_UP)


def d(v) -> date:
    return v.date() if isinstance(v, datetime) else v


def num(v) -> D:
    return D(str(v))


def sheet(wb, name):
    rows = list(wb[name].iter_rows(values_only=True))
    head = rows[0]
    return [dict(zip(head, r)) for r in rows[1:] if any(v is not None for v in r)]


def load(path=INPUT):
    wb = openpyxl.load_workbook(path, data_only=True)
    return {n: sheet(wb, n) for n in ("部门", "小组", "员工", "项目", "销售明细", "回款明细", "项目参与", "绩效评分",
                                      "目标值", "里程碑", "历史已发")}


# ---------- 数据校验（发布前校验） ----------
def validate(t) -> list[str]:
    errs = []
    ids = [r["sales_id"] for r in t["销售明细"]]
    if len(ids) != len(set(ids)):
        errs.append("DUPLICATE_KEY")
    if any(p["region"] not in REGION for p in t["项目"]):
        errs.append("UNKNOWN_REGION")
    shares = defaultdict(D)
    for r in t["项目参与"]:
        shares[r["project_id"]] += num(r["share"])
    if any(s != 1 for s in shares.values()):
        errs.append("SHARE_SUM_NOT_ONE")
    sales = {r["sales_id"]: num(r["sales_amount"]) for r in t["销售明细"]}
    if any(r["sales_id"] not in sales for r in t["回款明细"]):
        errs.append("UNKNOWN_SALES")
    if any(r["approved"] is not True for r in t["绩效评分"]):
        errs.append("SCORE_NOT_APPROVED")
    if any(num(r["sales_target"]) <= 0 for r in t["目标值"]):
        errs.append("NONPOSITIVE_TARGET")
    recv = defaultdict(D)
    for r in t["回款明细"]:
        recv[r["sales_id"]] += num(r["receipt_amount"])
    if any(num(r["receipt_amount"]) < 0 for r in t["回款明细"]) or any(v > sales.get(k, v) for k, v in recv.items()):
        errs.append("RECEIPT_EXCEEDS_SALES")
    if any(not (0 <= num(r[c]) <= 100) for r in t["绩效评分"] for c in ("contribution_score", "output_score")):
        errs.append("SCORE_OUT_OF_RANGE")
    return errs


# ---------- 分配：最大余数，余数相同按编号升序；人工指定先扣除 ----------
def allocate(pool: D, weights: dict[str, D], fixed: dict[str, D]) -> dict[str, D]:
    rest = pool - sum(fixed.values(), D(0))
    if rest < 0:
        raise ValueError("POOL_EXCEEDED")
    free = {k: w for k, w in weights.items() if k not in fixed}
    total = sum(free.values(), D(0))
    out = dict(fixed)
    if not free:
        return out
    if total == 0:
        raise ValueError("ZERO_WEIGHT_SUM")
    cents = int(rest * 100)
    exact = {k: D(cents) * w / total for k, w in free.items()}
    floor = {k: int(v.to_integral_value(ROUND_DOWN)) for k, v in exact.items()}
    left = cents - sum(floor.values())
    for k in sorted(free, key=lambda k: (-(exact[k] - floor[k]), k))[:left]:
        floor[k] += 1
    out.update({k: D(v) / 100 for k, v in floor.items()})
    return out


def full_years(start: date, end: date) -> int:
    return end.year - start.year - ((end.month, end.day) < (start.month, start.day))


def compute(t):
    groups = {g["group_id"]: g for g in t["小组"]}
    dept_of_group = {g: r["dept_id"] for g, r in groups.items()}
    emps = {e["employee_id"]: e for e in t["员工"]}
    projs = {p["project_id"]: p for p in t["项目"]}
    target = {(r["period"], r["entity_type"], r["entity_id"]): num(r["sales_target"]) for r in t["目标值"]}
    parts = defaultdict(list)                       # 员工 -> [(项目, share)]
    for r in t["项目参与"]:
        parts[r["employee_id"]].append((r["project_id"], num(r["share"])))
    out = {}

    def project_figures(per):
        lo, hi = PERIODS[per]
        f = defaultdict(lambda: defaultdict(D))
        in_period = {}
        for s in t["销售明细"]:
            ok = lo <= d(s["recognition_date"]) <= hi
            in_period[s["sales_id"]] = (s["project_id"], ok)
            if ok:
                p = f[s["project_id"]]
                p["sales"] += num(s["sales_amount"])
                p["cost"] += num(s["cost_amount"])
                p["expense"] += num(s["expense_amount"])
        for r in t["回款明细"]:     # 只算期间内销售单、截止日前到账的回款
            pid, ok = in_period[r["sales_id"]]
            if ok and d(r["receipt_date"]) <= hi:
                f[pid]["recv"] += num(r["receipt_amount"])
        for pid, p in f.items():
            pr = projs[pid]
            margin = (p["sales"] - p["cost"]) / p["sales"] if p["sales"] else D(0)
            mc = D("0") if margin < D("0.15") else D("0.8") if margin < D("0.25") else D("1") if margin < D("0.35") else D("1.2")
            no = D("1.10") if d(pr["contract_date"]) >= NEW_SINCE else D("0.95")
            p.update(margin=margin, region=REGION[pr["region"]], newold=no, mcoef=mc,
                     eff=p["sales"] * REGION[pr["region"]] * no * mc, net=p["sales"] - p["cost"] - p["expense"])
        return f

    def milestones():
        lo, hi = PERIODS["Q3"]
        m = defaultdict(D)
        for r in t["里程碑"]:
            if r["status"] == "已验收" and lo <= d(r["acceptance_date"]) <= hi:
                pr = projs[r["project_id"]]
                no = D("1.10") if d(pr["contract_date"]) >= NEW_SINCE else D("0.95")
                m[r["project_id"]] += money(num(r["base_bonus"]) * REGION[pr["region"]] * no * DIFF[pr["difficulty"]])
        return m

    def coef(rate):
        return D("0.8") if rate < D("0.8") else D("1") if rate < 1 else D("1.1")

    def q_tier(x):
        return x * (D("0.08") if x < 10_000_000 else D("0.10") if x < 20_000_000 else D("0.12"))

    def a_tier(x):
        return (min(x, D(10_000_000)) * D("0.05") + max(min(x, D(20_000_000)) - 10_000_000, 0) * D("0.08")
                + max(x - 20_000_000, 0) * D("0.10"))

    def scheme(name, per, revs, prior_quarter_paid=None):
        pf = project_figures(per)
        ms = milestones() if name == "PROCESS" else {}
        cutoff = PERIODS[per][1]
        dept_of_proj = {p: dept_of_group[pr["group_id"]] for p, pr in projs.items()}
        dept = {}
        for did in sorted({g["dept_id"] for g in t["小组"]}):
            ps = [p for p in pf if dept_of_proj[p] == did]
            sales = sum((pf[p]["sales"] for p in ps), D(0))
            eff = sum((pf[p]["eff"] for p in ps), D(0))
            net = sum((pf[p]["net"] for p in ps), D(0))
            rate = sales / target[(per, "department", did)]
            if name == "PROCESS":
                pool = sum((v for p, v in ms.items() if dept_of_proj[p] == did), D(0))
            else:
                inc = q_tier(eff) if name == "QUARTER" else a_tier(eff)
                pool = money((inc + max(net, 0) * D("0.05")) * coef(rate))
            dept[did] = dict(sales=sales, eff=eff, net=net, rate=rate, pool=pool)
        gw = {}
        for gid, g in groups.items():
            gs = sum((pf[p]["sales"] for p in pf if projs[p]["group_id"] == gid), D(0))
            perf = num(g["perf_fy" if per == "FY" else "perf_q3"])
            gw[gid] = r6(perf / 100 * min(gs / target[(per, "group", gid)], D("1.2")))
        score = {r["employee_id"]: r for r in t["绩效评分"] if r["period"] == per and r["approved"] is True}
        ew, info = {}, {}
        for eid, e in emps.items():
            ps = parts[eid]
            share_sum = sum((s for _, s in ps), D(0))
            attrib = sum((pf[p]["sales"] * s for p, s in ps if p in pf), D(0))
            attrib_recv = sum((pf[p]["recv"] * s for p, s in ps if p in pf), D(0))
            diff = sum((projs[p]["difficulty"] * s for p, s in ps), D(0)) / share_sum / 3
            sales_n = min(attrib / target[(per, "employee", eid)], D("1.5")) / D("1.5")
            sc = score[eid]
            raw = (W["tenure"] * min(full_years(d(e["join_date"]), cutoff), 5) / 5
                   + W["contribution"] * num(sc["contribution_score"]) / 100 + W["difficulty"] * diff
                   + W["sales"] * sales_n + W["count"] * D(min(len({p for p, _ in ps}), 8)) / 8
                   + W["output"] * num(sc["output_score"]) / 100)
            ew[eid] = r6(raw)
            ratio = min(attrib_recv / attrib, D(1)) if attrib else D(0)
            info[eid] = dict(attrib=attrib, attrib_recv=attrib_recv, ratio=ratio)
        res = {}
        for rev in revs:
            fixed = defaultdict(dict)
            for lvl, k, amt in ADJ.get(rev, []):
                fixed[lvl][k] = amt
            gamt = {}
            for did, dv in dept.items():
                gamt.update(allocate(dv["pool"], {g: gw[g] for g in groups if dept_of_group[g] == did},
                                     {g: a for g, a in fixed["小组"].items() if dept_of_group[g] == did}))
            eamt = {}
            for gid, ga in gamt.items():
                members = [e for e, r in emps.items() if r["group_id"] == gid]
                eamt.update(allocate(ga, {e: ew[e] for e in members},
                                     {e: a for e, a in fixed["个人"].items() if e in members}))
            pay = {}
            for eid, final in eamt.items():
                cap = money(final * info[eid]["ratio"])
                offset = D(0)
                if name == "ANNUAL":
                    offset = sum((num(r["paid_amount"]) for r in t["历史已发"] if r["employee_id"] == eid
                                  and r["annual_bonus_year"] == 2026 and r["bonus_item"] == "年度奖金预付"), D(0))
                    offset += prior_quarter_paid[eid]
                pay[eid] = dict(final=final, cap=cap, paid=max(cap - offset, 0), held=final - cap, offset=offset,
                                overpaid=max(offset - cap, 0))
            res[rev] = dict(proj=pf, ms=ms, dept=dept, group_weight=gw, group=gamt, emp_weight=ew, emp=eamt, pay=pay, info=info)
        return res

    out["QUARTER"] = scheme("QUARTER", "Q3", ["R0", "R1", "R2"])
    out["PROCESS"] = scheme("PROCESS", "Q3", ["R0"])
    # 年度抵扣：本年度预付 + 季度 R2 最终可发（确认单 Q_OFFSET），过程奖不抵扣
    out["ANNUAL"] = scheme("ANNUAL", "FY", ["R0"], {e: p["paid"] for e, p in out["QUARTER"]["R2"]["pay"].items()})
    return out


# ---------- 与 05 期望结果比对 ----------
def check_expected(res, path) -> list[str]:
    wb = openpyxl.load_workbook(path, data_only=True)

    def rows(name):
        rs = list(wb[name].iter_rows(min_row=5, values_only=True))
        return [dict(zip(rs[0], r)) for r in rs[1:] if r[0]]

    lines, n = [], 0

    def cmp(where, exp, ours, tol=D("0.005")):
        nonlocal n
        n += 1
        if exp is None and ours == 0:
            return
        if exp is None or abs(num(exp) - ours) > tol:
            lines.append(f"{where}：期望 {exp}，参考实现 {ours}")

    for r in rows("项目计算"):
        x = res[r["方案"]]["R0"]
        pv = x["proj"][r["项目"]]
        where = f"项目 {r['方案']} {r['项目']}"
        for col, k in (("销售额", "sales"), ("成本", "cost"), ("费用", "expense"), ("净利润", "net"),
                       ("有效销售额", "eff"), ("期间销售单回款", "recv")):
            cmp(f"{where} {col}", r[col], pv[k])
        cmp(where + " 毛利率", r["毛利率"], pv["margin"], D("1e-12"))
        cmp(where + " 毛利系数", r["毛利系数"], pv["mcoef"], D("0"))
        if r["方案"] == "PROCESS":
            cmp(where + " 里程碑奖金", r["里程碑奖金"], x["ms"].get(r["项目"], D(0)))
    for r in rows("部门结果"):
        x = res[r["方案"]][r["修订"]]
        dv = x["dept"][r["部门编号"]]
        emps = [e for e, g in x["emp"].items() if x_group(e) and DEPT[x_group(e)] == r["部门编号"]]
        where = f"部门 {r['方案']} {r['修订']} {r['部门编号']}"
        cmp(where + " 销售额", r["销售额"], dv["sales"])
        cmp(where + " 净利润", r["净利润"], dv["net"])
        cmp(where + " 有效销售额", r["有效销售额"], dv["eff"])
        cmp(where + " 应得总包", r["期望应得总包"], dv["pool"])
        for col, k in (("本次可发", "paid"), ("暂缓金额", "held"), ("已发抵扣", "offset"), ("超发待核对", "overpaid")):
            cmp(f"{where} {col}", r[col], sum((x["pay"][e][k] for e in emps), D(0)))
    for r in rows("小组结果"):
        x = res[r["方案"]][r["修订"]]
        where = f"小组 {r['方案']} {r['修订']} {r['小组编号']}"
        cmp(where + " 原始权重", r["原始权重"], x["group_weight"][r["小组编号"]], D("0"))
        cmp(where + " 应得", r["期望应得"], x["group"][r["小组编号"]])
    for r in rows("个人结果"):
        x = res[r["方案"]][r["修订"]]
        e = r["员工编号"]
        where = f"个人 {r['方案']} {r['修订']} {e}"
        cmp(where + " 权重", r["原分配权重"], x["emp_weight"][e], D("0"))
        cmp(where + " 应得", r["应得金额"], x["emp"][e])
        cmp(where + " 归属销售额", r["归属销售额"], x["info"][e]["attrib"])
        cmp(where + " 归属回款额", r["归属回款额"], x["info"][e]["attrib_recv"])
        cmp(where + " 回款比例", r["回款比例"], x["info"][e]["ratio"], D("1e-12"))
        for col, k in (("可发上限", "cap"), ("本次可发", "paid"), ("暂缓金额", "held"), ("已发抵扣", "offset"),
                       ("超发待核对", "overpaid")):
            cmp(f"{where} {col}", r[col], x["pay"][e][k])
    for r in rows("汇总"):
        x = res[r["方案"]][r["修订"]]
        where = f"汇总 {r['方案']} {r['修订']}"
        cmp(where + " 应得总包", r["应得总包"], sum((v["pool"] for v in x["dept"].values()), D(0)))
        for col, k in (("本次可发", "paid"), ("暂缓金额", "held"), ("已发抵扣", "offset"), ("超发待核对", "overpaid")):
            cmp(f"{where} {col}", r[col], sum((p[k] for p in x["pay"].values()), D(0)))
    lines.insert(0, f"共核对 {n} 个数值，不一致 {len(lines)} 个")
    return lines


GROUP: dict[str, str] = {}
DEPT: dict[str, str] = {}


def x_group(e):
    return GROUP.get(e)


def main():
    t = load()
    GROUP.update({e["employee_id"]: e["group_id"] for e in t["员工"]})
    DEPT.update({g["group_id"]: g["dept_id"] for g in t["小组"]})
    errs = validate(t)
    print("数据校验：", errs or "通过")
    res = compute(t)
    for s, revs in res.items():
        for rev, x in revs.items():
            pay = x["pay"].values()
            print(f"{s} {rev}: 应得 {sum((v['pool'] for v in x['dept'].values()), D(0)):,}  本次可发 "
                  f"{sum((p['paid'] for p in pay), D(0)):,}  暂缓 {sum((p['held'] for p in pay), D(0)):,}")
    if len(sys.argv) > 1:
        for line in check_expected(res, sys.argv[1]):
            print(line)
    expected = {s: {rev: {"dept": {k: str(v["pool"]) for k, v in x["dept"].items()},
                          "group": {k: str(v) for k, v in x["group"].items()},
                          "emp": {k: {kk: str(vv) for kk, vv in p.items()} for k, p in x["pay"].items()}}
                     for rev, x in revs.items()} for s, revs in res.items()}
    (HERE / "reference.json").write_text(json.dumps(expected, ensure_ascii=False, indent=1), encoding="utf-8")


if __name__ == "__main__":
    main()
